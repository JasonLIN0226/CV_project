"""Build the Phase 1 independent-confirmation report from verified compact data.

This script reads only the locally verified JSON extracted from the locked HPC
artifacts. It does not run a model, mutate an experiment, or open new data.
"""
from __future__ import annotations

import json
import hashlib
import math
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "phase1_independent_confirmation_20261020"
FIG = OUT / "figures"
SOURCE = (
    ROOT
    / "results"
    / "phase1_independent_confirmation"
    / "local_verified_results_20261020.json"
)
REUSED_FIGURES = ROOT / "reports" / "figures"


def load_results() -> dict:
    result = json.loads(SOURCE.read_text())
    assert result["status"] == "PHASE1_ORIGINAL_STANDARD_NOT_MET"
    assert result["phase1_go"] is False
    assert result["independent_audit"] == "PASS_INDEPENDENT_FRESH_PROTOCOL_AUDIT"
    assert result["original_prediction_checks"]["gate"]["passed"] == 51
    assert result["original_prediction_checks"]["motion"]["passed"] == 51
    assert result["fixed_reference_checks"]["gate"]["passed"] == 32
    assert result["fixed_reference_checks"]["motion"]["passed"] == 32
    assert result["deletion_checks"]["gate"]["passed"] == 72
    assert result["deletion_checks"]["motion"]["passed"] == 64
    return result


def configure_plotting() -> None:
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 9,
            "axes.titlesize": 11,
            "axes.labelsize": 9,
            "axes.edgecolor": "black",
            "axes.linewidth": 0.8,
            "xtick.color": "black",
            "ytick.color": "black",
            "text.color": "black",
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "savefig.facecolor": "white",
        }
    )


def save_figure(fig: plt.Figure, name: str) -> Path:
    FIG.mkdir(parents=True, exist_ok=True)
    path = FIG / name
    fig.savefig(path, dpi=240, bbox_inches="tight")
    plt.close(fig)
    return path


def make_protocol_figure(result: dict) -> Path:
    fig, ax = plt.subplots(figsize=(10.8, 4.4))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 5.6)
    ax.axis("off")
    steps = [
        (0.1, 4.0, "Fresh seed 9137\n512 train · 96 val"),
        (2.0, 4.0, "Fit train-only\ntransforms and 2 maps"),
        (3.9, 4.0, "Train 12 models\n120 epochs / 7,680 steps"),
        (5.8, 4.0, "Lock all models\nbefore any test"),
        (7.7, 4.0, "Open 3×256 test pairs\nper task"),
        (1.0, 1.05, "Evaluate 51 original\nprediction checks / task"),
        (3.25, 1.05, "Replay 32 fixed-reference\nchecks / task"),
        (5.5, 1.05, "Run 72 deletion\nchecks / task"),
        (7.75, 1.05, "Independent raw replay\nand immutable evidence"),
    ]
    for x, y, label in steps:
        w = 1.65 if y == 4.0 else 1.88
        h = 0.86
        ax.add_patch(
            plt.Rectangle(
                (x, y),
                w,
                h,
                facecolor="white" if y == 4.0 else "0.91",
                edgecolor="black",
                linewidth=0.9,
            )
        )
        ax.text(x + w / 2, y + h / 2, label, ha="center", va="center", fontsize=7.8)
    for x1, x2 in [(1.76, 1.97), (3.66, 3.87), (5.56, 5.77), (7.46, 7.67)]:
        ax.annotate(
            "",
            xy=(x2, 4.43),
            xytext=(x1, 4.43),
            arrowprops=dict(arrowstyle="->", lw=1.1, color="black"),
        )
    for x, label in [
        (1.94, "1  Data"),
        (3.84, "2  Fit"),
        (5.74, "3  Lock"),
        (7.64, "4  Test"),
    ]:
        ax.text(x, 5.18, label, ha="center", va="bottom", fontsize=9, fontweight="bold")
    ax.text(
        5.0,
        2.85,
        "Only after every model is locked is the fresh 9137 test opened",
        ha="center",
        va="center",
        fontsize=10,
        fontweight="bold",
    )
    ax.annotate(
        "",
        xy=(5.0, 2.1),
        xytext=(5.0, 2.62),
        arrowprops=dict(arrowstyle="->", lw=1.1, color="black"),
    )
    ax.text(
        5.0,
        0.45,
        "Counts below are registered comparison counts; repeats within a pair are averaged.",
        ha="center",
        fontsize=8.2,
    )
    return save_figure(fig, "fig03_protocol.png")


def make_readout_figure(result: dict) -> Path:
    fig, axes = plt.subplots(1, 2, figsize=(10.2, 3.7), sharey=True)
    labels = ["Ordinary", "Long gap", "Distractor\nshift"]
    x = np.arange(3)
    width = 0.34
    for ax, task in zip(axes, ("gate", "motion")):
        d = result["feature_readout"][task]
        ax.bar(
            x - width / 2,
            d["historical"],
            width,
            color="0.25",
            edgecolor="black",
            label="Historical evidence",
        )
        ax.bar(
            x + width / 2,
            d["recent_only"],
            width,
            color="0.88",
            edgecolor="black",
            hatch="////",
            label="Recent only",
        )
        ax.set_xticks(x, labels)
        ax.set_ylim(0, 1.08)
        ax.set_title(task.title())
        ax.set_ylabel("Feature-readout accuracy")
        ax.grid(axis="y", color="0.86", linewidth=0.6)
        ax.set_axisbelow(True)
        for xi, value in zip(x, d["historical"]):
            ax.text(
                xi - width / 2,
                value + 0.025,
                f"{value:.0%}",
                ha="center",
                va="bottom",
                fontsize=8,
            )
        for xi, value in zip(x, d["recent_only"]):
            ax.text(
                xi + width / 2,
                value + 0.025,
                f"{value:.0%}",
                ha="center",
                va="bottom",
                fontsize=8,
            )
    handles, labels = axes[0].get_legend_handles_labels()
    fig.legend(
        handles,
        labels,
        loc="lower center",
        bbox_to_anchor=(0.5, -0.09),
        ncol=2,
        frameon=False,
        fontsize=8,
    )
    fig.suptitle("Evidence remains readable from the frozen features", y=1.02)
    fig.tight_layout()
    return save_figure(fig, "fig04_evidence_readout.png")


def make_fixed_reference_figure(result: dict) -> Path:
    records = result["fixed_reference_effects_seed_mean"]
    fig, axes = plt.subplots(1, 2, figsize=(10.3, 4.15), sharey=True)
    for ax, task in zip(axes, ("Gate", "Motion")):
        rows = [r for r in records if r["task"] == task]
        names = [f"{r['split']} h{r['horizon']}" for r in rows]
        vals = [100 * r["relative_reduction"] for r in rows]
        lows = [100 * (r["relative_reduction"] - r["low"]) for r in rows]
        highs = [100 * (r["high"] - r["relative_reduction"]) for r in rows]
        y = np.arange(len(rows))
        ax.errorbar(
            vals,
            y,
            xerr=np.vstack([lows, highs]),
            fmt="o",
            color="black",
            ecolor="black",
            capsize=3,
            markersize=5,
        )
        ax.axvline(0, color="black", linewidth=0.8, linestyle="--")
        ax.set_yticks(y, names)
        ax.set_xlabel("Lower feature MSE than fixed recent-only reference (%)")
        ax.set_title(task)
        ax.grid(axis="x", color="0.9", linewidth=0.6)
        ax.set_axisbelow(True)
    fig.suptitle("The corrected predictor improves over its frozen reference", y=1.02)
    fig.tight_layout()
    return save_figure(fig, "fig05_fixed_reference.png")


def make_deletion_score_figure(result: dict) -> Path:
    fig, ax = plt.subplots(figsize=(10.1, 4.4))
    conditions = ["Ordinary", "Long gap", "Distractor shift"]
    x = np.arange(3)
    width = 0.36
    for j, (task, shade, hatch) in enumerate(
        [("gate", "0.25", None), ("motion", "0.86", "////")]
    ):
        agg = {}
        for row in result["deletion_checks"][task]["by_slice"]:
            agg.setdefault(row["condition"], [0, 0])
            agg[row["condition"]][0] += row["passed"]
            agg[row["condition"]][1] += row["total"]
        rate = [agg[c][0] / agg[c][1] for c in conditions]
        labels = [f"{agg[c][0]}/{agg[c][1]}" for c in conditions]
        bars = ax.bar(
            x + (j - 0.5) * width,
            rate,
            width,
            color=shade,
            edgecolor="black",
            hatch=hatch,
            label=task.title(),
        )
        for bar, label in zip(bars, labels):
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                bar.get_height() + 0.018,
                label,
                ha="center",
                va="bottom",
                fontsize=8.5,
            )
    ax.set_xticks(x, conditions)
    ax.set_ylim(0, 1.08)
    ax.set_ylabel("Deletion checks passed")
    ax.set_title("Deletion support by condition")
    ax.legend(
        frameon=False, ncol=2, loc="upper center", bbox_to_anchor=(0.5, -0.16)
    )
    ax.grid(axis="y", color="0.88", linewidth=0.6)
    ax.set_axisbelow(True)
    return save_figure(fig, "fig06_deletion_by_condition.png")


def make_motion_failure_figure(result: dict) -> Path:
    failures = result["deletion_checks"]["motion_long_gap_failures"]
    labels = [
        f"s{r['seed']} · {r['score_horizon']}←{r['selection_horizon']}"
        for r in failures
    ]
    x = np.arange(len(failures))
    task_gain = 100 * np.array([r["task_relative_gain"] for r in failures])
    feature_gain = 100 * np.array([r["feature_relative_gain"] for r in failures])

    fig, ax = plt.subplots(figsize=(10.3, 4.25))
    width = 0.38
    ax.bar(
        x - width / 2,
        feature_gain,
        width,
        color="0.25",
        edgecolor="black",
        label="Feature error",
    )
    ax.bar(
        x + width / 2,
        task_gain,
        width,
        color="0.86",
        edgecolor="black",
        hatch="////",
        label="Position error",
    )
    ax.axhline(0, color="black", linewidth=0.9)
    ax.set_xticks(x, labels, rotation=28, ha="right")
    ax.set_ylim(-2.0, 12.5)
    ax.set_ylabel("Reduction vs reject-newest (%)")
    ax.set_title("Why the eight Motion long-gap checks fail")
    ax.legend(
        frameon=False, ncol=2, loc="upper center", bbox_to_anchor=(0.5, -0.18)
    )
    ax.grid(axis="y", color="0.9", linewidth=0.6)
    ax.set_axisbelow(True)
    ax.text(
        0.02,
        0.94,
        "← selection horizon  ·  score horizon →",
        transform=ax.transAxes,
        fontsize=8,
        ha="left",
    )
    fig.tight_layout()
    return save_figure(fig, "fig07_motion_failure.png")


def make_scoreboard_figure(result: dict) -> Path:
    groups = [
        ("Evidence read\nfrom features", [3, 3, 3, 3]),
        ("Original\nprediction", [51, 51, 51, 51]),
        ("Fixed-reference\nprediction", [32, 32, 32, 32]),
        ("Deletion\nreliability", [72, 72, 64, 72]),
    ]
    # Each row is Gate passed/total followed by Motion passed/total.
    labels = [g[0] for g in groups]
    gate_pct = [100 * g[1][0] / g[1][1] for g in groups]
    motion_pct = [100 * g[1][2] / g[1][3] for g in groups]
    gate_text = [f"{g[1][0]}/{g[1][1]}" for g in groups]
    motion_text = [f"{g[1][2]}/{g[1][3]}" for g in groups]

    fig, ax = plt.subplots(figsize=(10.2, 4.1))
    x = np.arange(len(groups))
    width = 0.36
    b1 = ax.bar(
        x - width / 2, gate_pct, width, color="0.25", edgecolor="black", label="Gate"
    )
    b2 = ax.bar(
        x + width / 2,
        motion_pct,
        width,
        color="0.88",
        edgecolor="black",
        hatch="////",
        label="Motion",
    )
    for bars, texts in ((b1, gate_text), (b2, motion_text)):
        for bar, text in zip(bars, texts):
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                bar.get_height() + 1.5,
                text,
                ha="center",
                va="bottom",
                fontsize=8.5,
            )
    ax.set_xticks(x, labels)
    ax.set_ylim(0, 110)
    ax.set_ylabel("Registered checks passed (%)")
    ax.set_title("Fresh confirmation scoreboard")
    ax.legend(
        frameon=False, ncol=2, loc="upper center", bbox_to_anchor=(0.5, -0.17)
    )
    ax.grid(axis="y", color="0.9", linewidth=0.6)
    ax.set_axisbelow(True)
    return save_figure(fig, "fig08_scoreboard.png")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="D9D9D9", size="5") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    font_size: float = 8.8,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns  # Ensure column property materialization.
    for idx, width in enumerate(widths):
        for cell in table.columns[idx].cells:
            cell.width = Inches(width)
    hdr = table.rows[0].cells
    for i, (cell, text) in enumerate(zip(hdr, headers)):
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "404040")
        set_cell_border(cell)
        set_cell_margins(cell)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(font_size)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, (cell, value) in enumerate(zip(cells, values)):
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, "F2F2F2")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
                )
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = RGBColor(0, 0, 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_figure(doc: Document, path: Path, caption: str, width: float = 6.7) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next = True
    para.add_run().add_picture(str(path), width=Inches(width))
    caption_para = doc.add_paragraph(caption, style="Caption")
    caption_para.alignment = WD_ALIGN_PARAGRAPH.CENTER


def body(doc: Document, text: str, bold_lead: str | None = None) -> None:
    para = doc.add_paragraph(style="Body Text")
    if bold_lead:
        run = para.add_run(bold_lead + " ")
        run.bold = True
    para.add_run(text)


def bullet(doc: Document, text: str) -> None:
    para = doc.add_paragraph(style="List Bullet")
    para.add_run(text)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    body_style = styles["Body Text"]
    body_style.font.name = "Arial"
    body_style.font.size = Pt(10.5)
    body_style.font.color.rgb = RGBColor(0, 0, 0)
    body_style.paragraph_format.space_after = Pt(6)
    body_style.paragraph_format.line_spacing = 1.1

    for style_name, size, before, after in [
        ("Title", 24, 0, 16),
        ("Subtitle", 13, 0, 10),
        ("Heading 1", 16, 16, 6),
        ("Heading 2", 12.5, 12, 4),
        ("Heading 3", 11, 9, 3),
        ("Caption", 8.8, 4, 10),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(size)
        style.font.bold = style_name != "Caption"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = style_name.startswith("Heading")

    title_pPr = styles["Title"]._element.get_or_add_pPr()
    for border in title_pPr.findall(qn("w:pBdr")):
        title_pPr.remove(border)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Phase 1 independent confirmation  |  ")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0, 0, 0)
    add_page_number(footer)


def make_report(result: dict, visual_paths: dict[str, Path]) -> tuple[Path, dict]:
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Phase 1 Independent Confirmation")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        "Carry-over of evidence in bounded visual memory\n"
        "Fresh seed 9137  |  Controlled I-JEPA study  |  20 October 2026"
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Decision: PHASE1_ORIGINAL_STANDARD_NOT_MET")
    r.bold = True
    r.font.size = Pt(13)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(
        "Gate passed every registered check. Motion passed every prediction and "
        "fixed-reference check, but 8 of 72 deletion checks failed. The independent "
        "audit passed, so this is a scientific threshold result, not an execution "
        "or evidence-completeness problem."
    )
    add_table(
        doc,
        ["Result at a glance", "Gate", "Motion", "Overall"],
        [
            ["Prediction and fixed reference", "83/83", "83/83", "166/166"],
            ["Deletion reliability", "72/72", "64/72", "136/144"],
            ["Registered checks", "155/155", "147/155", "302/310"],
        ],
        [2.55, 1.1, 1.1, 1.1],
        9.2,
    )
    doc.add_page_break()

    doc.add_heading("1  Research objective", level=1)
    body(
        doc,
        "The broader project asks which observations a memory-limited model should "
        "retain. A practical retention policy cannot be learned until historical "
        "visual evidence survives encoding, the predictor uses that evidence, and "
        "deleting a useful memory causes a larger prediction cost than deleting a "
        "nuisance memory.",
    )
    body(
        doc,
        "Phase 1 tests those prerequisites. It is not a test of a learned retention "
        "policy and does not establish transfer to V-JEPA or real video. The purpose "
        "of this final study is narrower: repeat the complete agreed test once on "
        "untouched seed 9137 data, with newly fitted preprocessing, predictors and "
        "readouts, after all models are locked.",
        bold_lead="Stage objective.",
    )
    body(
        doc,
        "Strict result. Phase 1 does not pass. Gate passes; Motion does not meet its "
        "72/72 deletion requirement. Because the original rule requires every "
        "criterion for both tasks, the project cannot move to retention-policy "
        "training under the current proposal.",
        bold_lead="Overall outcome.",
    )

    doc.add_heading("2  Experimental design", level=1)
    doc.add_heading("2.1  Fresh data and information boundary", level=2)
    body(
        doc,
        "Fresh confirmation used seed 9137. Each task contains 512 training pairs, "
        "96 validation pairs, and three separate 256-pair test conditions: ordinary, "
        "long gap, and distractor shift. Across both tasks the study contains 5,504 "
        "episodes. The 51 original prediction checks, 32 fixed-reference checks, and "
        "72 deletion checks per task were frozen before test access.",
    )
    add_table(
        doc,
        ["Split per task", "Matched pairs", "Episodes", "Role"],
        [
            ["Training", "512", "1,024", "Fit normalization, basis, readouts and predictors"],
            ["Validation", "96", "192", "Select checkpoint within the frozen schedule"],
            ["Ordinary test", "256", "512", "Familiar timing and scenery"],
            ["Long-gap test", "256", "512", "20 to 32 steps before the recent window"],
            ["Distractor-shift test", "256", "512", "Changed irrelevant scenery"],
            ["Total per task", "1,376", "2,752", "Both tasks: 5,504 episodes"],
        ],
        [1.55, 1.0, 0.8, 3.15],
        8.5,
    )
    body(
        doc,
        "No development weight was evaluated on the confirmation test. The test was "
        "generated only after all 12 new models and the two train-only observation maps "
        "were locked. The training and validation cache was encoded once and reused.",
    )

    doc.add_heading("2.2  What the tasks ask", level=2)
    add_figure(
        doc,
        visual_paths["tasks"],
        "Figure 1. Each matched pair has the same recent image but different earlier "
        "evidence, known actions, and future outcomes. Gate requires a transient cue. "
        "Motion requires the trajectory observed before occlusion. The drawings explain "
        "the task and are not model outputs.",
        6.35,
    )
    body(
        doc,
        "Ordinary uses a short gap and familiar scenery. Long gap holds the task rules "
        "fixed but delays the useful evidence by 20–32 steps. Distractor shift keeps the "
        "timing but changes irrelevant objects and colours. The recent four-frame window "
        "is identical within each pair, so a recent-only predictor cannot distinguish "
        "the two futures.",
    )
    add_figure(
        doc,
        visual_paths["conditions"],
        "Figure 2. The three conditions change the amount of missing time or the "
        "irrelevant visual context. Prediction horizons remain three and six steps "
        "after the four-frame query window.",
        6.55,
    )

    doc.add_heading("2.3  Processing and predictors", level=2)
    add_figure(
        doc,
        visual_paths["pipeline"],
        "Figure 3. Frozen I-JEPA features are cached. A small predictor reads four recent "
        "frames, legal historical entries, timestamps and action continuations. It never "
        "receives future images, hidden simulator state or evidence labels.",
        6.6,
    )
    body(
        doc,
        "The encoder is the frozen I-JEPA implementation used throughout the study. "
        "Images are processed without a center crop, spatially pooled into four cells, "
        "and represented by 5,120-dimensional features. Normalization, the 128-dimensional "
        "input basis, readouts, and the two observation maps are fitted on training data "
        "only. Full normalized future-feature MSE remains the primary prediction loss.",
    )
    body(
        doc,
        "Each task has a reference that predicts from recent frames plus legal actions "
        "and a candidate that adds a corrected historical-evidence branch. The fresh "
        "study used three seeds (17, 29 and 43), two roles, two tasks, 120 epochs and "
        "7,680 optimizer steps per model: 12 models in total.",
    )
    body(doc, "The locked Gate candidate disables the temporal extrapolation branch. "
         "The locked Motion candidate combines frozen-reference visual prediction, "
         "guarded position, and a fixed 0.5-weight training-only observation innovation "
         "map. The position readout is supervised with training data; the guard uses "
         "observed RGB. These restrictions matter: this result does not establish a "
         "fully self-supervised video memory system.")

    doc.add_heading("2.4  Tests and recorded evidence", level=2)
    add_table(
        doc,
        ["Test group", "What is compared", "What it establishes"],
        [
            [
                "Evidence readable from features",
                "Linear readout from historical evidence versus the recent-only control",
                "Useful information has not already been erased by encoding",
            ],
            [
                "Original prediction checks",
                "History benefit, paired-content intervention, and feature probes",
                "The predictor uses historical content at both horizons",
            ],
            [
                "Additional fixed-reference checks",
                "Candidate versus frozen recent-only and same-memory references",
                "Any gain is measured against a stable outside baseline",
            ],
            [
                "Critical deletion contrasts",
                "Delete useful evidence versus delete nuisance content at equal capacity",
                "Deletion cost separates useful from uninformative memories",
            ],
            [
                "Cross-horizon deletion checks",
                "Select a deletion at one horizon; score the choice at the other",
                "The deletion signal is not confined to one future horizon",
            ],
            [
                "Learned-policy advantage and fresh confirmation",
                "Policy comparison and untouched-data replication",
                "Fresh confirmation is completed here; learned policy is not tested",
            ],
        ],
        [1.55, 2.55, 2.45],
        8.3,
    )
    add_figure(
        doc,
        visual_paths["protocol"],
        "Figure 4. The execution order prevents test leakage. Fresh data and train-only "
        "transforms are fixed first; all 12 model checkpoints are locked before the 9137 "
        "test is opened. The final audit independently replays decisions from raw rows.",
        6.75,
    )

    doc.add_heading("3  Results for each test", level=1)
    doc.add_heading("3.1  Evidence readable from features", level=2)
    body(
        doc,
        "Before testing prediction, can a fixed linear readout recover the "
        "history-dependent distinction from the frozen features?",
        bold_lead="Question.",
    )
    body(
        doc,
        "Gate historical evidence is decoded at 100% under all three conditions, "
        "while the matched recent-only control stays at chance (50%). Motion history is "
        "100% in ordinary and long-gap conditions and 92.6% under distractor shift; its "
        "recent-only control is again 50%. The registered evidence-probe count is 3/3 for "
        "both tasks.",
        bold_lead="Result.",
    )
    add_figure(
        doc,
        visual_paths["readout"],
        "Figure 5. Taller bars mean the answer is easier to decode. Black bars use the "
        "historical evidence; hatched bars use only recent frames. The 50% recent-only "
        "floor is expected for matched pairs. These are classification readouts, not "
        "future-prediction scores.",
        6.65,
    )
    body(
        doc,
        "The distinction required by the tasks survives the encoder and "
        "is available in features. This enables the later prediction tests but does not "
        "show that a dynamics predictor will use it.",
        bold_lead="Interpretation.",
    )

    doc.add_heading("3.2  Original prediction checks", level=2)
    body(
        doc,
        "The original 51 checks combine three feature probes, 24 history-benefit checks "
        "and 24 paired-content intervention checks. Each task passed all 51: Gate 51/51 "
        "and Motion 51/51. The checks cover every condition, both three- and six-step "
        "horizons, all three model seeds, and the seed mean.",
    )
    body(
        doc,
        "This is a substantial change from the earlier development run. The fresh "
        "confirmation shows that the corrected predictor now uses historical content in "
        "both tasks under the registered prediction protocol.",
    )

    doc.add_heading("3.3  Additional fixed-reference checks", level=2)
    body(
        doc,
        "The 32 checks compare the candidate with a fixed recent-only reference and with "
        "a frozen same-memory reference. Both tasks pass 32/32. All intervals exclude no "
        "improvement. The seed-mean relative feature-error reductions range from 13.3% "
        "to 27.8% for Gate and 38.2% to 55.3% for Motion.",
    )
    add_figure(
        doc,
        visual_paths["fixed"],
        "Figure 6. Points to the right are lower feature error than the fixed "
        "recent-only reference. Horizontal bars show 95% paired intervals over 256 base "
        "pairs, after averaging the three model seeds. Every interval lies above zero.",
        6.65,
    )
    add_table(
        doc,
        ["Task", "Ordinary h3", "Ordinary h6", "Long gap h3", "Long gap h6", "Distractor h3", "Distractor h6"],
        [
            [
                "Gate",
                "27.8%",
                "26.6%",
                "26.5%",
                "25.0%",
                "14.1%",
                "13.3%",
            ],
            [
                "Motion",
                "53.8%",
                "55.3%",
                "43.6%",
                "40.4%",
                "38.2%",
                "38.4%",
            ],
        ],
        [1.15, 0.88, 0.88, 0.9, 0.9, 0.98, 0.98],
        8.3,
    )
    body(
        doc,
        "The candidate's historical branch provides a clear and "
        "consistent prediction advantage against a fixed reference in both tasks. This "
        "supports the predictor mechanism, but it is not by itself the memory-retention "
        "criterion.",
        bold_lead="Interpretation.",
    )

    doc.add_heading("3.4  Critical deletion contrasts across all conditions", level=2)
    body(
        doc,
        "Each deletion fork keeps exactly two historical entries and the same four recent "
        "frames. Gate compares retaining the cue with retaining a nuisance frame across "
        "three cue-age pools. Motion's registered reliability test asks whether the "
        "chosen deletion continues to help at the other prediction horizon.",
    )
    body(doc, "Deletion cost is the change in future-feature mean squared error when "
         "one stored observation is removed. The assay compares deletions with the same "
         "remaining memory capacity. Gate requires at least a 10% critical-versus-nuisance "
         "effect and a positive paired 95% lower bound for both metrics. Motion requires "
         "positive paired 95% gain lower bounds in feature and decoded-position error "
         "against uniform, FIFO, and reject-newest choices. Its 72 checks are three "
         "conditions times two horizons times three baselines times four seed summaries "
         "(three seeds and their mean).")
    add_figure(
        doc,
        visual_paths["deletion"],
        "Figure 7. Bars show registered deletion checks passed by condition. Gate passes "
        "72/72. Motion passes 24/24 ordinary, 16/24 long-gap, and 24/24 distractor checks. "
        "The eight missing Motion checks are not counted as a rounding loss: each has a "
        "95% interval that fails the positive-both-metrics rule.",
        6.55,
    )
    add_table(
        doc,
        ["Task", "Ordinary h3", "Ordinary h6", "Long gap h3", "Long gap h6", "Distractor h3", "Distractor h6", "Total"],
        [
            ["Gate", "12/12", "12/12", "12/12", "12/12", "12/12", "12/12", "72/72"],
            ["Motion", "12/12", "12/12", "8/12", "8/12", "12/12", "12/12", "64/72"],
        ],
        [0.85, 0.82, 0.82, 0.85, 0.85, 0.94, 0.94, 0.78],
        8.0,
    )
    body(
        doc,
        "Gate has complete deletion support. Motion's deletion signal is reliable in "
        "ordinary and distractor conditions and unreliable only at long gap.",
        bold_lead="Interpretation.",
    )

    doc.add_heading("3.5  Cross-horizon deletion checks", level=2)
    body(
        doc,
        "The eight failures are exactly the Motion long-gap comparisons against the "
        "strong reject-newest baseline: three seeds and their mean at each of the two "
        "score horizons. Feature error still improves by 9.4–11.1% relative to that "
        "baseline. The task-position error change is close to zero, with every 95% "
        "interval crossing zero and seed 29 slightly negative.",
    )
    add_figure(
        doc,
        visual_paths["motion_failure"],
        "Figure 8. Black bars show feature-error reduction relative to reject-newest; "
        "hatched bars show decoded position error. Feature error improves by roughly "
        "9–11%, while position improvement clusters near zero. The deletion choice found "
        "a useful feature-space direction that did not transfer reliably to position in "
        "the long-occlusion condition.",
        6.8,
    )
    add_table(
        doc,
        ["Seed / horizon", "Feature gain vs reject", "Task gain vs reject"],
        [
            ["17, score h3 / select h6", "9.4%", "+0.94%"],
            ["29, score h3 / select h6", "9.8%", "−0.51%"],
            ["43, score h3 / select h6", "10.1%", "+1.34%"],
            ["mean, score h3 / select h6", "9.8%", "+0.59%"],
            ["17, score h6 / select h3", "10.0%", "+1.04%"],
            ["29, score h6 / select h3", "10.7%", "−0.13%"],
            ["43, score h6 / select h3", "11.1%", "+1.29%"],
            ["mean, score h6 / select h3", "10.6%", "+0.74%"],
        ],
        [2.1, 1.55, 1.55],
        8.4,
    )
    add_table(doc, ["Seed and score horizon", "Absolute task gain 95% interval"],
              [[f"{r['seed']} h{r['score_horizon']}",
                f"[{r['task_task_gain_low']:.6f}, {r['task_task_gain_high']:.6f}]"]
               for r in result['deletion_checks']['motion_long_gap_failures']],
              [2.5, 3.7], 8.4)
    body(doc, "Intervals above use absolute decoded-position error reduction, not "
         "percent reduction. Every interval contains zero. This means superiority "
         "over reject-newest is not established; it does not prove that the two "
         "choices are equivalent or that historical memory has no value.")
    body(
        doc,
        "This is the only scientific failure in the fresh confirmation. The predictor "
        "passes all its own checks; the future-informed deletion choice remains useful in "
        "feature space but is not reliably better than the simple reject-newest baseline "
        "for the decoded position after long occlusion.",
        bold_lead="Cause.",
    )

    doc.add_heading("3.6  Learned-policy advantage and fresh confirmation", level=2)
    body(
        doc,
        "The new-data study itself is complete. All 12 fresh models "
        "were locked before test access; the independent audit replayed all 310 registered "
        "decision checks and reported a maximum original-statistic difference of 0.0. "
        "The audit separately checked 248,832 prediction rows, 248,832 deletion rows, "
        "18,432 paired deletion summaries, and 31,772 actual-input traces.",
        bold_lead="Fresh confirmation.",
    )
    body(
        doc,
        "No retention policy was trained. A future-informed "
        "oracle is a diagnostic for available signal, not a deployable policy. Therefore "
        "policy learnability and online memory-value estimation remain untested.",
        bold_lead="Learned policy.",
    )

    doc.add_heading("4  Overall summary and next steps", level=1)
    add_figure(
        doc,
        visual_paths["scoreboard"],
        "Figure 9. Passing counts are shown as x/y. Gate clears every group. Motion clears "
        "all prediction and fixed-reference checks and 64/72 deletion checks. Under the "
        "pre-registered conjunction, the visual majority does not replace the missing "
        "eight Motion checks.",
        6.55,
    )
    add_table(
        doc,
        ["Check group", "Gate", "Motion", "Overall"],
        [
            ["Evidence readable (included in original)", "3/3", "3/3", "6/6"],
            ["Original prediction", "51/51", "51/51", "102/102"],
            ["Fixed-reference", "32/32", "32/32", "64/64"],
            ["Deletion reliability", "72/72", "64/72", "136/144"],
            ["Total registered checks", "155/155", "147/155", "302/310"],
        ],
        [2.45, 1.25, 1.25, 1.25],
        9.0,
    )
    body(
        doc,
        "The result is 302/310 checks passed, or 97.4%. This count is descriptive; the "
        "original decision rule requires all checks, so the strict Phase 1 decision is "
        "not met. The audit was not inconclusive, and the GPU ledger stayed within the "
        "approved budget: 53,544 of 69,299 seconds.",
        bold_lead="What the numbers mean.",
    )
    doc.add_heading("4.1  What was achieved", level=2)
    body(
        doc,
        "The study provides strong fresh evidence that history is encoded, decoded, and "
        "used by the predictor. Gate is fully supported. Motion's prediction and fixed-"
        "reference behavior are also fully supported. These are meaningful prerequisites "
        "for a future retention study, but they are not a substitute for the deletion "
        "reliability requirement.",
    )
    doc.add_heading("4.2  Why Motion's deletion test failed", level=2)
    body(
        doc,
        "The failure is concentrated in long-gap Motion and in the comparison with the "
        "reject-newest baseline. The future-informed choice lowers feature MSE, yet it "
        "does not reliably lower decoded position error. A feature error and a downstream "
        "task error can rank deletions differently. The confirmation therefore shows a "
        "mismatch between the current deletion score and the position objective in "
        "this assay. This is evidence of limited task alignment, not a causal proof "
        "of a unique underlying failure mechanism.",
    )
    doc.add_heading("4.3  Can the project enter the retention-policy stage?", level=2)
    body(
        doc,
        "No, not under the registered Phase 1 rule. The Motion deletion requirement is "
        "not met, and every acceptance criterion is conjunctive. Training a retention "
        "policy now would optimize against a score that has not passed the confirmation "
        "test. The appropriate next study is a separate, validation-only examination of "
        "motion-state alignment and baseline-relative deletion value. It must not reuse "
        "these fresh test outcomes to tune a threshold.",
    )
    doc.add_heading("Evidence and reproducibility", level=1)
    body(
        doc,
        "Primary artifacts: final_outcome.json, primary_outcome.json, "
        "independent_final_audit.json, deletion/summary.json, and probes/metrics.json "
        "under /scratch/xl4358/phase1_independent_confirmation_execution_v9_cpu_incarnation/fresh. "
        "The independent audit report hash is "
        "de0638ab1f0141e15003a5bae004ccfc09df47046b5366c27754853c66ae4e09.",
    )
    body(
        doc,
        "Method references: [1] Assran et al., Self-Supervised Learning from Images with "
        "a Joint-Embedding Predictive Architecture, CVPR 2023, arXiv:2301.08243. "
        "[2] Lee et al., Set Transformer, ICML 2019. [3] Hooker et al., A Benchmark for "
        "Interpretability Methods in Deep Neural Networks, NeurIPS 2019.",
    )

    doc.add_heading("4.4  Conclusion", level=2)
    body(
        doc,
        "Phase 1 does not pass. Gate passes completely; Motion passes 147 of 155 checks "
        "and misses exactly eight long-gap deletion checks. The final report therefore "
        "records a scientific standard failure, not an execution failure. The memory "
        "mechanism is supported well enough to justify a new independently designed "
        "validation study, but not enough to begin retention-policy training.",
    )

    output = OUT / "Phase1_Independent_Confirmation_Report.docx"
    doc.save(output)
    manifest = {
        "report": str(output.relative_to(ROOT)),
        "report_sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
        "status": result["status"],
        "phase1_go": result["phase1_go"],
        "independent_audit": result["independent_audit"],
        "figure_files": [str(p.relative_to(ROOT)) for p in visual_paths.values()],
        "source": str(SOURCE.relative_to(ROOT)),
        "source_sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
        "scientific_conclusion": (
            "Gate passes every registered check. Motion passes 147/155; the eight "
            "long-gap deletion checks fail. Strict Phase 1 is not met."
        ),
        "retention_policy_started": False,
    }
    (OUT / "report_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    return output, manifest


def main() -> None:
    configure_plotting()
    result = load_results()
    OUT.mkdir(parents=True, exist_ok=True)
    FIG.mkdir(parents=True, exist_ok=True)
    visual_paths = {
        "tasks": REUSED_FIGURES / "01_tasks.png",
        "conditions": REUSED_FIGURES / "11_condition_timeline.png",
        "pipeline": REUSED_FIGURES / "02_pipeline.png",
        "protocol": make_protocol_figure(result),
        "readout": make_readout_figure(result),
        "fixed": make_fixed_reference_figure(result),
        "deletion": make_deletion_score_figure(result),
        "motion_failure": make_motion_failure_figure(result),
        "scoreboard": make_scoreboard_figure(result),
    }
    for path in visual_paths.values():
        assert path.exists(), path
    output, manifest = make_report(result, visual_paths)
    print(json.dumps(manifest, indent=2))
    print(output)


if __name__ == "__main__":
    main()
