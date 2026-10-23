"""Build a detailed, reader-friendly Phase 1 experimental walkthrough.

The report explains the purpose, data, examples, experiment procedure, acceptance
rules, and observed results. It reads only verified local result artifacts.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "phase1_experimental_walkthrough_20261020"
RESULT_PATH = (
    ROOT
    / "results"
    / "phase1_independent_confirmation"
    / "local_verified_results_20261020.json"
)
PROTOCOL_PATH = ROOT / "experiments" / "independent_confirmation" / "protocol.json"
CONFIG_PATH = ROOT / "experiments" / "independent_confirmation" / "config.json"
REUSED_FIGURES = ROOT / "reports" / "figures"
RESULT_FIGURES = (
    ROOT / "reports" / "phase1_independent_confirmation_20261020" / "figures"
)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="D9D9D9", size="5") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    for style_name, size, before, after, bold in [
        ("Normal", 10.5, 0, 6, False),
        ("Body Text", 10.5, 0, 6, False),
        ("Title", 24, 0, 16, True),
        ("Subtitle", 13, 0, 10, True),
        ("Heading 1", 16, 16, 6, True),
        ("Heading 2", 12.5, 12, 4, True),
        ("Heading 3", 11, 9, 3, True),
        ("Caption", 8.8, 4, 10, False),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(size)
        style.font.bold = bold
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1
        style.paragraph_format.keep_with_next = style_name.startswith("Heading")

    title_p_pr = styles["Title"]._element.get_or_add_pPr()
    for border in title_p_pr.findall(qn("w:pBdr")):
        title_p_pr.remove(border)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    font_size: float = 8.8,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx, width in enumerate(widths):
        for cell in table.columns[idx].cells:
            cell.width = Inches(width)

    for i, (cell, text) in enumerate(zip(table.rows[0].cells, headers)):
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "404040")
        set_cell_border(cell)
        set_cell_margins(cell)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(font_size)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    header_tr_pr = table.rows[0]._tr.get_or_add_trPr()
    header_repeat = OxmlElement("w:tblHeader")
    header_tr_pr.append(header_repeat)

    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, (cell, value) in enumerate(zip(cells, values)):
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, "F2F2F2")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT
                    if i == 0
                    else WD_ALIGN_PARAGRAPH.CENTER
                )
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = RGBColor(0, 0, 0)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_figure(doc: Document, path: Path, caption: str, width: float = 6.7) -> None:
    assert path.exists(), path
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    picture = paragraph.add_run().add_picture(str(path), width=Inches(width))
    doc_pr = picture._inline.docPr
    doc_pr.set("descr", caption)
    doc_pr.set("title", caption)
    caption_paragraph = doc.add_paragraph(caption, style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def body(doc: Document, text: str, bold_lead: str | None = None) -> None:
    paragraph = doc.add_paragraph(style="Body Text")
    if bold_lead:
        run = paragraph.add_run(bold_lead + " ")
        run.bold = True
    paragraph.add_run(text)


def load_inputs() -> tuple[dict, dict, dict]:
    result = json.loads(RESULT_PATH.read_text())
    protocol = json.loads(PROTOCOL_PATH.read_text())
    config = json.loads(CONFIG_PATH.read_text())

    assert result["status"] == "PHASE1_ORIGINAL_STANDARD_NOT_MET"
    assert result["phase1_go"] is False
    assert result["scientific_failure_not_execution_failure"] is True
    assert result["independent_audit"] == "PASS_INDEPENDENT_FRESH_PROTOCOL_AUDIT"
    assert result["data"]["fresh_models"] == 12
    assert protocol["fresh_seed"] == 9137
    assert config["decision"]["probe_accuracy"] == 0.75
    assert config["decision"]["readout_accuracy"] == 0.8
    return result, protocol, config


def set_footer(doc: Document) -> None:
    footer = doc.sections[0].footer.paragraphs[0]
    for run in list(footer.runs):
        run._element.getparent().remove(run._element)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Phase 1 Experimental Walkthrough  |  ")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0, 0, 0)
    add_page_number(footer)


def add_result_callout(
    doc: Document,
    result: str,
    pass_status: str,
    meaning: str,
) -> None:
    body(doc, result, bold_lead="Result.")
    body(doc, pass_status, bold_lead="Small-step decision.")
    body(doc, meaning, bold_lead="What this proves.")


def add_title(doc: Document, result: dict) -> None:
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Phase 1 Experimental Walkthrough and Evaluation Guide")

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        "Purpose, data, examples, experiment procedures, pass rules, and results\n"
        "Fresh seed 9137  |  Controlled I-JEPA study  |  20 October 2026"
    )

    decision = doc.add_paragraph()
    decision.alignment = WD_ALIGN_PARAGRAPH.CENTER
    decision_run = decision.add_run("Overall decision: PHASE 1 NOT PASSED")
    decision_run.bold = True
    decision_run.font.name = "Arial"
    decision_run.font.size = Pt(14)
    decision_run.font.color.rgb = RGBColor(0, 0, 0)

    body(
        doc,
        "This report explains Phase 1 for a reader who has not seen the experiments. "
        "It starts with the scientific purpose and the data, then walks through each "
        "experiment using the same structure: question, data and example, procedure, "
        "pass rule, result, and interpretation. The final section explains why a study "
        "that passed 302 of 310 checks still failed the preregistered Phase 1 gate.",
        bold_lead="Purpose.",
    )

    add_table(
        doc,
        ["Result", "Gate", "Motion", "Overall"],
        [
            ["Prediction and fixed reference", "83/83", "83/83", "166/166"],
            ["Deletion reliability", "72/72", "64/72", "136/144"],
            ["All registered checks", "155/155", "147/155", "302/310"],
        ],
        [2.6, 1.0, 1.0, 1.0],
        9.2,
    )

    body(
        doc,
        "Gate passes every registered check. Motion passes every prediction and "
        "fixed-reference check but misses eight long-gap deletion checks. The independent "
        "audit passes, so the failure is a scientific threshold result rather than an "
        "execution failure or missing evidence.",
        bold_lead="Why the final decision is negative.",
    )
    doc.add_page_break()


def add_roadmap(doc: Document) -> None:
    doc.add_heading("Reader roadmap", level=1)
    body(
        doc,
        "The report follows the logic used to define the experiments. Sections 1 and 2 "
        "establish the objective and the evaluation vocabulary. Section 3 explains the "
        "data and tasks. Section 4 explains the model pipeline and leakage controls. "
        "Sections 5 through 10 explain the six experiment groups one at a time. Section "
        "11 combines the results. Section 12 states what the study does and does not "
        "support.",
    )
    add_table(
        doc,
        ["Part", "Question answered", "Main output"],
        [
            ["1 and 2", "What must Phase 1 prove, and what counts as a pass?", "Objective and evaluation rules"],
            ["3", "What data, examples, and task conditions exist?", "Split sizes, task examples, stress conditions"],
            ["4", "How do features, predictors, locking, and test access work?", "Pipeline and leakage safeguards"],
            ["5", "Can historical evidence be read from frozen features?", "Evidence-readout result"],
            ["6", "Does the predictor use history?", "Original 51 prediction checks"],
            ["7", "Does the gain survive a fixed baseline?", "32 fixed-reference checks"],
            ["8", "Does deletion distinguish useful and nuisance memory?", "Gate 72 checks; Motion 72 checks"],
            ["9", "Does deletion transfer across horizons?", "Eight Motion failure checks"],
            ["10", "Does fresh confirmation rule out a development artifact?", "Audit and test-lock result"],
            ["11 and 12", "What is the overall decision and next step?", "302/310 result and no-go decision"],
        ],
        [1.05, 3.1, 2.55],
        8.5,
    )


def add_objective(doc: Document, result: dict) -> None:
    doc.add_heading("1  Phase 1 objective", level=1)
    body(
        doc,
        "The larger project will eventually ask which observations a memory-limited "
        "visual world model should retain. Before training that retention policy, the "
        "project must show that history is present in the visual representation, that "
        "the predictor actually uses it, and that removing useful history has a larger "
        "cost than removing nuisance content.",
    )
    body(
        doc,
        "Phase 1 tests these prerequisites on controlled tasks. It does not train a "
        "learned retention policy. It also does not establish transfer to V-JEPA or to "
        "real video. The encoder and the evidence are frozen I-JEPA features from "
        "controlled Gate and Motion trajectories.",
        bold_lead="Scope.",
    )
    body(
        doc,
        "The planned full retention stage remains out of scope until both tasks clear "
        "their frozen acceptance rules. A result that works on most checks does not "
        "clear the gate when the registered rule requires every check.",
        bold_lead="Decision boundary.",
    )
    doc.add_page_break()

    doc.add_heading("1.1  Three prerequisites", level=2)
    add_table(
        doc,
        ["Prerequisite", "Plain-language question", "Evidence required"],
        [
            [
                "Readable",
                "Can we recover the historical clue from the frozen representation?",
                "Historical and recent-only readouts must separate on matched pairs.",
            ],
            [
                "Used",
                "Does prediction improve when the legal history is present?",
                "History interventions and privileged-history tests must pass at both horizons.",
            ],
            [
                "Valuable",
                "Is deleting useful history more harmful than deleting nuisance content?",
                "At equal memory capacity, the deletion-cost ranking must favor useful content.",
            ],
        ],
        [1.2, 2.65, 2.85],
        8.7,
    )
    body(
        doc,
        "The chain is directional. A readable clue is not enough if the predictor ignores "
        "it. A prediction gain is not enough if deleting the useful clue is indistinguishable "
        "from deleting a nuisance frame. Phase 1 therefore tests the chain in stages.",
        bold_lead="Why all three matter.",
    )


def add_evaluation_rules(doc: Document, result: dict, protocol: dict, config: dict) -> None:
    doc.add_heading("2  Evaluation rules and small-step passes", level=1)
    body(
        doc,
        "The report uses x/y notation throughout. The numerator is the number of "
        "registered decision checks that pass. The denominator is the number frozen in "
        "the protocol before test access. A check is not a single image pair: one check "
        "usually aggregates 256 base pairs, with repeated seeds and clustering handled "
        "by the preregistered analysis.",
    )

    doc.add_heading("2.1  Levels of success", level=2)
    add_table(
        doc,
        ["Level", "Definition", "What the reader should conclude"],
        [
            [
                "Passed check",
                "The observed effect is in the required direction and the paired 95% interval satisfies the frozen rule.",
                "This single registered decision supports the mechanism.",
            ],
            [
                "Small-step pass",
                "Every check in one experiment block passes for both tasks under the same rule.",
                "That prerequisite is usable evidence, but the full Phase 1 gate remains open.",
            ],
            [
                "Final Phase 1 pass",
                "Gate and Motion both reach 155/155, the test is locked before access, and the independent audit passes.",
                "The project may proceed to retention-policy work under the original proposal.",
            ],
        ],
        [1.2, 3.35, 2.15],
        8.6,
    )

    doc.add_heading("2.2  How the statistics are interpreted", level=2)
    body(
        doc,
        "Prediction and fixed-reference checks use 2,000 bootstrap resamples. Deletion "
        "checks use 10,000 resamples. The bootstrap resamples base pairs, averages sibling "
        "or pool repeats within a pair, keeps each model seed separate, and averages the "
        "three seeds only for the explicit seed-mean cells.",
        bold_lead="Confidence intervals.",
    )
    body(
        doc,
        "For improvements, the paired 95% lower bound must remain above zero. Some checks "
        "also require a minimum effect size. The Gate deletion test requires a relative "
        "effect of at least 10% for both feature and task error. Motion deletion requires "
        "positive feature-error and decoded-position-error lower bounds against three "
        "baselines.",
        bold_lead="Direction and threshold.",
    )

    doc.add_page_break()
    doc.add_heading("2.3  Scoring blocks and how the 155 checks are formed", level=2)
    add_table(
        doc,
        ["Scoring block", "Checks per task", "Role"],
        [
            ["Original prediction", "51", "Includes the 3 feature-readability probes, 24 history-benefit checks, and 24 paired-content intervention checks"],
            ["Fixed reference", "32", "Candidate versus frozen recent-only and same-memory references"],
            ["Deletion reliability", "72", "Gate critical-versus-nuisance and Motion cross-horizon baseline comparisons"],
            ["Total", "155", "51 + 32 + 72"],
        ],
        [1.55, 1.1, 4.1],
        8.6,
    )
    body(
        doc,
        "The three feature-readability checks are included in the 51 prediction checks. "
        "They are explained separately because they answer a different scientific question, "
        "but they must not be added again when totaling 155.",
        bold_lead="Counting rule.",
    )

    doc.add_heading("2.4  Current small-step scoreboard", level=2)
    add_table(
        doc,
        ["Evidence block", "Gate", "Motion", "Interpretation"],
        [
            ["Feature readability", "3/3", "3/3", "Historical information survives encoding"],
            ["Original prediction", "51/51", "51/51", "The predictor uses history across conditions and horizons"],
            ["Fixed reference", "32/32", "32/32", "Prediction gains survive a stable baseline"],
            ["Deletion reliability", "72/72", "64/72", "Gate passes; Motion fails only at long gap"],
        ],
        [1.6, 0.9, 0.9, 3.55],
        8.7,
    )
    body(
        doc,
        "Gate has passed all four evidence blocks. Motion has passed readability, prediction, "
        "and fixed-reference, but has not passed deletion reliability. The final gate therefore "
        "does not pass.",
        bold_lead="Current status.",
    )


def add_data_and_tasks(doc: Document, result: dict, protocol: dict) -> None:
    doc.add_heading("3  Data, tasks, and examples", level=1)
    doc.add_heading("3.1  Dataset size and splits", level=2)
    data = result["data"]
    add_table(
        doc,
        ["Split per task", "Matched pairs", "Episodes", "Purpose"],
        [
            ["Training", str(data["train_pairs_per_task"]), "1,024", "Fit train-only normalization, basis, readouts, predictors, and maps"],
            ["Validation", str(data["validation_pairs_per_task"]), "192", "Select the earliest equal-condition validation checkpoint"],
            ["Ordinary test", "256", "512", "Short gap and familiar visual context"],
            ["Long-gap test", "256", "512", "Useful evidence appears 20 to 32 steps before the recent window"],
            ["Distractor-shift test", "256", "512", "Change irrelevant objects and colours while task rules stay fixed"],
            ["Total per task", "1,376", "2,752", "Both tasks contain 5,504 episodes"],
        ],
        [1.42, 1.0, 0.85, 3.45],
        8.4,
    )
    body(
        doc,
        "Each training or validation pair contains two sibling episodes. The two siblings share "
        "the recent four-frame query and legal actions but differ in the earlier evidence and "
        "future outcome. Pair identity is preserved through normalization, batching, bootstrap, "
        "and test reporting.",
        bold_lead="Pair structure.",
    )

    doc.add_heading("3.2  Gate and Motion examples", level=2)
    add_figure(
        doc,
        REUSED_FIGURES / "01_tasks.png",
        "Figure 1. The recent image is matched within each pair. The earlier evidence differs, "
        "and so does the future outcome. Gate uses a delayed left or right cue. Motion uses the "
        "direction observed before an occlusion.",
        6.4,
    )
    body(
        doc,
        "In Gate, the two episodes reach the same recent frame but one earlier marker appeared "
        "on the left and the other on the right. The correct future side depends on that earlier "
        "marker. A predictor that sees only the recent frame cannot distinguish the siblings.",
        bold_lead="Gate example.",
    )
    body(
        doc,
        "In Motion, a ball moves toward the centre from different directions before an occlusion. "
        "The recent frames are matched after the disappearance. The future reappearance side "
        "depends on the pre-occlusion trajectory, which lies outside the recent window.",
        bold_lead="Motion example.",
    )

    doc.add_heading("3.3  The three stress conditions", level=2)
    add_figure(
        doc,
        REUSED_FIGURES / "11_condition_timeline.png",
        "Figure 2. Ordinary changes the delay from 5 to 10 steps. Long gap delays the useful "
        "evidence to 20 to 32 steps before the recent window. Distractor shift changes irrelevant "
        "visual context while preserving the task rule.",
        6.55,
    )
    body(
        doc,
        "Ordinary places the clue 5 to 10 steps before the recent window. Long gap keeps "
        "the task rule but moves the clue to 20 to 32 steps earlier. Distractor shift keeps "
        "the timing while changing irrelevant objects and colours.",
    )
    body(
        doc,
        "The test contains three separate 256-pair conditions per task. Each condition uses "
        "256 base pairs, 512 episodes, three model seeds, two horizons, and the frozen bootstrap "
        "grid. A condition score therefore reflects many repeated paired comparisons rather "
        "than one selected example.",
        bold_lead="Condition scope.",
    )


def add_pipeline(doc: Document, result: dict, protocol: dict) -> None:
    doc.add_heading("4  Feature pipeline and test safeguards", level=1)
    add_figure(
        doc,
        RESULT_FIGURES / "fig03_protocol.png",
        "Figure 3. The fresh data and train-only transforms are fixed first. Twelve models are "
        "trained and locked before the 9137 test is generated or opened. Evaluation and the "
        "independent raw replay follow the model lock.",
        6.75,
    )
    doc.add_heading("4.1  What enters the predictor", level=2)
    body(
        doc,
        "The encoder is the frozen facebook/ijepa_vith14_1k model. Images are resized to 224 by "
        "224 without a center crop, passed through the frozen encoder, and pooled on a 2 by 2 grid. "
        "The resulting feature input has 5,120 dimensions.",
        bold_lead="Encoder.",
    )
    body(
        doc,
        "The predictor receives four recent frames, legal historical entries, timestamps, and "
        "known action continuations. It never receives future images, hidden simulator state, "
        "evidence labels, sibling identifiers, or future RGB. The target is the normalized future "
        "feature representation.",
        bold_lead="Predictor input.",
    )
    doc.add_heading("4.2  Models, losses, and selection", level=2)
    body(
        doc,
        "The study trains twelve fresh models: two tasks, two roles, and three seeds. Each task "
        "has a reference model and a candidate model for seeds 17, 29, and 43. Each model trains "
        "for 120 epochs and 7,680 optimizer steps.",
        bold_lead="Model grid.",
    )
    body(
        doc,
        "The primary loss is normalized future-feature mean squared error. Training also applies "
        "the registered pair and rank terms. The candidate uses AdamW with learning rate 0.0003, "
        "weight decay 0.0001, gradient clipping 1, five warmup epochs, and cosine decay. The "
        "checkpoint rule selects the earliest minimum validation loss across the four balanced "
        "conditions and two horizons.",
        bold_lead="Training and selection.",
    )
    doc.add_heading("4.3  Lock and fresh confirmation", level=2)
    body(
        doc,
        "Normalization, the 128-dimensional basis, all readouts, both observation maps, and the "
        "twelve model checkpoints are fitted on training or validation data. Their hashes are "
        "fixed before the test set is generated. The independent audit replays the registered "
        "decisions from raw rows without forwarding or fitting another model.",
        bold_lead="No test leakage.",
    )


def add_experiment_one(doc: Document, result: dict, config: dict) -> None:
    readout = result["feature_readout"]
    doc.add_heading("5  Experiment 1  Evidence readable from features", level=1)
    body(
        doc,
        "Can a fixed linear readout recover the history-dependent answer from frozen features?",
        bold_lead="Question.",
    )
    body(
        doc,
        "The experiment uses the three 256-pair test conditions for each task. The readout is "
        "fitted with training data only and then frozen. Historical evidence is compared with a "
        "matched recent-only control.",
        bold_lead="Data used.",
    )
    body(
        doc,
        "Take a Gate pair with identical recent frames. A readout receives the frozen feature "
        "representation associated with the earlier cue and predicts which future side follows. "
        "The recent-only control cannot see the cue, so a successful readout must exceed chance.",
        bold_lead="Example.",
    )
    body(
        doc,
        "The registered probe passes when accuracy is at least 75% and the Wilson lower bound is "
        "above 50%. The recent-only structural control must remain at approximately 50%. Gate also "
        "requires at least 80% state-decoder accuracy. Motion requires decoded-position RMSE at "
        "or below 0.05.",
        bold_lead="Pass rule.",
    )
    body(
        doc,
        "Gate reaches 100% under ordinary, long gap, and distractor shift. Motion reaches 100%, "
        "100%, and 92.6%. The recent-only control remains 50% in every condition. Each task passes "
        "3 of 3 registered probes.",
        bold_lead="Result.",
    )
    add_figure(
        doc,
        RESULT_FIGURES / "fig04_evidence_readout.png",
        "Figure 4. Historical evidence is easy to decode; the matched recent-only control lies at "
        "the 50% floor. This is a feature-readout test, not a future-prediction score.",
        5.0,
    )
    add_result_callout(
        doc,
        "Gate 3/3 and Motion 3/3.",
        "Passed for both tasks.",
        "The useful distinction survives the frozen encoder and remains available in the "
        "representation. It does not yet show that the predictor will use it.",
    )


def add_experiment_two(doc: Document, result: dict, config: dict) -> None:
    doc.add_heading("6  Experiment 2  Original prediction checks", level=1)
    body(
        doc,
        "Does the predictor produce a better future-feature prediction when the legal historical "
        "content is present, and does swapping the content change the prediction in the expected "
        "direction?",
        bold_lead="Question.",
    )
    body(
        doc,
        "The 51 checks cover three conditions, horizons of three and six steps, seeds 17, 29, and "
        "43, and the seed mean. The count contains three feature probes, 24 privileged-history "
        "benefit checks, and 24 paired-content intervention checks.",
        bold_lead="Data and check structure.",
    )
    body(
        doc,
        "Consider one Gate slice at horizon 6 for seed 17. The privileged model keeps the valid "
        "earlier cue. The wrong-evidence model receives a swapped cue. If the cue is used, the "
        "privileged run should have lower latent MSE and better pair ranking, while the "
        "wrong-evidence run should lose at least 0.05 in pair margin.",
        bold_lead="Example.",
    )
    body(
        doc,
        "A privileged-benefit check passes only if the latent and task metrics meet their "
        "registered thresholds and both paired lower bounds are positive. A content-intervention "
        "check passes only if swapping to the wrong future lowers the pair margin by at least "
        "0.05 and the lower bound remains above zero. Full-history performance may not degrade "
        "normalized MSE by more than 5% from the recent-only control.",
        bold_lead="Pass rule.",
    )
    body(
        doc,
        "Gate passes 51/51 and Motion passes 51/51. All three conditions and both horizons "
        "support the expected direction.",
        bold_lead="Result.",
    )
    add_table(
        doc,
        ["Subgroup", "Checks per task", "What it establishes"],
        [
            ["Feature probes", "3", "Historical evidence remains readable"],
            ["Privileged history benefit", "24", "Legal history improves latent and task metrics"],
            ["Paired-content intervention", "24", "Wrong evidence degrades the prediction"],
            ["Total", "51", "Predictor behavior depends on historical content"],
        ],
        [1.65, 1.1, 3.95],
        8.7,
    )
    add_result_callout(
        doc,
        "Gate 51/51 and Motion 51/51.",
        "Passed for both tasks.",
        "The predictor uses historical content across all registered conditions, horizons, and "
        "seed summaries. This supports the predictor mechanism but does not yet establish that "
        "deletion value is reliable.",
    )


def add_experiment_three(doc: Document, result: dict, config: dict) -> None:
    records = result["fixed_reference_effects_seed_mean"]
    doc.add_heading("7  Experiment 3  Fixed-reference checks", level=1)
    body(
        doc,
        "Does the candidate's historical branch remain better than a stable outside baseline, "
        "not merely better than a weak internal control?",
        bold_lead="Question.",
    )
    body(
        doc,
        "The test compares the candidate with a frozen recent-only reference and a frozen "
        "same-memory reference using the 256 base pairs in each condition. The three model seeds "
        "are retained separately and averaged only in the seed-mean checks.",
        bold_lead="Data used.",
    )
    body(
        doc,
        "Take Motion ordinary at horizon 3. The candidate sees the valid history; the fixed "
        "recent-only reference does not. The candidate's feature MSE is compared against the "
        "reference on the same base pairs rather than across unrelated samples.",
        bold_lead="Example.",
    )
    body(
        doc,
        "Every one of the 32 checks requires the appropriate improvement threshold and a "
        "paired 95% lower bound strictly above zero.",
        bold_lead="Pass rule.",
    )
    body(
        doc,
        "Gate passes 32/32 and Motion passes 32/32. Seed-mean relative feature-error reductions "
        "range from 13.3% to 27.8% for Gate and 38.2% to 55.3% for Motion.",
        bold_lead="Result.",
    )
    add_figure(
        doc,
        RESULT_FIGURES / "fig05_fixed_reference.png",
        "Figure 5. Points to the right indicate lower feature MSE than the fixed recent-only "
        "reference. All displayed confidence intervals remain above zero.",
        6.65,
    )
    add_table(
        doc,
        ["Task", "Ordinary h3", "Ordinary h6", "Long gap h3", "Long gap h6", "Distractor h3", "Distractor h6"],
        [
            [
                "Gate",
                f"{100 * records[0]['relative_reduction']:.1f}%",
                f"{100 * records[1]['relative_reduction']:.1f}%",
                f"{100 * records[2]['relative_reduction']:.1f}%",
                f"{100 * records[3]['relative_reduction']:.1f}%",
                f"{100 * records[4]['relative_reduction']:.1f}%",
                f"{100 * records[5]['relative_reduction']:.1f}%",
            ],
            [
                "Motion",
                f"{100 * records[6]['relative_reduction']:.1f}%",
                f"{100 * records[7]['relative_reduction']:.1f}%",
                f"{100 * records[8]['relative_reduction']:.1f}%",
                f"{100 * records[9]['relative_reduction']:.1f}%",
                f"{100 * records[10]['relative_reduction']:.1f}%",
                f"{100 * records[11]['relative_reduction']:.1f}%",
            ],
        ],
        [1.0, 0.84, 0.84, 0.86, 0.86, 0.96, 0.96],
        8.2,
    )
    add_result_callout(
        doc,
        "Gate 32/32 and Motion 32/32.",
        "Passed for both tasks.",
        "The historical branch provides a stable prediction advantage against the outside "
        "reference. This is still a prediction result, not the retention criterion.",
    )


def add_experiment_four(doc: Document, result: dict, config: dict) -> None:
    doc.add_heading("8  Experiment 4  Critical deletion contrasts", level=1)
    body(
        doc,
        "At equal remaining memory capacity, does deleting useful evidence produce a larger "
        "cost than deleting nuisance content?",
        bold_lead="Question.",
    )
    body(
        doc,
        "Each deletion fork keeps two historical entries and the same four recent frames. "
        "Gate compares keeping the cue with keeping a nuisance frame across three cue-age pools. "
        "Motion compares a future-feature-selected or cross-horizon choice with uniform, FIFO, "
        "and reject-newest policies across six time pools.",
        bold_lead="Data used.",
    )
    body(
        doc,
        "Imagine a two-slot history. Fork A keeps the useful cue and one nuisance frame. Fork B "
        "removes the cue and keeps two nuisance frames. The memory count is identical. A useful "
        "deletion signal should assign a larger prediction cost to Fork B.",
        bold_lead="Example.",
    )
    body(
        doc,
        "Gate has 72 checks: three conditions, two horizons, three cue pools, and four seed "
        "summaries. Each check must reach a relative critical-versus-nuisance effect of at least "
        "10% for both feature and task error, with a positive paired lower bound.",
        bold_lead="Gate pass rule.",
    )
    body(
        doc,
        "Motion also has 72 checks: three conditions, two horizons, three baselines, and four "
        "seed summaries. Each cross-horizon choice must show positive paired lower bounds for "
        "both future-feature error and decoded-position error.",
        bold_lead="Motion pass rule.",
    )
    body(
        doc,
        "Gate passes 72/72. Motion passes 24/24 ordinary, 16/24 long gap, and 24/24 distractor "
        "shift, for 64/72 overall. The eight missing checks are exactly Motion long gap against "
        "the reject-newest baseline.",
        bold_lead="Result.",
    )
    add_figure(
        doc,
        RESULT_FIGURES / "fig06_deletion_by_condition.png",
        "Figure 6. Gate passes all deletion checks. Motion passes ordinary and distractor shift, "
        "but only 16 of 24 checks pass in the long-gap condition.",
        5.2,
    )
    add_table(
        doc,
        ["Task", "Ordinary h3", "Ordinary h6", "Long gap h3", "Long gap h6", "Distractor h3", "Distractor h6", "Total"],
        [
            ["Gate", "12/12", "12/12", "12/12", "12/12", "12/12", "12/12", "72/72"],
            ["Motion", "12/12", "12/12", "8/12", "8/12", "12/12", "12/12", "64/72"],
        ],
        [0.85, 0.8, 0.8, 0.84, 0.84, 0.94, 0.94, 0.75],
        8.0,
    )
    add_result_callout(
        doc,
        "Gate 72/72 and Motion 64/72.",
        "Gate passes this block; Motion does not.",
        "Deletion cost ranks critical and nuisance memory correctly for Gate. Motion's signal is "
        "reliable in ordinary and distractor conditions but not under long occlusion.",
    )


def add_experiment_five(doc: Document, result: dict, config: dict) -> None:
    failures = result["deletion_checks"]["motion_long_gap_failures"]
    doc.add_heading("9  Experiment 5  Cross-horizon deletion transfer", level=1)
    body(
        doc,
        "Does a deletion chosen using the three-step score remain useful when the same memory "
        "choice is scored at six steps, and vice versa?",
        bold_lead="Question.",
    )
    body(
        doc,
        "This section explains a subset of the same 72 Motion deletion checks used in Experiment "
        "4. It does not add another 72 checks. The eight failed cross-horizon cells contain three "
        "seeds and their mean at both score horizons.",
        bold_lead="Data relationship.",
    )
    body(
        doc,
        "Choose a deletion because it minimizes feature error at horizon 3. Then remove that "
        "entry and score the resulting future at horizon 6. Repeat in the reverse direction. "
        "The paired comparison asks whether the cross-horizon choice beats reject-newest for "
        "both feature and decoded-position targets.",
        bold_lead="Example.",
    )
    body(
        doc,
        "Each check passes only when both feature-error and decoded-position-error paired 95% "
        "lower bounds stay above zero. The position metric cannot be omitted merely because the "
        "feature metric improves.",
        bold_lead="Pass rule.",
    )
    body(
        doc,
        "Feature error improves by 9.4% to 11.1% against reject-newest. The decoded-position "
        "change ranges from -0.5% to +1.3%. Every position-error interval includes zero, so all "
        "eight cross-horizon long-gap checks fail.",
        bold_lead="Result.",
    )
    add_figure(
        doc,
        RESULT_FIGURES / "fig07_motion_failure.png",
        "Figure 7. Feature error improves by roughly 9% to 11%, while decoded-position gains "
        "cluster near zero and several seed summaries are slightly negative.",
        6.8,
    )
    add_table(
        doc,
        ["Seed and horizon", "Feature gain", "Position gain", "Position 95% interval"],
        [
            [
                f"{row['seed']}, score h{row['score_horizon']} / select h{row['selection_horizon']}",
                f"{100 * row['feature_relative_gain']:.1f}%",
                f"{100 * row['task_relative_gain']:+.2f}%",
                f"[{row['task_task_gain_low']:.6f}, {row['task_task_gain_high']:.6f}]",
            ]
            for row in failures
        ],
        [2.35, 1.0, 1.0, 1.7],
        8.1,
    )
    add_result_callout(
        doc,
        "0/8 long-gap cross-horizon checks pass.",
        "This is the exact failure that blocks Phase 1.",
        "The future-informed deletion score captures a useful feature-space direction but does not "
        "reliably transfer to the decoded position objective after long occlusion. The study does "
        "not prove that all history is useless; it shows that the current deletion score is not "
        "reliable enough for this task-state metric.",
    )


def add_experiment_six(doc: Document, result: dict) -> None:
    audit = result["audit"]
    doc.add_heading("10  Experiment 6  Fresh confirmation and leakage audit", level=1)
    body(
        doc,
        "Does the result reproduce on untouched seed 9137 data after every model and train-only "
        "transform is locked?",
        bold_lead="Question.",
    )
    body(
        doc,
        "Fresh confirmation uses twelve new models: two tasks, two roles, and three seeds. Each "
        "model trains for 120 epochs and 7,680 steps. The train/validation cache contains 62,418 "
        "encoded frames. The test conditions use the separate 256-pair grids described above.",
        bold_lead="Data used.",
    )
    body(
        doc,
        "The test generator cannot open the 9137 test until all twelve model hashes, the two "
        "observation-map hashes, normalization, basis, and readouts match their locked records.",
        bold_lead="Procedure.",
    )
    body(
        doc,
        "The experiment passes only if the model-lock and transform-lock checks pass before test "
        "access, the registered decision replay matches the original statistics, and the audit "
        "can reconstruct the outputs without changing a model.",
        bold_lead="Pass rule.",
    )
    body(
        doc,
        "The audit passes. It replays all 310 registered decisions with a maximum original-statistic "
        "difference of 0.0. It checks 248,832 prediction rows, 248,832 deletion rows, 18,432 paired "
        "deletion summaries, and 31,772 actual-input traces. The GPU ledger remains within budget.",
        bold_lead="Result.",
    )
    add_table(
        doc,
        ["Audit item", "Count"],
        [
            ["Registered decisions replayed", str(audit["decision_checks_replayed"])],
            ["Prediction rows checked", f"{audit['prediction_scalar_rows']:,}"],
            ["Deletion rows checked", f"{audit['deletion_raw_rows']:,}"],
            ["Paired deletion summaries", f"{audit['deletion_paired_rows']:,}"],
            ["Actual-input traces", "31,772"],
        ],
        [3.9, 1.45],
        8.8,
    )
    add_result_callout(
        doc,
        "Audit PASS; all model and transform locks verified.",
        "Passed for execution integrity and fresh confirmation.",
        "The result is not a development artifact and the test was not opened early. This makes "
        "the Motion failure a valid negative scientific result rather than an inconclusive execution.",
    )


def add_summary(doc: Document, result: dict) -> None:
    doc.add_heading("11  Overall results", level=1)
    add_figure(
        doc,
        RESULT_FIGURES / "fig08_scoreboard.png",
        "Figure 8. Gate passes all four evidence blocks. Motion passes feature readability, "
        "original prediction, and fixed-reference checks, but only 64 of 72 deletion checks.",
        6.65,
    )
    add_table(
        doc,
        ["Check group", "Gate", "Motion", "Overall"],
        [
            ["Evidence readable", "3/3", "3/3", "6/6"],
            ["Original prediction", "51/51", "51/51", "102/102"],
            ["Fixed reference", "32/32", "32/32", "64/64"],
            ["Deletion reliability", "72/72", "64/72", "136/144"],
            ["All registered checks", "155/155", "147/155", "302/310"],
        ],
        [2.0, 1.0, 1.0, 1.0],
        8.7,
    )
    body(
        doc,
        "The descriptive pass rate is 97.4%. The final gate is conjunctive: both Gate and Motion "
        "must reach 155/155. Motion misses exactly eight long-gap cross-horizon deletion checks, "
        "so the project does not enter the retention-policy stage.",
        bold_lead="Decision.",
    )
    body(
        doc,
        "The result supports history encoding, predictor use, and stable prediction gains. It "
        "does not support the deletion-objective alignment required to train a retention policy "
        "for long-occlusion Motion.",
        bold_lead="What the result means.",
    )


def add_implications(doc: Document) -> None:
    doc.add_heading("12  Interpretation and next step", level=1)
    doc.add_heading("12.1  What succeeded", level=2)
    body(
        doc,
        "The frozen features contain the historical distinction. The predictor uses that history "
        "across short delay, long gap, and distractor shift. Its gain remains above a fixed "
        "reference at both horizons. Gate also passes deletion reliability.",
    )
    doc.add_heading("12.2  What failed", level=2)
    body(
        doc,
        "Motion's long-gap deletion signal is inconsistent with decoded-position performance. "
        "The score improves a feature-space error but does not reliably improve the downstream "
        "position metric. The final report should describe this as a deletion-score alignment "
        "failure within this assay, not as evidence that historical Motion information has no value.",
    )
    doc.add_heading("12.3  Next experiment", level=2)
    body(
        doc,
        "The next study should examine Motion state alignment and baseline-relative deletion "
        "value using validation data only. It should not reuse the opened fresh test outcomes to "
        "tune a threshold or choose a new weight. A deployable retention policy remains unproven "
        "and outside Phase 1.",
    )
    body(
        doc,
        "The practical conclusion for a first-time reader is that most of the scientific chain is "
        "supported, but the project's own all-checks rule is not met. Phase 1 can justify a focused "
        "follow-up study. It does not justify retention-policy training yet.",
        bold_lead="Reader takeaway.",
    )

    doc.add_heading("Evidence and reproducibility", level=1)
    body(
        doc,
        "The report reads the locally verified compact result file and the frozen independent-"
        "confirmation protocol. The primary result is stored in "
        "results/phase1_independent_confirmation/local_verified_results_20261020.json. The full "
        "scientific source is locked by the project's protocol and source lock. Raw weights, "
        "feature arrays, and large logs remain in the HPC project directories and are not "
        "duplicated in this report.",
    )


def make_manifest(output: Path, result: dict) -> dict:
    pdf_path = OUT / "Phase1_Experimental_Walkthrough_and_Evaluation_Report.pdf"
    manifest = {
        "report": str(output.relative_to(ROOT)),
        "report_sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
        "source": str(RESULT_PATH.relative_to(ROOT)),
        "source_sha256": hashlib.sha256(RESULT_PATH.read_bytes()).hexdigest(),
        "protocol": str(PROTOCOL_PATH.relative_to(ROOT)),
        "protocol_sha256": hashlib.sha256(PROTOCOL_PATH.read_bytes()).hexdigest(),
        "status": result["status"],
        "phase1_go": result["phase1_go"],
        "independent_audit": result["independent_audit"],
        "scientific_conclusion": (
            "Gate 155/155. Motion 147/155. Eight Motion long-gap "
            "cross-horizon deletion checks fail, so strict Phase 1 does not pass."
        ),
        "figures": [
            str((REUSED_FIGURES / "01_tasks.png").relative_to(ROOT)),
            str((REUSED_FIGURES / "11_condition_timeline.png").relative_to(ROOT)),
            str((RESULT_FIGURES / "fig03_protocol.png").relative_to(ROOT)),
            str((RESULT_FIGURES / "fig04_evidence_readout.png").relative_to(ROOT)),
            str((RESULT_FIGURES / "fig05_fixed_reference.png").relative_to(ROOT)),
            str((RESULT_FIGURES / "fig06_deletion_by_condition.png").relative_to(ROOT)),
            str((RESULT_FIGURES / "fig07_motion_failure.png").relative_to(ROOT)),
            str((RESULT_FIGURES / "fig08_scoreboard.png").relative_to(ROOT)),
        ],
        "page_count": 14,
        "revision": "Detailed experiment walkthrough with per-experiment examples and pass rules",
    }
    if pdf_path.exists():
        manifest["rendered_pdf"] = str(pdf_path.relative_to(ROOT))
        manifest["rendered_pdf_sha256"] = hashlib.sha256(pdf_path.read_bytes()).hexdigest()
    (OUT / "report_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    return manifest


def main() -> None:
    result, protocol, config = load_inputs()
    OUT.mkdir(parents=True, exist_ok=True)

    doc = Document()
    configure_document(doc)
    set_footer(doc)

    add_title(doc, result)
    add_roadmap(doc)
    add_objective(doc, result)
    add_evaluation_rules(doc, result, protocol, config)
    doc.add_page_break()
    add_data_and_tasks(doc, result, protocol)
    add_pipeline(doc, result, protocol)
    add_experiment_one(doc, result, config)
    add_experiment_two(doc, result, config)
    add_experiment_three(doc, result, config)
    add_experiment_four(doc, result, config)
    add_experiment_five(doc, result, config)
    add_experiment_six(doc, result)
    doc.add_page_break()
    add_summary(doc, result)
    add_implications(doc)

    output = OUT / "Phase1_Experimental_Walkthrough_and_Evaluation_Report.docx"
    doc.save(output)
    manifest = make_manifest(output, result)
    print(json.dumps(manifest, indent=2))
    print(output)


if __name__ == "__main__":
    main()
