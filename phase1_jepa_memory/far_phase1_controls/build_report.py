from pathlib import Path
import csv,json,numpy as np
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
D=Path(__file__).resolve().parent;OUT=D/'report';A=D/'native_completed/analysis/output/job_7456'
S=json.loads((A/'summary.json').read_text());F=json.loads((D/'native_completed/final/final_native_audit.json').read_text());assert F['status']=='PASS_FINAL_INDEPENDENT_144_RECORD_AUDIT'
rows=list(csv.DictReader((A/'record_summary.csv').open()));order=['Current_only','Recent_only','FIFO','Reservoir','FAR_current','Archive_reference'];labels=['Current only','Recent only','FIFO with FAR','Reservoir','FAR current','Archive reference']
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.line_spacing=1.05;doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(19);doc.styles['Title'].paragraph_format.space_after=Pt(4)
doc.styles['Subtitle'].font.size=Pt(10);doc.styles['Subtitle'].paragraph_format.space_after=Pt(8)
doc.styles['Heading 1'].font.size=Pt(11.5);doc.styles['Heading 1'].paragraph_format.space_before=Pt(8);doc.styles['Heading 1'].paragraph_format.space_after=Pt(4)
doc.styles['Caption'].font.size=Pt(9);doc.styles['Caption'].font.bold=False;doc.styles['Caption'].paragraph_format.space_after=Pt(5)
def text(t):return doc.add_paragraph(t)
def heading(t,page=False):
 p=doc.add_heading(t,level=1);p.paragraph_format.page_break_before=page;return p
def table(headers,body,widths,size=9.5):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,width in enumerate(widths):t.columns[i].width=Inches(width)
 for i,h in enumerate(headers):t.rows[0].cells[i].text=h
 for values in body:
  cells=t.add_row().cells
  for i,v in enumerate(values):cells[i].text=str(v)
 for ri,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'D9D9D9');b.append(e)
   pr.append(b);pad=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');pad.append(e)
   pr.append(pad);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'333333' if ri==0 else ('F2F2F2' if ri%2==0 else 'FFFFFF'));pr.append(shade)
   for p in c.paragraphs:
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
    for run in p.runs:run.font.name='Arial';run.font.size=Pt(size);run.font.bold=ri==0;run.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1;p.add_run().font.size=Pt(1)
 return t
def figure(n,caption,alt,width=6.8):
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(1);p.paragraph_format.keep_with_next=True
 pic=p.add_run().add_picture(str(OUT/n),width=Inches(width));pic._inline.docPr.set('descr',alt)
 doc.add_paragraph(caption,style='Caption')
def v(policy,key):return S['policies'][policy][key]['mean']
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');rr=OxmlElement('w:r');pr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'18');pr.append(sz);rr.append(pr);tt=OxmlElement('w:t');tt.text='1';rr.append(tt);field.append(rr);footer._p.append(field)
doc.add_paragraph('Historical Visual Memory in Phase1',style='Title')
doc.add_paragraph('A controlled FAR study on eight LoopNav scenes',style='Subtitle')
heading('Purpose')
text('We tested historical evidence for return prediction in a fixed FAR world model. Eight LoopNav scenes, four ABA and four ABCA, and three common seeds produced 144 returns and 10494 predictions. The 96 pilot records were preserved and 48 controls were added.')
heading('Controlled design')
text('Current only supplies one real current frame and no explicit past. Recent only stores 16 past entries and directly reads the latest 11 plus current. FIFO, Reservoir and FAR current each store 16 past entries; the same pretrained FAR reader selects 11 per query. Archive reference retains all arrived history and has greater storage.')
text('All conditions use the released far_multicue predictor, DDIM 20, stride 4 and one frame chunks, adapted from the original ten. Initial current and external actions are shared. Each policy then uses its own generated current. Return ground truth never enters memory, and evicted payloads cannot reenter the policy buffer. Retention cannot inspect future routes or queries.')
heading('Six condition results')
text('The primary score is DreamSim over the last ceil(n/4) actual predictions. We average three seeds within each scene, then weight eight scenes equally. Lower DreamSim and LPIPS and higher PSNR are better.')
body=[]
for label,p in zip(labels,order):
 q=S['policies'][p];body.append([label,f"{v(p,'dreamsim_tail'):.4f} ({q['dreamsim_tail']['seed_SD']:.4f})",f"{v(p,'dreamsim_full'):.4f}",f"{v(p,'lpips_tail'):.4f}",f"{v(p,'psnr_tail'):.3f}"])
table(['Condition','DreamSim tail\nmean (seed SD)','DreamSim\nfull','LPIPS\ntail','PSNR tail\ndB'],body,[1.65,1.8,1.05,1.1,1.1])
figure('six_condition_error.png','Figure 1  Final quarter error. Each dot is one scene averaged over three seeds; bars show the equal scene mean. Archive reference has greater storage.','Six inference conditions with their primary DreamSim means and all eight individual scene means.',width=6.3)

heading('Paired memory benefit',True)
text('Reservoir reduced final quarter DreamSim error by 56.8% relative to current only and by 44.6% relative to recent only. Both comparisons favored Reservoir in all eight scene averages and in all three seed averages. It was also 46.1% lower than FIFO, which uses the same reader as Reservoir. The declared two control criterion was met.')
body=[]
for label,b in [('Reservoir over current','Current_only'),('Reservoir over recent','Recent_only')]:
 q=S['paired_comparisons']['Reservoir_over_'+b]['dreamsim_tail'];lo,hi=q['paired_scene_bootstrap95'];body.append([label,f"{q['mean_improvement']:+.4f}",f"[{lo:.4f}, {hi:.4f}]",f"{q['exact_one_sided_signflip_p']:.5f}",f"{q['Holm_adjusted_p']:.5f}"])
table(['Comparison','Mean\nimprovement','Paired 95%\ninterval','Exact\np','Holm\np'],body,[2.0,1.0,1.7,.9,.9])
figure('primary_paired_effects.png','Figure 2  Positive values favor Reservoir. Gray dots are paired scene effects; black diamonds and intervals show the mean and 95% interval from 10000 paired bootstrap draws.','Two positive primary Reservoir effects, with paired bootstrap intervals and all eight scene effects.')
text('Exact one sided sign flip tests enumerate all 256 sign patterns across eight scenes. Holm correction covers the two primary comparisons. We required both positive lower 95% bounds and both adjusted p values below 0.05. These thresholds were fixed before the new controls ran.')
heading('Relevance did not improve retention')
text('FAR current reached 0.3389 versus 0.3238 for recent only. Its improvement over recency was -0.0151, with a paired 95% interval of [-0.0437, 0.0092], which provides no supported advantage. FAR current was worse than Reservoir by 0.1595, with the entire improvement interval below zero. The original finding against adopting this retention rule remains unchanged.')
heading('Limits of the inference')
text('The controls were fixed before their own inference but added after the pilot results were examined; Reservoir was selected from that pilot. All conditions reuse the same eight scenes. The results are exploratory evidence within this set, not independent confirmation. Current only removes explicit history but generated current can carry past information, and one frame inference differs from the usual twelve frame context. Recent only bypasses the reader, so that comparison changes both retention and reading. FIFO shares Reservoir\'s reader and provides a separate check. Closed loop current images can diverge; these effects describe complete policies. Capacity and deletion are enforced at the policy interface, with no claim of physical deletion from evaluator memory.')

heading('Scene and seed detail',True)
text('The tables report final quarter DreamSim. Every scene row averages three seeds; village names are abbreviated and the test indices preserve exact scene identities. All 144 record summaries and all six full and tail metrics are retained in the accompanying compact results.')
body=[]
for idx in dict.fromkeys(r['clip_index'] for r in rows):
 rr=[r for r in rows if r['clip_index']==idx];scene=rr[0]['scene'].replace('_village_',' ').title();label=f"{idx} {scene} {rr[0]['route']}"
 body.append([label]+[f"{np.mean([float(r['dreamsim_tail']) for r in rr if r['policy']==p]):.4f}" for p in order])
table(['Scene','Current','Recent','FIFO','Reservoir','FAR','Archive'],body,[2.05,.81,.81,.81,.81,.81,.81],size=9)
body=[]
for seed in [0,1,2]:body.append(['Seed '+str(seed)]+[f"{S['policies'][p]['dreamsim_tail']['seed_means'][seed]:.4f}" for p in order])
for route in ['ABA','ABCA']:body.append([route+'  four scenes']+[f"{S['route_means'][route][p]:.4f}" for p in order])
table(['Aggregate','Current','Recent','FIFO','Reservoir','FAR','Archive'],body,[2.05,.81,.81,.81,.81,.81,.81],size=9.25)
doc.add_paragraph('Seed rows average eight scenes. Route rows average three seeds within each of four scenes. Seeds share scenes and are not independent population units.',style='Caption')
heading('Execution and verification')
text('The added controls consumed 36 minutes 26 seconds on one A100; the independent CPU audit took 2 minutes 9 seconds. No model was trained. A second implementation verified all 144 finite metric records and paired frame axes, every storage trace and control read, all 24 Algorithm R replays, and all 42 paired statistics and the Holm calculation. Current hashes matched 187 source files and 54 selected data, model and metric cache files. The 487 loaded floating tensors were finite before prediction.')
heading('Conclusion')
text('Within these eight LoopNav scenes, bounded memory that preserves earlier observations improves return prediction beyond both current only and recent only controls. Reservoir meets the declared memory benefit criterion and is the strongest equal capacity condition observed. FAR current provides no supported gain over recency and performs worse than Reservoir. This completes the explicit memory controls for the FAR based Phase1 test. It does not establish a superior learned retention rule, independent generalization, an I JEPA result or the original Phase1 GO decision.')
doc.core_properties.title='Historical Visual Memory in Phase1';doc.core_properties.subject='Independently verified six condition FAR memory comparison';doc.core_properties.author=''
doc.save(OUT/'_draft_report.docx');print(OUT/'_draft_report.docx')
