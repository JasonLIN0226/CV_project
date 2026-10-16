import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE=Path(__file__).resolve().parent
OUT=HERE/'report'
S=json.loads((HERE/'native_completed7426_7435/analysis_outputs/summary.json').read_text())
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing=1.05
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(19)
doc.styles['Title'].paragraph_format.space_after=Pt(4)
doc.styles['Subtitle'].font.size=Pt(10)
doc.styles['Subtitle'].paragraph_format.space_after=Pt(9)
doc.styles['Heading 1'].font.size=Pt(11.5)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(8)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(4)
doc.styles['Caption'].font.size=Pt(9)
doc.styles['Caption'].paragraph_format.space_after=Pt(5)

def text(s):return doc.add_paragraph(s)
def heading(s):return doc.add_heading(s,level=1)
def table(headers,body,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,h in enumerate(headers):t.rows[0].cells[i].text=h
    for row in body:
        cells=t.add_row().cells
        for i,v in enumerate(row):cells[i].text=str(v)
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            pad=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'70');e.set(qn('w:type'),'dxa');pad.append(e)
            pr.append(pad)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'333333' if ri==0 else ('F2F2F2' if ri%2==0 else 'FFFFFF'));pr.append(shade)
            for p in c.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
                for run in p.runs:
                    run.font.size=Pt(9.5);run.font.bold=ri==0;run.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t
def figure(name,caption,alt):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1)
    pic=p.add_run().add_picture(str(OUT/name),width=Inches(6.9));pic._inline.docPr.set('descr',alt)
    doc.add_paragraph(caption,style='Caption')

doc.add_paragraph('Memory Retention for Visual World Models',style='Title')
doc.add_paragraph('A controlled Phase 1 pilot on eight LoopNav scenes',style='Subtitle')
heading('Purpose')
text('We tested whether relevance to the current state is a useful rule for retaining a small visual memory. The experiment adapts a pretrained FAR world model to irreversible online storage. It compares retention rules while keeping the predictor and memory reader fixed, so differences reflect the retained evidence and the resulting closed loop predictions.')
heading('Controlled design')
text('We selected eight distinct LoopNav scenes before inspecting their scores, with four ABA and four ABCA routes. Three common seeds and four conditions produced 96 complete return records and 6996 prediction steps. FIFO, true Reservoir sampling and FAR current each retain 16 past entries plus a separate current entry. Archive reference keeps all arrived history and serves only as a diagnostic with greater storage.')
text('FAR current ranks candidates using the newly arrived state and a zero motion and zero time query. Retention cannot inspect future routes, images or queries. Deleted payloads cannot return. Return observations are each policy\'s generated images. All conditions use DDIM 20, stride 4 and context 12. Temporal chunks are one frame, instead of the original ten, to accommodate the same 16 entry capacity.')
heading('Evaluation')
text('The primary score is DreamSim error over the final quarter of each return, using the last ceil(n/4) actual predictions. Lower DreamSim and LPIPS are better; higher PSNR is better. We average seeds within each scene, then weight the eight scenes equally.')
body=[]
for label,key in [('FIFO','FIFO'),('Reservoir','Reservoir'),('FAR current','FAR_current'),('Archive reference','Archive_reference')]:
    p=S['policies'][key]
    body.append([label,f"{p['dreamsim_tail']['mean']:.4f} ({p['dreamsim_tail']['seed_SD']:.4f})",f"{p['lpips_tail']['mean']:.4f}",f"{p['psnr_tail']['mean']:.3f}"])
table(['Retention rule','DreamSim\nmean (seed SD)','LPIPS\nmean','PSNR\ndB'],body,[1.65,2.05,1.5,1.5])
figure('prediction_error.png','Figure 1  Primary prediction error. Bars show equal scene means; each dot is one scene averaged over three seeds. Archive reference has greater storage.','DreamSim final return quarter error for four retention policies, with eight individual scene means per policy.')

heading('Paired comparisons').paragraph_format.page_break_before=True
text('FAR current did not improve on FIFO. Its mean error was 0.3389 versus 0.3327 for FIFO, while Reservoir achieved 0.1795. Reservoir error was descriptively 46.1% lower than FIFO. The predefined comparison also found FAR current worse than Reservoir. Archive reference reached 0.1954, which shows that keeping every arrived entry was not the best configuration in this pilot.')
comparisons=[]
for label,key in [('FAR versus FIFO','FIFO'),('FAR versus Reservoir','Reservoir')]:
    p=S['paired_comparisons'][key]['dreamsim_tail'];lo,hi=p['paired_scene_bootstrap95']
    comparisons.append([label,f"{p['mean_improvement']:+.4f}",f"[{lo:+.4f}, {hi:+.4f}]",f"{p['exact_one_sided_signflip_p']:.4f}"])
table(['Comparison','Mean\nimprovement','Paired 95%\ninterval','One sided\np'],comparisons,[2.05,1.25,2.1,1.3])
figure('paired_improvement.png','Figure 2  Positive values favor FAR current. Intervals use 10000 paired scene bootstrap draws. Exact one sided sign flip tests enumerate all 256 patterns over eight scenes.','FAR minus baseline benefit expressed as baseline DreamSim error minus FAR error, with intervals spanning zero against FIFO and entirely negative against Reservoir.')
heading('What the memory traces show')
text('An additional descriptive audit examined which retained entries came from the observed outward route. At the end of return, those entries occupied 89.3% of Reservoir memory, 4.7% of FAR current memory and 0% of FIFO memory, averaged across the 24 scene and seed cases. This is consistent with current state relevance losing older evidence needed on return. It does not establish a causal explanation: generated current images can diverge across policies. This audit was added after results and did not change the primary test.')
heading('Scope and verification')
text('The GPU stage completed in 1 hour 46 minutes 50 seconds on one A100. A separate CPU audit and an independent replay verified all 96 paired frame axes, finite scores, source and asset hashes, storage transitions, and exact Algorithm R Reservoir choices. All reported summary statistics reproduced. This is an eight scene pilot using an existing FAR checkpoint. It is neither a full FAR reproduction nor an I JEPA evaluation, and it does not satisfy or replace the original Phase 1 GO criteria.')
heading('Conclusion')
text('The tested current state scoring rule should not replace the simpler baselines. Reservoir is the strongest equal capacity method in this pilot and should remain the reference for future retention studies. The results support testing preservation of earlier evidence rather than assuming that current relevance predicts later usefulness. A broader causal claim or a new retention mechanism requires a separate prospective experiment.')
doc.core_properties.title='Memory Retention for Visual World Models'
doc.core_properties.subject='Audited FAR to Phase 1 online memory pilot'
doc.core_properties.author=''
doc.save(OUT/'_draft_report.docx')
print(OUT/'_draft_report.docx')
