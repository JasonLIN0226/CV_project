from pathlib import Path
import json, csv, subprocess, statistics
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from reportlab.lib.colors import black, white

BASE = Path(__file__).resolve().parent
NATIVE = BASE / 'fast_execution/native_final_7506'
E = NATIVE / 'evaluation/job_7486'
OUT = BASE / 'report'
OUT.mkdir(exist_ok=True)
outcome = json.loads((E / 'outcome.json').read_text())
audit = json.loads((E / 'independent_scalar_audit.json').read_text())
proof = json.loads((NATIVE / 'final_native_verification_7506.json').read_text())
transfer = json.loads((NATIVE / 'transfer_integrity.json').read_text())
assert transfer['native_before_local_native_after_all_match']
assert outcome['status'] == 'DEVELOPMENT_REJECTED'
assert audit['status'] == 'INDEPENDENT_RAW_AXIS_FINITE_AND_64_FIXED_COMPARISONS_PASS'
assert proof['all_six_GPU_CPU_proofs_equal']
means = list(csv.DictReader((E / 'final_condition_means.csv').open()))
splits = ['test_id','test_long_gap','test_distractor_shift']
labels = {'test_id':'ID', 'test_long_gap':'Long gap', 'test_distractor_shift':'Distractor'}

# Recheck every reported fixed statistic against the separately implemented audit.
for task, result in outcome['tasks'].items():
    independent = {(x['slice'],x['kind']):x for x in audit['tasks'][task]['paired_checks']}
    assert len(independent) == 32 and audit['tasks'][task]['raw_records'] == 34560
    for x in result['fixed_reference_checks']:
        y = independent[(x['slice'],x['kind'])]
        v = x.get('improvement',x.get('absolute_mse_reduction'))
        assert x['passed'] == y['passed']
        assert all(abs(v[k]-y['independent'][k]) < 1e-12 for k in ['mean','low','high','pairs'])

fig = OUT / 'pooled_memory_gain.pdf'
c = canvas.Canvas(str(fig),pagesize=(490,285))
c.setTitle('Paired normalized MSE reduction on development data')
left,right,top,bottom = 140,460,251,40
xmin,xmax=-10,55
def X(v): return left+(v-xmin)/(xmax-xmin)*(right-left)
c.setFont('Helvetica',9)
for tick in [-10,0,10,20,30,40,50]:
    c.setLineWidth(.3);c.setStrokeColor(black)
    if tick in [0,10]:
        c.setDash(2,2) if tick==10 else c.setDash()
        c.line(X(tick),bottom,X(tick),top)
        c.setDash()
    c.drawCentredString(X(tick),bottom-13,str(tick))
c.drawCentredString((left+right)/2,12,'Reduction versus fixed recent reference in percent')
ys=[]
for ti,task in enumerate(['gate','motion']):
    for si,split in enumerate(splits):
        for hi,h in enumerate([3,6]):
            y=top-11-(ti*6+si*2+hi)*16.5
            x=next(x for x in outcome['tasks'][task]['fixed_reference_checks'] if x['kind']=='candidate_vs_reference_recent' and x['slice']==f'{task}/{split}/h{h}/seedall_seeds_mean')
            v=x['improvement'];a,m,b=[v[k]*100 for k in ['low','mean','high']]
            c.setFont('Helvetica',8.5);c.drawRightString(left-10,y-3,f'{task.title()}  {labels[split]}  h{h}')
            c.setLineWidth(.8);c.line(X(a),y,X(b),y);c.line(X(a),y-3,X(a),y+3);c.line(X(b),y-3,X(b),y+3)
            c.setFillColor(black if x['passed'] else white);c.circle(X(m),y,2.7,stroke=1,fill=1);c.setFillColor(black)
c.setFont('Helvetica',8)
c.drawString(left,270,'Paired 95 percent bootstrap intervals   Filled marker meets fixed gain rule')
c.save()
pdftoppm='/Users/linxiangyu/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/pdftoppm'
if not Path(pdftoppm).exists():
    pdftoppm='/Users/linxiangyu/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/fallback/pdftoppm'
subprocess.run([pdftoppm,'-png','-r','180','-singlefile',str(fig),str(OUT/'pooled_memory_gain')],check=True,capture_output=True)

doc=Document();sec=doc.sections[0]
sec.top_margin=Inches(.63);sec.bottom_margin=Inches(.62);sec.left_margin=sec.right_margin=Inches(.68)
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.05
doc.styles['Title'].font.size=Pt(23)
doc.styles['Heading 1'].font.size=Pt(15)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Caption'].font.size=Pt(9)
doc.styles['Caption'].font.bold=False
for border in list(doc.styles.element.xpath('.//w:pBdr')):
    border.getparent().remove(border)
for name in ['Heading 1','Heading 2']:
    doc.styles[name].paragraph_format.space_before=Pt(10)
    doc.styles[name].paragraph_format.space_after=Pt(5)

def para(text,style=None): return doc.add_paragraph(text,style)
def heading(text): doc.add_heading(text,level=1)
def table(headers,rows,widths=None,size=9):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    if widths:
        for column,w in zip(t.columns,widths):column.width=Inches(w)
        for cell,w in zip(t.rows[0].cells,widths):cell.width=Inches(w)
    for cell,text in zip(t.rows[0].cells,headers):
        cell.text=text
        for run in cell.paragraphs[0].runs:run.bold=True
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEDED');cell._tc.get_or_add_tcPr().append(sh)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for i,text in enumerate(row):
            cells[i].text=str(text)
            if widths:cells[i].width=Inches(widths[i])
    for row in t.rows:
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                for run in p.runs:run.font.size=Pt(size)
    return t
def page():doc.add_page_break()

doc.add_heading('Proposal Phase 1 Validation',0)
para('Gate and occluded Motion development evaluation   15 October 2026','Subtitle')
para('This experiment asks whether a frozen JEPA representation can support reliable future prediction when useful evidence lies outside the recent context. It tests one fixed history conditioned predictor on paired controlled tasks. The final decision and the boundary on independent confirmation appear at the end.')
heading('Question and controlled design')
para('Each sibling pair has identical recent images and actions but different earlier evidence and future outcomes. Gate requires a transient cue; Motion requires motion information before an occlusion. This construction isolates information supplied by history. It does not test a learned memory selector or a deletion policy.')
table(['Item','Prespecified setting'],[
('Tasks and seeds','Gate and occluded Motion; training seeds 17, 29 and 43'),
('Data per task','Seed 6301; 512 training pairs, 96 validation pairs; 128 pairs each for ID, long gap and distractor tests'),
('Prediction','Four recent frames; horizons 3 and 6; original frozen JEPA, train only normalization and 128 dimensional basis; error over all 5120 output dimensions'),
('Candidate','Frozen v3 recent predictor plus a zero initialized causal history correction; all retained history enters the predictor'),
('Training','Six models; each 120 epochs and 7680 steps; balanced four condition training; earliest minimum of equal condition and horizon validation MSE'),
('Evaluation','Six conditions including wrong evidence and persistence; all six models locked and checked before development tests')
],[1.35,5.75],9.5)
para('The candidate received additional training on top of a previously trained frozen reference. These comparisons do not match total historical training compute.')
heading('Conditions and acceptance rules')
para('Recent uses no explicit older evidence. Full uses all legal history. Privileged B2 uses two legally available task relevant history frames. Random B2 uses two sampled history frames. Wrong evidence replaces evidence with its sibling counterpart as a causal intervention. Persistence copies the current representation. Privileged access is an upper bound diagnostic, not a deployed selector.')
para('Both tasks must satisfy the original task rules, 24 fixed recent reference comparisons and 8 critical absolute error checks per task. Each seed and the paired mean across seeds must pass at both horizons. A relative gain requires a point estimate of at least 10 percent and a positive lower confidence bound. Critical absolute reductions must have a positive mean and lower bound. The full original rules also test decoded task performance and response to wrong evidence.')
heading('Execution and audit')
para('Training job 7485 completed successfully. Evaluation job 7486 produced 69120 finite raw scalar records across both tasks. Independent job 7487 checked pairing, axes, hashes and all 64 fixed comparisons using a separate implementation. Current source and selected checkpoint hashes matched their frozen locks. Exact input caching accelerated execution without changing the scientific samples or training values.')

page();heading('Observed development results')
table(['Task','Original rules','Fixed checks passed','Decision'],[
('Gate','38 of 51 checks; REVISE','24 of 32','Rejected'),
('Motion','40 of 51 checks; REVISE','14 of 32','Rejected')
],[1,2.3,2.2,1.6],10)
para('These counts are a checklist, not a probability of success. Acceptance is a conjunction of the prespecified rules, so a good pooled average cannot cancel a failed seed or hard condition.')
doc.add_picture(str(OUT/'pooled_memory_gain.png'),width=Inches(6.85))
para('Figure 1. Privileged B2 reduction in normalized MSE relative to the fixed recent reference. Seeds are averaged within each pair before 2000 paired bootstrap draws with seed 0. Each interval uses 128 independent pairs. The dashed line marks the 10 percent threshold; zero marks no improvement. These are development data, not independent confirmation.',style='Caption')
heading('Six condition means at horizon 6')
cond=['recent_only','full_history','privileged_small_budget','random_small_budget','wrong_evidence','persistence']
data=[]
for task in ['gate','motion']:
    for split in splits:
        vals=[]
        for co in cond:
            v=[float(r['normalized_mse']) for r in means if r['task']==task and r['split']==split and r['horizon']=='6' and r['condition']==co]
            assert len(v)==3;vals.append(f'{statistics.mean(v):.4f}')
        data.append([task.title(),labels[split],*vals])
table(['Task','Split','Recent','Full','B2','Random','Wrong','Persist'],data,[.65,.85,.85,.85,.85,.85,.85,.85],8.5)
para('Normalized MSE is lower when prediction is better. Means above average the three training seeds. The next two pages retain every seed and confidence interval used by the fixed adoption checks. Original task checks are also required and are not replaced by these latent error summaries.')

for task in ['gate','motion']:
    page();heading(task.title()+' fixed reference comparisons')
    result=outcome['tasks'][task];checks=result['fixed_reference_checks']
    critical='Distractor' if task=='gate' else 'Long gap'
    para('R compares candidate privileged B2 with the fixed reference recent predictor. A compares candidate privileged B2 with the earlier reference privileged B2 on the critical '+critical.lower()+' split. Values show point estimates and paired 95 percent confidence intervals. R is percent error reduction; A is an absolute normalized MSE reduction. Avg averages the three seeds within each pair. Positive values favor the candidate.')
    rows=[]
    for split in splits:
        for h in [3,6]:
            for seed in ['17','29','43','all_seeds_mean']:
                name=f'{task}/{split}/h{h}/seed{seed}'
                x=next(x for x in checks if x['kind']=='candidate_vs_reference_recent' and x['slice']==name)
                v=x['improvement'];r=f"{100*v['mean']:.1f} [{100*v['low']:.1f}, {100*v['high']:.1f}]"
                a=next((x for x in checks if x['kind']=='critical_absolute_mse_reduction' and x['slice']==name),None)
                av='';flags='R pass' if x['passed'] else 'R FAIL'
                if a:
                    z=a['absolute_mse_reduction'];av=f"{z['mean']:.4f} [{z['low']:.4f}, {z['high']:.4f}]";flags+='; A '+('pass' if a['passed'] else 'FAIL')
                rows.append([labels[split],h,'Avg' if seed=='all_seeds_mean' else seed,r,av,flags])
    table(['Split','h','Seed','R percent and 95 CI','A MSE and 95 CI','Rule'],rows,[.85,.32,.65,1.65,2.17,1.45],9)
    if task=='gate':
        para('Gate seed 17 fails all six R comparisons and fails wrong content intervention checks. Seed 43 additionally fails the distractor h6 R and A rules. Strong gains in seeds 29 and 43 on other splits do not establish stability across the required seeds.')
    else:
        para('Motion has 18 failed fixed checks. Long gap mean R is only 3.15 percent at h3 and 0.62 percent at h6, with intervals crossing zero. Its pooled A reductions are negative at both horizons, so the candidate is worse than the earlier privileged reference on this critical split. Distractor pooled R remains below 10 percent.')

page();heading('What the experiment establishes')
table(['Proposal premise','Evidence and present boundary'],[
('Useful information is missing from recent context','The controlled sibling construction keeps recent inputs and actions identical while earlier evidence and future outcomes differ. This supports the intended information gap in these constructed tasks.'),
('The frozen representation contains the evidence','All three evidence probe checks pass for each task. This supports readability of the historical cue and two frame motion information in this setting; it is not proof that every possible predictor can exploit it.'),
('A predictor reliably uses that evidence','Not established. Gate is sensitive to the training seed; Motion fails long gap and some distractor requirements. Wrong evidence and decoded task criteria remain part of the failed original decision.')
],[2.2,4.9],10)
heading('Interpretation and next decision')
para('The result identifies a predictor reliability problem in the tested configuration. History helps on several ordinary and pooled comparisons, but the improvement does not survive all required seeds and hard settings. It does not show that history is useless or that the research direction is impossible. It also provides no evidence that learning a memory selector will fix the current problem.')
para('The appropriate next research step is a bounded development diagnosis of the failed predictor: preserve the paired tasks, inspect Gate seed 17 and Motion long gap predictions, and separate failure to use evidence from interference caused by full history. Any revised candidate needs a new fixed protocol and development test. Independent confirmation should begin only after both tasks satisfy every unchanged development rule. No additional candidate or confirmation training was launched by this run.')
heading('Budget and evidence boundary')
para('Actual A100 allocation time was 6021 seconds for job 7485 plus 226 seconds from the two earlier budget stopped attempts: 6247 seconds in total, or 1 hour 44 minutes 7 seconds. The shared 5 hour ceiling leaves 11753 seconds unused. CPU evaluation, auditing and final evidence collection used no GPU. The failures above are scientific criterion failures, not exhaustion of the GPU allowance.')
para('Evidence consists of the frozen protocol and 48 source files; six current checkpoint completion proofs; 720 training scalar rows; both development comparison files; 69120 raw evaluation rows retained on HPC; the independent 64 comparison audit; and the native transfer ledger. Compact artifacts were checked with matching native before, local and native after SHA256 values. Full raw records, feature arrays and model weights remain on HPC.')
heading('Final conclusion')
para('Proposal Phase 1 has not passed. Both complete development tasks reject this candidate under the original rules. The condition for independent confirmation was therefore not met, and fresh seed 9137 data were neither opened nor generated. This is a completed negative development result, not an incomplete run and not an independent confirmation failure. A Phase 1 GO cannot be claimed from these results.')

doc.core_properties.title='Proposal Phase 1 Validation'
doc.core_properties.subject='Gate and occluded Motion development evidence and decision'
doc.core_properties.author=''
target=OUT/'Proposal_Phase1_Validation_Report.docx';doc.save(target)
(OUT/'report_validation.json').write_text(json.dumps({'source_outcome':str(E/'outcome.json'),'all64_independent_fixed_comparisons_match':True,'all69120_raw_records_verified_by_independent_audit':True,'native_three_SHA_transfer':True,'phase1_passed':False,'fresh_confirmation_opened':False,'GPU_used_seconds':6247,'GPU_unused_seconds':11753,'document':str(target)},indent=2)+'\n')
print(target)
