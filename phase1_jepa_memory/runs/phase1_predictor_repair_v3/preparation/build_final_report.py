"""Build the final budget-boundary report only from native verified artifacts."""
from pathlib import Path
import json,hashlib,csv,statistics
from reportlab.graphics.shapes import Drawing,String,Line
from reportlab.graphics.charts.barcharts import VerticalBarChart,HorizontalBarChart
from reportlab.graphics import renderSVG
from reportlab.lib import colors
import subprocess
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

BASE=Path(__file__).resolve().parents[1]
NATIVE=BASE/'native_final7586'
OUT=BASE/'report'
def read(p):return json.loads(Path(p).read_text())
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def hms(n):
    n=int(round(n));return f'{n//3600}h {n//60%60:02d}m {n%60:02d}s'

def table(doc,headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,(h,w) in enumerate(zip(headers,widths)):t.rows[0].cells[i].text=h;t.columns[i].width=Inches(w)
    for row in rows:
        cells=t.add_row().cells
        for cell,value in zip(cells,row):cell.text=str(value)
    for ri,row in enumerate(t.rows):
        for ci,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            margins=OxmlElement('w:tcMar');
            for edge in ('top','left','bottom','right'):
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8E8E8' if ri==0 else ('F7F7F7' if ri%2==0 else 'FFFFFF'));pr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1)
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:run.font.size=Pt(9);run.bold=ri==0;run.font.color.rgb=RGBColor(0,0,0)
    trpr=t.rows[0]._tr.get_or_add_trPr();rep=OxmlElement('w:tblHeader');trpr.append(rep)
    return t

def main():
    proof=read(NATIVE/'native_budget_stop.json');transfer=read(NATIVE/'transfer_integrity.json')
    assert transfer['native_before_local_native_after_all_match'] and sha(NATIVE/'native_budget_stop.json')==transfer['files']['native_budget_stop.json']['sha256']
    assert proof['no_model_forward_or_training'] and proof['complete_models']==0 and len(proof['partials'])==4
    assert proof['GPU_cap_seconds']==28800 and proof['outcome']['status']=='COMPUTE_BUDGET_STOP'
    assert all(v['model_and_optimizer_finite'] and v['checkpoint_stable'] for v in proof['partials'].values())
    forecast=proof['forecast'];diag=read(BASE/'native_diagnosis7580/diagnostic_results_v2.json')
    OUT.mkdir(exist_ok=True)
    with (OUT/'all_120_diagnostic_slices.csv').open('w',newline='') as f:
        keys=sorted(set().union(*(r.keys() for r in diag['slices'])));w=csv.DictWriter(f,fieldnames=keys);w.writeheader();w.writerows(diag['slices'])
    def mean(split,h,key):
        rows=[x for x in diag['slices'] if x['task']=='motion' and x['split']==split and x['condition']=='full_history' and x['horizon']==h]
        assert len(rows)==3;return statistics.mean(x[key] for x in rows)
    node='/Users/linxiangyu/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node'
    sharp='/Users/linxiangyu/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/node_modules/sharp'
    def save_chart(d,name):
        svg=OUT/(name+'.svg');png=OUT/(name+'.png');renderSVG.drawToFile(d,str(svg))
        code="require(process.argv[1])(process.argv[2],{density:220}).png().toFile(process.argv[3]).catch(e=>{console.error(e);process.exit(1)})"
        subprocess.run([node,'-e',code,sharp,str(svg),str(png)],check=True)
    d=Drawing(504,190)
    for pos,title,key,basekey,legends,ymax in [
        (30,'Direct feature extrapolation','feature_secant_MSE',None,('ID','Long gap'),6),
        (292,'Long gap common error','common_MSE','reference_common_MSE',('Recent reference','Second repair'),.12)]:
        chart=VerticalBarChart();chart.x=pos;chart.y=48;chart.width=185;chart.height=105
        if basekey:
            values=[[mean('test_long_gap',h,basekey) for h in (3,6)],[mean('test_long_gap',h,key) for h in (3,6)]]
        else:values=[[mean('test_id',h,key) for h in (3,6)],[mean('test_long_gap',h,key) for h in (3,6)]]
        chart.data=values;chart.categoryAxis.categoryNames=['h3','h6'];chart.categoryAxis.labels.fontName='Helvetica';chart.categoryAxis.labels.fontSize=9
        chart.valueAxis.valueMin=0;chart.valueAxis.valueMax=ymax;chart.valueAxis.valueStep=ymax/3;chart.valueAxis.labels.fontSize=8
        chart.bars[0].fillColor=colors.white;chart.bars[0].strokeColor=colors.black
        chart.bars[1].fillColor=colors.HexColor('#999999');chart.bars[1].strokeColor=colors.black
        chart.barSpacing=4;chart.groupSpacing=17;d.add(chart)
        d.add(String(pos-10,173,title,fontName='Helvetica-Bold',fontSize=10))
        d.add(String(pos-10,159,'MSE',fontSize=8))
        d.add(String(pos-10,8,'White  '+legends[0]+'    Gray  '+legends[1],fontSize=8))
    save_chart(d,'motion_diagnosis')
    d=Drawing(504,112);chart=HorizontalBarChart();chart.x=116;chart.y=30;chart.width=340;chart.height=52
    chart.data=[[proof['GPU_used_seconds']/3600,forecast['projected_cumulative_gpu_seconds']/3600]]
    chart.categoryAxis.categoryNames=['Actual used','Full plan forecast'];chart.categoryAxis.labels.fontSize=9
    chart.valueAxis.valueMin=0;chart.valueAxis.valueMax=9;chart.valueAxis.valueStep=1;chart.valueAxis.labels.fontSize=8
    chart.bars[0].fillColor=colors.HexColor('#BBBBBB');chart.bars[0].strokeColor=colors.black;d.add(chart)
    capx=116+340*8/9;line=Line(capx,25,capx,89);line.strokeDashArray=[3,2];d.add(line)
    d.add(String(300,98,'Approved cumulative cap 8h',fontSize=9));d.add(String(225,7,'Cumulative A100 hours',fontSize=9))
    save_chart(d,'budget_boundary')
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.72);sec.left_margin=sec.right_margin=Inches(.85)
    normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10.5);normal.font.color.rgb=RGBColor(0,0,0)
    normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.06
    for name,size in [('Title',22),('Heading 1',14),('Heading 2',11)]:
        s=doc.styles[name];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
    doc.add_paragraph('Proposal Phase1 third repair and budget evidence','Title')
    doc.add_paragraph('16 October 2026   |   Gate and occluded Motion',style='Subtitle')
    doc.add_paragraph('This report records the focused diagnosis of Motion long occlusion, the remaining Gate check, and the measured resource boundary of one new predictor. It preserves the original acceptance rules and separates completed scientific evidence from an incomplete new trial.')
    doc.add_heading('Completed evidence before the third repair',1)
    table(doc,['Task','Original checks','Fixed reference checks','Decision'],[
        ['Gate','50 of 51','32 of 32','Rejected'],['Motion','35 of 51','8 of 32','Rejected']], [1.1,1.65,2.1,1.1])
    after_table=doc.add_paragraph('These are the second repair results from all six completed models, each with 120 epochs and 7680 steps. Their 69120 raw records and 64 fixed comparisons were independently checked. All checks are required; a near pass is not a pass.')
    after_table.paragraph_format.space_before=Pt(6)
    doc.add_heading('What the focused diagnosis found',1)
    doc.add_paragraph('Motion: no geometric reflections occurred in the diagnosed subset. Directly extending visual features produced much larger errors at long gaps. The second repair also increased the common part of long gap prediction error, despite improving the difference between paired futures. Removing its transport component reduced h6 MSE from 0.1871 to 0.1727. These observations motivate a different dynamics representation; they do not identify one unique cause.')
    doc.add_picture(str(OUT/'motion_diagnosis.png'),width=Inches(6.65))
    doc.add_paragraph('Figure 1  Read only diagnostics on the first 64 pairs per split, averaged across three seeds. These subset summaries do not replace any acceptance test.',style='Caption')
    doc.add_paragraph('Gate: the remaining hard failure was seed43 on ID at h6. Privileged memory improved normalized latent MSE, but relative state classification error reduction was 10.94% with a confidence interval of -4.69% to 26.56%. The interval crossing zero prevented acceptance. The proposed change therefore targets how the remembered cue combines with the future action.')
    doc.add_page_break()
    doc.add_heading('One fixed literature informed correction',1)
    doc.add_paragraph('LatentFlowPredictor learns a small nonlinear coordinate map from observed visual features, extrapolates differences over elapsed time in that coordinate space, and decodes nonlinearly into the frozen output basis. Post pooling FiLM conditions the correction on legal future actions. This borrows coordinate learning principles from E2C and nonlinear dynamical embeddings, and featurewise conditioning from FiLM. It is not a reproduction of those full methods or evidence that a physical state was learned.')
    doc.add_paragraph('The frozen recent predictor remains exact when history is empty. All retained observations are read. No hidden simulator state, labels, sibling identity, or future image enters the predictor. Tasks, seeds17/29/43, samples, 120 epochs, losses, optimizer, selection rule, and all acceptance thresholds remain unchanged. The revealed development data make this an exploratory repair; independent confirmation is still required.')
    doc.add_heading('Checks completed and work retained',1)
    table(doc,['Evidence','Verified scope'],[
        ['Software boundaries','10 predictor fixtures and 4 staged data fixtures'],
        ['Native input preflight','9728 exact input target info RNG samples and 24 validation batches'],
        ['Frozen current sources',f"58 science files and 3 independent audit files"],
        ['Retained scientific work','Gate17 and Motion17 each 1 of 120 epochs and 64 of 7680 steps'],
        ['Partial checkpoint audit','Four best latest files finite with AdamW RNG and frozen reference verified'],
        ['Complete new models and tests','0 of 6 models complete no new held out result no fresh9137 opened']], [2.15,4.25])
    doc.add_heading('Measured complete sequence budget',1)
    doc.add_picture(str(OUT/'budget_boundary.png'),width=Inches(6.65))
    doc.add_paragraph('Figure 2  Actual consumption and a conservative cost forecast are different quantities. The forecast includes all remaining development, 12 new confirmation models, encoding, safety margin, and overhead.',style='Caption')
    table(doc,['GPU accounting','Seconds','Time'],[
        ['Prior settled use','12604',hms(12604)],
        ['This partial GPU allocation',proof['GPU_ledger']['7582'],hms(proof['GPU_ledger']['7582'])],
        ['Total actual used',proof['GPU_used_seconds'],hms(proof['GPU_used_seconds'])],
        ['Unused within approved cap',proof['GPU_remaining_seconds'],hms(proof['GPU_remaining_seconds'])],
        ['Projected complete cumulative plan',f"{forecast['projected_cumulative_gpu_seconds']:.2f}",hms(forecast['projected_cumulative_gpu_seconds'])]], [3.45,1.3,1.65])
    doc.add_page_break()
    doc.add_heading('Why the budget guard stopped',1)
    doc.add_paragraph(f"The first full Gate and Motion epochs took {forecast['measured_full_epochs']['gate']:.3f} and {forecast['measured_full_epochs']['motion']:.3f} seconds. The complete plan forecast was {hms(forecast['projected_cumulative_gpu_seconds'])}, exceeding the approved eight hour cumulative cap by {hms(forecast['projected_cumulative_gpu_seconds']-28800)}. This is an estimate, not a promise that a small budget increase would suffice.")
    doc.add_paragraph('The forecast reserves six complete development candidates, six new confirmation references and six new confirmation candidates, all 120 epochs; 198656 encoded confirmation frames; a 15% safety factor; and fixed proof and loading overhead. Leaving GPU time unused avoids completing development at the expense of the required independent confirmation. No sample, seed, epoch, or threshold was shortened to fit.')
    doc.add_heading('Evidence integrity and limits',1)
    doc.add_paragraph(f"The native partial verifier checked the terminal GPU ledger, all current source files, and four real partial checkpoints without a model forward or optimizer update. All {len(transfer['files'])} explicitly selected compact files, including the verifier output, have matching native before transfer, local, and native after transfer SHA256 values. Large checkpoints and raw feature arrays remain on HPC.")
    doc.add_paragraph('The new trial has no complete evaluated model and no new scientific score. Its epoch validation values are not acceptance results. The diagnostic subset was selected by fixed manifest order and is descriptive. Learning dynamics coordinates and action conditioning may help, but their effect has not been established. Fresh confirmation would require new train only transforms and all 12 newly trained models locked before opening the tests; none of that confirmation was run.')
    doc.add_paragraph('Evidence directory: runs/phase1_predictor_repair_v3/native_final7586. Diagnostic records: native_diagnosis7580 and diagnostic_review.json. Frozen implementation: execution/control/source_lock.json. Full 120 slice summaries accompany this report as a CSV. The earlier complete second repair remains unchanged in phase1_predictor_repair_v2.')
    doc.add_heading('Primary method references',1)
    for s in [
        'Watter et al 2015. Embed to Control. https://arxiv.org/abs/1506.07365',
        'Lusch Kutz and Brunton 2018. Deep learning for universal linear embeddings of nonlinear dynamics. https://arxiv.org/abs/1712.09707',
        'Perez et al 2018. FiLM Visual Reasoning with a General Conditioning Layer. https://arxiv.org/abs/1709.07871']:
        p=doc.add_paragraph(s);p.paragraph_format.space_after=Pt(5)
        for run in p.runs:run.font.size=Pt(9)
    doc.add_heading('Conclusion',1)
    doc.add_paragraph('Proposal Phase1 has not passed. The last complete development experiment failed the unchanged criteria. The third repair was stopped by the authorized resource guard before a complete model or independent confirmation existed. Its method therefore remains untested by the full acceptance procedure. This is insufficient evidence to certify Phase1 under the approved eight hour budget, rather than a demonstrated scientific failure of the new predictor.')
    # Strip template theme colors and paragraph rules for the requested monochrome report.
    for style in doc.styles:
        if hasattr(style,'font'):
            style.font.color.rgb=RGBColor(0,0,0)
        for element in list(style.element.iter()):
            if element.tag==qn('w:pBdr'):
                element.getparent().remove(element)
            if element.tag==qn('w:color'):
                element.set(qn('w:val'),'000000')
                for attr in ('themeColor','themeTint','themeShade'):
                    element.attrib.pop(qn('w:'+attr),None)
    for element in list(doc.element.iter()):
        if element.tag==qn('w:pBdr'):element.getparent().remove(element)
    path=OUT/'Proposal_Phase1_Third_Repair_Budget_Report.docx';doc.save(path)
    manifest={'docx':str(path),'docx_sha256':sha(path),'native_proof_sha256':sha(NATIVE/'native_budget_stop.json'),'native_transfer_files':len(transfer['files']),'conclusion':'PHASE1_NOT_PASSED_NEW_REPAIR_INCOMPLETE_BUDGET_STOP','requires_render_and_all_pages_visual_review':True}
    (OUT/'report_manifest.json').write_text(json.dumps(manifest,indent=2)+'\n');print(json.dumps(manifest))

if __name__=='__main__':main()
