"""Create a report from transferred native evidence, without model execution."""
from pathlib import Path
import csv
import hashlib
import json
import math
import numpy as np
import subprocess
from reportlab.pdfgen import canvas
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path(__file__).resolve().parents[1]
NATIVE=BASE/'approved_8h/native_final7553'
REPORT=BASE/'report'
REPORT.mkdir(exist_ok=True)
def read(p): return json.loads(Path(p).read_text())
def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
v=read(NATIVE/'final_native_verification.json')
transfer=read(NATIVE/'transfer_integrity.json')
assert transfer['native_before_local_native_after_all_match'] and len(transfer['files'])==42
for rel,item in transfer['files'].items():
    assert sha(NATIVE/rel)==item['sha256']
outcome=read(NATIVE/'evaluation/job_7536/outcome.json')
independent=read(NATIVE/'evaluation/job_7536/independent_scalar_audit.json')
assert outcome['status']=='DEVELOPMENT_REJECTED' and not outcome['confirmation_opened']
assert independent['status']=='INDEPENDENT_RAW_AXIS_FINITE_AND_64_FIXED_COMPARISONS_PASS'
assert all(v['jobs'][str(j)]['state']=='COMPLETED' and v['jobs'][str(j)]['exit']=='0:0' for j in (7535,7536,7537))
rows=0
for task in ('gate','motion'):
    for seed in (17,29,43):
        folder=NATIVE/f'training/job_7535/predictor/{task}/seed_{seed}'
        t=read(folder/'training.json')
        with (folder/'training.csv').open(newline='') as f: scalar=list(csv.DictReader(f))
        assert len(t['logs'])==len(scalar)==120
        for i,(a,b) in enumerate(zip(t['logs'],scalar),1):
            assert a['epoch']==int(b['epoch'])==i and a['step']==int(b['step'])==64*i
            for k,value in a.items():
                if isinstance(value,(int,float)):
                    assert math.isfinite(value) and math.isclose(value,float(b[k]),abs_tol=1e-12,rel_tol=0)
        proof=v['six_current_checkpoint_proofs'][f'{task}/{seed}']
        assert proof['best_epoch']==min(t['logs'],key=lambda r:r['validation_mse'])['epoch']
        assert t['signature']==proof['signature']
        rows+=len(scalar)
assert rows==720
fixed=[]; original=[]
for task,x in outcome['tasks'].items():
    for c in x['original_checks']: original.append({'task':task,**c})
    for c in x['fixed_reference_checks']:
        z=c.get('improvement',c.get('absolute_mse_reduction'))
        fixed.append({'task':task,'slice':c['slice'],'kind':c['kind'],'passed':c['passed'],
                      'mean':z['mean'],'low':z['low'],'high':z['high'],'pairs':z['pairs'],
                      'scale':z.get('scale','absolute_mse_reduction')})
for name,data in [('all_64_fixed_checks.csv',fixed),('all_102_original_checks.csv',original)]:
    fields=sorted(set().union(*(r.keys() for r in data)))
    with (REPORT/name).open('w',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=fields);writer.writeheader();writer.writerows(data)
summary={'status':'COMPLETE_DEVELOPMENT_CRITERIA_NOT_MET','phase1_go':False,'confirmation_opened':False,
    'GPU_used_seconds':v['GPU_used_seconds'],'GPU_remaining_seconds':v['GPU_remaining_seconds'],
    'GPU_cap_seconds':28800,'models_completed':6,'epochs_each':120,'local_training_rows_verified':rows,
    'raw_records':69120,'independent_fixed_comparisons':64,'compact_three_SHA_files':42,
    'tasks':v['tasks'],'native_verification_sha256':sha(NATIVE/'final_native_verification.json')}
(REPORT/'final_summary.json').write_text(json.dumps(summary,indent=2)+'\n')

pdf=REPORT/'critical_error_intervals.pdf'
c=canvas.Canvas(str(pdf),pagesize=(700,340))
c.setFont('Helvetica-Bold',11)
c.drawCentredString(350,322,'Critical absolute error reductions with paired 95 percent intervals')
for panel,(task,title) in enumerate((('gate','Gate under distractor shift'),('motion','Motion under long occlusion'))):
    data=[r for r in fixed if r['task']==task and r['kind']=='critical_absolute_mse_reduction']
    left=70+panel*350;right=330+panel*350;bottom=47;top=285
    lo=min(r['low'] for r in data);hi=max(r['high'] for r in data)
    xmin=math.floor(min(lo,0)/.01)*.01;xmax=math.ceil(max(hi,0)/.01)*.01
    xp=lambda z:left+(z-xmin)/(xmax-xmin)*(right-left)
    c.setFont('Helvetica-Bold',10);c.drawCentredString((left+right)/2,301,title)
    c.setFont('Helvetica',8)
    ticks=np.linspace(xmin,xmax,5)
    for tick in ticks:
        x=xp(tick);c.setStrokeGray(.88);c.setLineWidth(.5);c.line(x,bottom,x,top)
        c.setFillGray(0);c.drawCentredString(x,bottom-14,f'{tick:.3f}')
    c.setStrokeGray(.4);c.setDash(3,2);c.line(xp(0),bottom,xp(0),top);c.setDash()
    for i,r in enumerate(data):
        y=top-12-i*28
        label=r['slice'].split('/')[2]+' '+r['slice'].split('/')[3].replace('seedall_seeds_mean','mean').replace('seed','s')
        c.setFillGray(0);c.drawRightString(left-7,y-3,label)
        c.setStrokeGray(0);c.setLineWidth(.9);c.line(xp(r['low']),y,xp(r['high']),y)
        for z in (r['low'],r['high']):c.line(xp(z),y-3,xp(z),y+3)
        c.setFillGray(0 if r['passed'] else 1);c.circle(xp(r['mean']),y,2.7,stroke=1,fill=1)
    c.setFillGray(0);c.setFont('Helvetica',8);c.drawCentredString((left+right)/2,17,'Fixed reference MSE minus candidate MSE')
c.showPage();c.save()
subprocess.run(['/Users/linxiangyu/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/pdftoppm',
    '-png','-singlefile','-r','180',str(pdf),str(REPORT/'critical_error_intervals')],check=True,stdout=subprocess.DEVNULL)

doc=Document()
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
sec=doc.sections[0];sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.62);sec.left_margin=Inches(.72);sec.right_margin=Inches(.72)
for style in ('Normal','Title','Heading 1','Heading 2'):
    doc.styles[style].font.name='Arial';doc.styles[style].font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(23)
doc.styles['Heading 1'].font.size=Pt(16)
doc.styles['Heading 2'].font.size=Pt(12)
def p(text):return doc.add_paragraph(text)
def h(text):return doc.add_heading(text,1)
def table(headers,rows,widths=None,small=9):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
    if widths:
        for column,width in zip(t.columns,widths):column.width=Inches(width)
    for c,text in zip(t.rows[0].cells,headers):c.text=str(text)
    for row in rows:
        for c,text in zip(t.add_row().cells,row):c.text=str(text)
    for row in t.rows:
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        for j,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[j])
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(1);para.paragraph_format.space_before=Pt(1)
                for run in para.runs:run.font.size=Pt(small)
    for cell in t.rows[0].cells:
        for run in cell.paragraphs[0].runs:run.bold=True
        shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'EEEEEE');cell._tc.get_or_add_tcPr().append(shading)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    return t

doc.add_paragraph('Proposal Phase1 Predictor Repair Validation',style='Title')
p('Completed development evaluation and independent verification\n15 October 2026  America New York')
h('Purpose and fixed experimental scope')
p('Phase1 asks whether a JEPA predictor can use earlier observed evidence to improve future prediction when recent frames are insufficient. The current test evaluates the preregistered repair on delayed Gate cues and occluded Motion. It does not train a memory selector or test the full research proposal.')
p('The repair combines attention pooling inspired by Set Transformer [1] with observed frame relations inspired by relational reasoning [2], bounded relative age, and linear elapsed transport. It keeps a frozen recent reference and one history correction. This is an adapted composite, not a reproduction of either paper. It reads all retained history without TopK retrieval.')
table(['Boundary','Preserved setting'],[
    ('Data and seeds','Development data 6301; seeds 17 29 43; two tasks'),
    ('Sample sizes per task','512 training pairs; 96 validation pairs; 128 pairs in each of three test splits'),
    ('Training','6 models; each 120 epochs and 7680 steps; four conditions balanced'),
    ('Context and targets','4 recent frames; memory budget 2; horizons 3 and 6; full 5120 dimensional output'),
    ('Selection','Earliest equal condition validation MSE minimum; no test based checkpoint selection'),
    ('Acceptance','All original task rules and all 32 fixed reference checks per task required')],[1.35,5.65])
h('Execution and verification')
p('Training 7535, CPU evaluation 7536, and independent scalar audit 7537 each completed with exit 0:0. All six current best and latest checkpoints passed finite model and optimizer, RNG, scheduler, frozen reference, source signature, 120 epoch and earliest best checks. Native and local training logs agree across 720 rows. Gate17 and Motion17 transparently resume the original epoch 1 state; the other four runs start from their fixed seeds.')
p('The audit checked all 69120 raw records for finite values and complete paired axes, then independently reproduced all 64 fixed reference intervals and pass flags. A separate final CPU verification repeated these checks. All 42 compact files have matching native before transfer, local, and native after transfer SHA256 values. These are execution and arithmetic checks, not evidence that scientific acceptance criteria were met.')
table(['Task','Original rules met','Fixed checks met','Raw records'],[
    ('Gate','39 of 51','19 of 32','34560'),('Motion','49 of 51','26 of 32','34560')],[1.0,2.0,2.0,2.0])
p('Check counts summarize different requirements. They are not prediction accuracies, success probabilities, or substitutes for the all required rule.')

doc.add_page_break();h('Where the repaired predictor remains unreliable')
doc.add_picture(str(REPORT/'critical_error_intervals.png'),width=Inches(7.0))
p('Figure 1. Positive values favor the repaired privileged memory predictor over the fixed reference privileged memory predictor. A critical absolute check requires a strictly positive mean and lower interval bound. Filled markers pass; open markers fail. Each interval uses 128 episode pairs, grouped at pair level with 2000 fixed bootstrap resamples. The mean label averages scalar errors across the three seeds within each paired episode; it does not add independent pairs.')
p('Gate failures concentrate on distractor shift. All eight privileged benefit slices in that split fail, as do four content replacement intervention slices. On the fixed absolute comparison, seed43 is worse at both horizons; seed29 intervals cross zero. Thus the repair has not established stable use of the intended cue under changed distractors.')
p('Motion is closer to the acceptance boundary, but it still misses the seed17 distractor benefit checks at both horizons. Long occlusion also fails four required absolute comparisons: h3 for seeds17 and29 and their three seed mean, and h6 for seed29. The h3 mean reduction is 0.002995 with interval minus 0.001239 to 0.007277, so a reliable absolute improvement is not established.')
table(['Task','Seed17 fixed','Seed29 fixed','Seed43 fixed','Three seed mean fixed'],[
    (t.title(),*[f"{sum(r['passed'] for r in fixed if r['task']==t and r['slice'].endswith('/seed'+s))}/8" for s in ('17','29','43','all_seeds_mean')]) for t in ('gate','motion')],[1.0,1.5,1.5,1.5,1.5])
p('The paired intervals describe this revealed development data. The repair was chosen after previous development failures were inspected, so these results cannot serve as independent confirmation. Differences across seeds and splits identify unresolved behavior; they do not isolate a unique causal mechanism.')
h('How to read the full comparison tables')
p('R compares candidate privileged memory with the fixed reference recent only predictor. Relative error reduction must be at least 10 percent and its interval lower bound must be above zero. A compares candidate privileged memory with fixed reference privileged memory in the preregistered critical split; absolute MSE reduction and its lower bound must both exceed zero. All reported R rows use the relative scale. ID means the original test distribution, Long means long occlusion, and Shift means changed distractors. The next two pages include every fixed comparison.')

for task in ('gate','motion'):
    doc.add_page_break();h(task.title()+' fixed reference comparisons')
    p('Values are mean reduction and paired 95 percent interval. R values are percent; A values are normalized MSE units. Each row uses 128 pairs. All rows are required.')
    rr=[]
    for r in fixed:
        if r['task']!=task:continue
        _,split,horizon,seed=r['slice'].split('/')
        kind='R' if r['kind']=='candidate_vs_reference_recent' else 'A'
        assert kind!='R' or r['scale']=='relative_error_reduction'
        factor=100 if kind=='R' else 1
        fmt=lambda x:(f'{x*factor:.2f}' if kind=='R' else f'{x:.5f}')
        rr.append([{'test_id':'ID','test_long_gap':'Long','test_distractor_shift':'Shift'}[split],horizon[1:],
            seed.replace('seedall_seeds_mean','Mean').replace('seed',''),kind,
            fmt(r['mean']),f"[{fmt(r['low'])}, {fmt(r['high'])}]",'PASS' if r['passed'] else 'FAIL'])
    table(['Split','h','Seed','Type','Mean','95 percent interval','Rule'],rr,[.63,.3,.52,.46,.8,2.57,.72],small=8.8)

doc.add_page_break();h('Budget and retained evidence')
table(['GPU job','Actual seconds','Purpose'],[(j,v['GPU_ledger'][j],purpose) for j,purpose in [
    ('7477','Earlier partial execution'),('7480','Earlier partial execution'),('7485','Previous complete development trial'),
    ('7524','Repair epoch 1 states and original budget stop'),('7535','Exact resume and complete repaired development')]], [1.1,1.2,4.7])
p('Cumulative GPU allocation time is 9797 seconds, or 2 hours 43 minutes 17 seconds, against the authorized 28800 seconds. The unused balance is 19003 seconds, or 5 hours 16 minutes 43 seconds. CPU evaluation, independent audit, and final evidence collection add no GPU time. No additional GPU training or fresh data generation was launched after the negative development result.')
p('The fixed 58 scientific files and 3 independent audit files match their locked SHA256 values at final native verification. Raw records, weights, feature arrays, and checkpoints remain on HPC. Local compact evidence includes both task outcomes, metrics, training histories, checkpoint certificates, the independent audit, the GPU ledger, and the 42 file transfer manifest. The full 102 original rule flags and 64 comparison values are also supplied as editable CSV files.')
h('What the evidence supports')
p('The completed run tests the third practical premise: whether this predictor reliably uses available earlier evidence across required conditions. It improves many slices but still fails the full criterion. Earlier controlled task and representation diagnostics motivated the trial; the new run does not independently reconfirm those premises. This result rejects the current candidate under the stated development rules, not every possible predictor or the broader research idea.')
h('Recommended next decision')
p('Do not describe Phase1 as passed or proceed to selector training. Preserve fresh9137 as unseen. Any further scientific change should begin with a bounded diagnosis of distractor dependence and Motion long gap calibration, using development data only, followed by a newly frozen candidate and the same full acceptance rules. Such another training trial is not started by this report. A fresh 12 model confirmation is eligible only after both development tasks satisfy all rules and pass independent audit.')
h('Method references')
p('[1] Lee et al. Set Transformer A Framework for Attention based Permutation Invariant Neural Networks. ICML 2019. https://proceedings.mlr.press/v97/lee19d.html')
p('[2] Santoro et al. A simple neural network module for relational reasoning. 2017. https://arxiv.org/abs/1706.01427')
h('Conclusion')
p('Proposal Phase1 has not passed. The repaired development trial completed correctly, but both tasks fail required original criteria. Independent confirmation was not opened because its prerequisite was not met. This is a scientific acceptance failure for the tested candidate, rather than an incomplete training run, a connection failure, or exhaustion of the authorized GPU budget.')
doc.save(REPORT/'Proposal_Phase1_Repair_Final_Report.docx')
print(json.dumps({'report':str(REPORT/'Proposal_Phase1_Repair_Final_Report.docx'),'rows_verified':rows,'fixed_checks':len(fixed),'GPU_seconds':v['GPU_used_seconds']}))
