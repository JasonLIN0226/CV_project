"""Create an editable plain-language report from already verified compact results."""
from pathlib import Path
import json, hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
DATA = ROOT / 'native7251_7268'
health = json.loads((DATA / 'visual_health.json').read_text())
deletion = json.loads((DATA / 'deletion_summary.json').read_text())
reference = json.loads((DATA / 'fixed_old_reference_comparison.json').read_text())
audit = json.loads((DATA / 'root_final_audit_7268.json').read_text())
assert sum(v['supported'][0] for h in ('16', '64') for v in health['full_vs_recent_paired_checks'][h].values()) == 6
assert sum(v['supported'][0] for h in ('1', '4') for v in health['full_vs_recent_paired_checks'][h].values()) == 0
assert sum(sum(v['supported']) for group in deletion['checks'].values() for v in group.values()) == 47
assert reference['primary_checks_passed'] == 9
assert audit['replay_scalar_rows'] == 812592 and audit['irreversible_retention_trace_rows'] == 456192
assert audit['all_fixed_FIFO_Reservoir_Uniform_transitions_match_no_reentry']

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(1.7)
section.left_margin = section.right_margin = Cm(1.9)
section.header_distance = section.footer_distance = Cm(.7)

def font(style, size, bold=False):
    style.font.name = 'Noto Sans CJK SC'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rpr.insert(0, fonts)
    fonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    fonts.set(qn('w:ascii'), 'Noto Sans CJK SC'); fonts.set(qn('w:hAnsi'), 'Noto Sans CJK SC')
    for key in ('eastAsiaTheme', 'asciiTheme', 'hAnsiTheme'):
        fonts.attrib.pop(qn('w:'+key), None)

font(doc.styles['Normal'], 10.5)
doc.styles['Normal'].paragraph_format.line_spacing = 1.12
doc.styles['Normal'].paragraph_format.space_after = Pt(5)
font(doc.styles['Title'], 20, True)
doc.styles['Title'].paragraph_format.space_after = Pt(5)
font(doc.styles['Subtitle'], 10)
doc.styles['Subtitle'].paragraph_format.space_after = Pt(10)
font(doc.styles['Heading 1'], 12.5, True)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(4)
doc.styles['Heading 1'].paragraph_format.keep_with_next = True
font(doc.styles['Heading 2'], 11, True)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)

def paragraph(text, bold=False):
    p = doc.add_paragraph()
    p.add_run(text).bold = bold
    return p

def heading(text):
    return doc.add_paragraph(text, 'Heading 1')

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Cm(width)
    for c, width in zip(t.rows[0].cells, widths): c.width = Cm(width)
    for i, h in enumerate(headers): t.rows[0].cells[i].text = h
    for values in rows:
        cells = t.add_row().cells
        for c, v, width in zip(cells, values, widths): c.text = v; c.width = Cm(width)
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge = OxmlElement('w:' + name)
        for k, v in {'val':'single', 'sz':'4', 'color':'D9D9D9'}.items(): edge.set(qn('w:'+k), v)
        borders.append(edge)
    props.append(borders)
    for ri, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if ri == 0: trpr.append(OxmlElement('w:tblHeader'))
        for ci, c in enumerate(row.cells):
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = c._tc.get_or_add_tcPr()
            mar = OxmlElement('w:tcMar')
            for side, val in [('top',75),('bottom',75),('left',100),('right',100)]:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),str(val));el.set(qn('w:type'),'dxa');mar.append(el)
            tcpr.append(mar)
            if ri == 0:
                shade = OxmlElement('w:shd');shade.set(qn('w:fill'),'E4EDF5');tcpr.append(shade)
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.08
                if ri == 0 or ci == 1: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(9.5)
                    if ri == 0: run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

doc.add_paragraph('视觉记忆实验阶段报告', 'Title')
doc.add_paragraph('2026年10月13日  纽约时间  真实组装视频实验简明版', 'Subtitle')
paragraph('阶段性结论  当前验证数据支持历史信息对8秒和32秒后的预测有帮助，理想删除选择在多数比较中有优势。训练间稳定性仍需补充，动作语义评价尚未完成。整体通过或不通过的判定暂缓。', True)

heading('1 我们要验证什么')
paragraph('我们想验证：模型的记忆空间有限时，保留哪些过去的信息能帮助它更好地预测未来，忘掉什么会损害预测。例如，当前画面中的手挡住了零件，早些时候看到的动作可能帮助判断接下来会发生什么。')
paragraph('本阶段检验历史是否有用、不同删法是否有区别，以及持续遗忘时会发生什么。理想删除选择用于检查研究潜力；自动学习删除策略还需要后续验证。')

heading('2 做了哪些实验')
paragraph('固定数据来自78段组装玩具录像，其中58段用于训练、20段用于验证；抽取198个片段，共50,688个观察时点。I-JEPA把画面转成描述内容的数字特征，小预测模型学习预测0.5秒、2秒、8秒和32秒后的特征。')
paragraph('同一个预测模型接受四种记忆条件：只看最近；记忆满时忘掉最老内容；从过去随机抽样保留；保留全部可用历史。使用同一批视频和目标作比较，独立训练三次，每次30轮。')
paragraph('实验A比较全部历史与只看最近的预测误差。实验B在624个固定记忆状态下尝试全部17种删法，并比较删除当下和之后32步的表现。实验C连续播放全部198个片段，让模型边看边删除，检查FIFO、Reservoir和Uniform三种固定策略；主记忆容量为16，另检查8和32，被删内容不能取回。')
paragraph('已完成两轮真实视频训练。最新一轮仅加强了正则化，尝试减少过拟合；模型、数据和训练轮数保持一致。我们另与上一轮固定模型的近期及完整历史预测作比较。')

heading('3 实验结果是什么')
table(['检查项目', '当前结果', '如何理解'], [
    ['8秒及32秒历史收益', '6/8项可靠改善', '两次独立训练和合并结果均有改善；另一次训练的两项收益仍不确定。'],
    ['0.5秒及2秒历史收益', '0/8项可靠正收益', '积极证据集中在较长预测时间；短期暂无稳定收益，部分训练误差增加。'],
    ['理想删除选择', '47/48项可靠改善', '多数比较支持删法有区别；一项32秒预测比较的收益区间略跨过零。'],
    ['与上一轮固定模型比较', '9/16项可靠改善', '加强正则化带来部分改善，但尚未在所有训练与长预测时间中稳定优于旧模型。'],
    ['持续不可逆遗忘', '全198个片段完成', '预测结果与保留轨迹已完整评估；固定策略、容量和删除不可恢复等执行规则核验正确。'],
    ['动作语义评价', '尚未完成', '标签与视频的精确对应仍待证明，动作判断、切换和恢复能力尚未评价。'],
], [4.2, 3.3, 9.7])
paragraph('6/8来自三次训练及合并结果，在8秒、32秒两个预测时间上共8项比较。47/48来自不同记忆来源、删除基线、训练与评价指标的组合；这些检查共享模型和视频，不能当作48次独立试验或真实场景成功率。')

heading('4 如何评价实验结果')
paragraph('先看预测误差。将预测特征与真实未来特征比较，计算归一化均方误差（MSE），数值越小越好。用“基线误差减去候选误差”表示收益，正数代表候选更准。删除实验同时考虑当下和后续预测，用同一个目标选出理想删法。')
paragraph('再看结果是否可靠。先在片段和录像内部汇总，再给每位验证参与者相同权重。对10位验证参与者重采样10,000次，计算95%置信区间；区间全在零以上，支持改善；跨零，记录为证据不确定；全在零以下，支持候选误差增加。三次训练分别报告，并补充合并结果。')
paragraph('例如，一次训练的8秒预测误差从0.3988降到0.3981，平均好了一点，但收益区间为−0.00065到＋0.00187，包含零。因此这项保留为“改善仍不确定”，同时保留另外两次训练和合并结果中的积极证据。')
paragraph('持续遗忘实验还检查每一步保留内容、容量、时间顺序和随机状态，确保被删内容没有回流。执行正确性的核验与预测效果的评价分别记录。理想删法会参考真实未来，用来估计潜力，尚不代表可部署的自动删除策略。')
paragraph('当前证据适合表述为“长程视觉预测与删除选择获得部分支持”。下一步需定位较弱训练的原因，并补齐60帧每秒视频与30帧每秒动作标注的精确映射，再评价动作语义。原始数据、统计区间和既定评价记录均保留。', True)

p=paragraph('资料依据  已核验的GPU训练记录、全量A B C评估摘要、原始标量与不可逆保留轨迹复核。此前可控模拟属于前置实验，本报告聚焦真实视频结果。')
for run in p.runs: run.font.size=Pt(9)

doc.core_properties.title = '视觉记忆实验阶段报告'
doc.core_properties.subject = '目的 实验 结果 评价方法'
doc.core_properties.author = ''
doc.core_properties.comments = ''
output = ROOT / 'CV视觉记忆实验简明报告.docx'
doc.save(output)
manifest = {
    'output': str(output), 'sha256': hashlib.sha256(output.read_bytes()).hexdigest(),
    'overall_decision': 'DEFERRED_AS_REQUESTED', 'evidence': 'PARTIAL_SUPPORT_STABILITY_AND_SEMANTICS_REQUIRE_FURTHER_VALIDATION',
    'long_horizon_supported':6, 'long_horizon_checks':8, 'short_horizon_supported':0, 'short_horizon_checks':8,
    'deletion_supported':47, 'deletion_checks':48, 'old_reference_supported':9, 'old_reference_checks':16,
    'original_machine_results_and_thresholds_modified':False,
    'source_files':[{'path':str(q),'sha256':hashlib.sha256(q.read_bytes()).hexdigest()} for q in (
        DATA/'visual_health.json', DATA/'deletion_summary.json', DATA/'fixed_old_reference_comparison.json', DATA/'root_final_audit_7268.json')]
}
(ROOT/'simple_report_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(manifest, ensure_ascii=False)[:1200])
