"""Author a monochrome English report from verified compact experiment results."""
from pathlib import Path
import csv
import datetime
import hashlib
import json
import math
import os
import shutil
import statistics
import subprocess

from reportlab.graphics.shapes import Drawing, String, Circle
from reportlab.graphics.charts.lineplots import LinePlot
from reportlab.graphics import renderSVG
from reportlab.lib import colors
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
PROJECT = ROOT.parent.parent
NEW = ROOT / 'native7251_7268'
OLD = PROJECT / 'runs/phase1_v5_real_evaluation/native_completed7209_diagnostic7226'
OUT = ROOT / 'english_detailed_report'
OUT.mkdir(exist_ok=True)
SOURCES = {}


def load(path):
    raw = path.read_bytes()
    SOURCES[str(path.relative_to(PROJECT))] = hashlib.sha256(raw).hexdigest()
    return json.loads(raw)


health = load(NEW / 'visual_health.json')
health_replay = load(NEW / 'root_health_scalar_replay_7268_20261014T0144.json')
old_health = load(OLD / 'root_health_scalar_replay_20261013T2330.json')
reference = load(NEW / 'fixed_old_reference_comparison.json')
deletion = load(NEW / 'deletion_summary.json')
replay = load(NEW / 'replay_participant_summary.json')
training = load(NEW / 'training_report.json')
protocol = load(NEW / 'protocol.json')
train_protocol = load(ROOT / 'software/protocol.json')
audit = load(NEW / 'root_final_audit_7268.json')
transfer = load(NEW / 'transfer_integrity_final.json')
original_report = ROOT / 'experiment_report.md'
original_report_hash = hashlib.sha256(original_report.read_bytes()).hexdigest()
assert original_report_hash == '06fd69c8e9db668790c6993700f0fedc032f5c6e017470864e9cfe6db11e04d2'
for item in transfer['files']:
    path = Path(item['local'])
    assert item['matched'] and path.stat().st_size == item['local_proof']['bytes']
    assert hashlib.sha256(path.read_bytes()).hexdigest() == item['local_proof']['sha256']

SEEDS = ('17', '29', '43')
GROUPS = (*SEEDS, 'pooled')
HORIZONS = (1, 4, 16, 64)
POLICIES = ('recent', 'FIFO', 'Reservoir', 'full')
TABLE_COUNTS = {}
TABLE_DATA = {}


def csv_rows(path):
    raw = path.read_bytes()
    SOURCES[str(path.relative_to(PROJECT))] = hashlib.sha256(raw).hexdigest()
    rows = list(csv.DictReader(raw.decode().splitlines()))
    assert len(rows) == 30 and [int(r['epoch']) for r in rows] == list(range(1, 31))
    return rows


new_curves = {s: csv_rows(NEW / f'seed_{s}_training.csv') for s in SEEDS}
old_curves = {s: csv_rows(OLD / f'seed_{s}_training.csv') for s in SEEDS}
for s, expected in zip(SEEDS, (5, 6, 7)):
    assert min(new_curves[s], key=lambda r: float(r['validation_full_mse']))['epoch'] == str(expected)
    assert sum(int(r['steps']) for r in new_curves[s]) == 9690
assert audit['all6_current_proof_recomputed_equals_executed']
assert audit['replay_scalar_rows'] == 812592
assert len(replay['participant_means']) == 216


def mean_mse(seed, h, policy, old=False):
    data = (old_health if old else health_replay)['participant_equal_weight_mse']
    if seed == 'pooled':
        return statistics.mean(mean_mse(s, h, policy, old) for s in SEEDS)
    return float(data[seed][str(h)][policy]['mse'])


def measure(seed, h, field):
    data = health_replay['participant_equal_weight_mse']
    if seed == 'pooled':
        return statistics.mean(measure(s, h, field) for s in SEEDS)
    return float(data[seed][str(h)]['recent'][field])


def fmt(value, places=8, signed=False):
    return format(float(value), f'{"+" if signed else ""}.{places}f')


def pct(delta, baseline):
    return f'{100.0 * delta / baseline:+.3f}'


def interval(check, i=0, places=8):
    return f'[{fmt(check["ci95_lower"][i], places)}, {fmt(check["ci95_upper"][i], places)}]'


def evidence(check, i=0):
    low, high = check['ci95_lower'][i], check['ci95_upper'][i]
    return 'I' if low > 0 else ('W' if high < 0 else 'U')


def replay_mse(seed, split, policy, capacity, h):
    if seed == 'pooled':
        return statistics.mean(replay_mse(s, split, policy, capacity, h) for s in SEEDS)
    values = replay['participant_means'][f'{seed}|{split}|{policy}|{capacity}|{h}']
    assert len(values) == (29 if split == 'train' else 10)
    return statistics.mean(values.values())


# Task-local fonts make rendering deterministic without changing user fonts.
fonts = OUT / 'fonts'
fonts.mkdir(exist_ok=True)
for filename in ('Times New Roman.ttf', 'Times New Roman Bold.ttf',
                 'Times New Roman Italic.ttf', 'Times New Roman Bold Italic.ttf'):
    target = fonts / filename
    if not target.exists():
        shutil.copyfile(Path('/System/Library/Fonts/Supplemental') / filename, target)
(fonts / 'fonts.conf').write_text(
    '<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd">'
    f'<fontconfig><dir>{fonts}</dir><cachedir>{fonts}/cache</cachedir>'
    '<alias><family>serif</family><prefer><family>Times New Roman</family>'
    '</prefer></alias></fontconfig>')

drawing = Drawing(510, 470)
for ri, s in enumerate(SEEDS):
    for ci, metric in enumerate(('train_full_mse', 'validation_full_mse')):
        chart=LinePlot()
        chart.x=40+255*ci; chart.y=342-150*ri
        chart.width=205; chart.height=92
        chart.data=[[(int(r['epoch']),float(r[metric])) for r in data]
                    for data in (old_curves[s],new_curves[s])]
        values=[v for data in chart.data for _,v in data]
        low=math.floor(min(values)*20)/20
        high=math.ceil(max(values)*20)/20
        if low==high: high=low+.05
        chart.xValueAxis.valueMin=1; chart.xValueAxis.valueMax=30
        chart.xValueAxis.valueSteps=[1,5,10,15,20,25,30]
        chart.yValueAxis.valueMin=low; chart.yValueAxis.valueMax=high
        chart.yValueAxis.valueSteps=[low,(low+high)/2,high]
        chart.yValueAxis.labelTextFormat='%.3f'
        for axis in (chart.xValueAxis,chart.yValueAxis):
            axis.labels.fontName='Times-Roman'; axis.labels.fontSize=8
            axis.strokeColor=colors.black
        chart.lines[0].strokeColor=chart.lines[1].strokeColor=colors.black
        chart.lines[0].strokeWidth=chart.lines[1].strokeWidth=1
        chart.lines[0].strokeDashArray=[4,3]
        drawing.add(chart)
        drawing.add(String(chart.x,chart.y+105,f'Seed {s}  {"training" if ci==0 else "validation"}',
                           fontName='Times-Roman',fontSize=10,fillColor=colors.black))
        drawing.add(String(chart.x,chart.y+94,'Normalized MSE',fontName='Times-Roman',
                           fontSize=8,fillColor=colors.black))
        drawing.add(String(chart.x+92,chart.y-22,'Epoch',fontName='Times-Roman',
                           fontSize=8,fillColor=colors.black))
        if ci == 1:
            best = min(new_curves[s], key=lambda r: float(r[metric]))
            cx=chart.x+(int(best['epoch'])-1)/29*chart.width
            cy=chart.y+(float(best[metric])-low)/(high-low)*chart.height
            drawing.add(Circle(cx,cy,2.3,strokeColor=colors.black,fillColor=colors.white,strokeWidth=.8))
figure = OUT / 'training_curves_monochrome.png'
svg=OUT/'training_curves_monochrome.svg'
renderSVG.drawToFile(drawing,str(svg))
svg.write_text(svg.read_text().replace('Times-Roman','Times New Roman'))
node='/Users/linxiangyu/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node'
sharp='/Users/linxiangyu/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/node_modules/sharp/dist/index.mjs'
script=f"import sharp from {json.dumps(sharp)}; await sharp(process.argv[1], {{density:220}}).png().toFile(process.argv[2]);"
subprocess.run([node,'--input-type=module','-e',script,str(svg),str(figure)],
               check=True,cwd=OUT,env={**os.environ,'FONTCONFIG_FILE':str(fonts/'fonts.conf')})

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.7), Cm(1.6)
sec.left_margin = sec.right_margin = Cm(1.9)
sec.header_distance = sec.footer_distance = Cm(.7)
for name, size, bold in [('Normal', 11, False), ('Title', 18, True),
                         ('Subtitle', 10.5, False), ('Heading 1', 13, True),
                         ('Heading 2', 11.5, True), ('Caption', 9.5, False)]:
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts'); rpr.insert(0, rfonts)
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'): rfonts.set(qn('w:' + key), 'Times New Roman')
    for key in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'): rfonts.attrib.pop(qn('w:' + key), None)
    for border in list(style.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
doc.styles['Normal'].paragraph_format.line_spacing = 1.12
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
for name in ('Heading 1', 'Heading 2'):
    pf = doc.styles[name].paragraph_format
    pf.space_before, pf.space_after, pf.keep_with_next = Pt(9), Pt(5), True
doc.styles['Title'].paragraph_format.space_after = Pt(5)
doc.styles['Subtitle'].paragraph_format.space_after = Pt(10)
foot = sec.footer.paragraphs[0]
foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = foot.add_run(); run.font.size = Pt(9)
begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
code = OxmlElement('w:instrText'); code.text = ' PAGE '
end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
run._r.extend([begin, code, end])


def paragraph(text, bold=False, size=None):
    after_table = len(doc._element.body) > 1 and doc._element.body[-2].tag == qn('w:tbl')
    p = doc.add_paragraph()
    if after_table: p.paragraph_format.space_before = Pt(5)
    r = p.add_run(text); r.bold = bold
    if size: r.font.size = Pt(size)
    return p


next_heading_new_page = False


def heading(text, level=1):
    global next_heading_new_page
    p = doc.add_paragraph(text, f'Heading {level}')
    if next_heading_new_page:
        p.paragraph_format.page_break_before = True
        next_heading_new_page = False
    return p


def page():
    global next_heading_new_page
    next_heading_new_page = True


table_number = 0


def table(title, headers, rows, widths, key, font_size=9):
    global table_number
    table_number += 1
    caption = doc.add_paragraph(f'Table {table_number}. {title}', 'Caption')
    caption.paragraph_format.keep_with_next = True
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(4)
    TABLE_COUNTS[key] = len(rows)
    TABLE_DATA[key] = {'columns': headers, 'rows': rows}
    headers = ['Time s' if h == 'Seconds' else 'Result' if h == 'Evidence' else 'Gain %' if h == 'Gain percent' else h for h in headers]
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, w in zip(t.columns, widths): c.width = Cm(w)
    for c, w, h in zip(t.rows[0].cells, widths, headers): c.width = Cm(w); c.text = h
    for values in rows:
        cells = t.add_row().cells
        for c, value, w in zip(cells, values, widths): c.width = Cm(w); c.text = str(value)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        for k, v in {'val': 'single', 'sz': '4', 'color': 'D9D9D9'}.items(): el.set(qn('w:' + k), v)
        borders.append(el)
    t._tbl.tblPr.append(borders)
    for ri, row in enumerate(t.rows):
        props = row._tr.get_or_add_trPr(); props.append(OxmlElement('w:cantSplit'))
        if ri == 0: props.append(OxmlElement('w:tblHeader'))
        for c in row.cells:
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = OxmlElement('w:tcMar')
            for side, value in [('top', 52), ('bottom', 52), ('left', 62), ('right', 62)]:
                el = OxmlElement('w:' + side); el.set(qn('w:w'), str(value)); el.set(qn('w:type'), 'dxa'); margins.append(el)
            c._tc.get_or_add_tcPr().append(margins)
            for p in c.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.02
                for run in p.runs:
                    run.font.name = 'Times New Roman'; run.font.size = Pt(font_size)
                    run.font.color.rgb = RGBColor(0, 0, 0); run.bold = ri == 0
    return t


doc.add_paragraph('Visual Memory Experiments', 'Title')
doc.add_paragraph('Experimental design  Evaluation scores  Detailed results', 'Subtitle')
paragraph('14 October 2026  |  Assembly101 development data  |  Real video trials v5 and v6', size=10)

heading('1 Research objective')
paragraph('We study whether a prediction model can use past visual evidence to predict a future scene, and whether choosing what to forget can preserve useful information when memory is limited. For example, an earlier view of a toy component may remain useful after the hand or another component obscures it.')
paragraph('This report describes the latest real video experiments and the preceding matched trial used as a reference. It first defines the data, model, experiments and evaluation scores, then gives every reported comparison at the seed and horizon level. The interpretation is presented in the final section. Earlier synthetic Gate and Motion trials remain separate historical experiments.')

heading('2 Data and prediction model')
table('Fixed data used in the two real video trials',
      ['Split', 'Videos', 'Participants', 'Episodes', 'Observations', 'Queries'],
      [['Training', 58, 29, 136, '34,816', '5,168'],
       ['Validation', 20, 10, 62, '15,872', '2,356'],
       ['Total development', 78, 39, 198, '50,688', '7,524']],
      [3.7, 1.7, 2.5, 2.3, 3.6, 3.4], 'data')
paragraph('Each episode contains 256 observations sampled 0.5 seconds apart. Prediction queries are at observation indices 40, 44, ..., 188, giving 38 queries per episode. The four prediction horizons are 1, 4, 16 and 64 observations, corresponding to 0.5, 2, 8 and 32 seconds. Training and validation participants are disjoint. The sealed test data were not evaluated.')
paragraph('A frozen I-JEPA encoder converts each image into a 5,120-dimensional visual feature vector. The encoder is not fine-tuned. A small causal Transformer, called the reader, is trained to predict future features from legal past observations. It uses train-only normalization and 128-dimensional input and output bases. Prediction error is measured after reconstruction in all 5,120 normalized dimensions.')
table('Reader and training specification', ['Setting', 'Value'], [
    ['Reader architecture', 'Hidden size 256; 2 Transformer layers; 4 attention heads; feed-forward size 1,024'],
    ['Trainable parameters', '1,779,072'],
    ['Recent observations and historical memory', '8 recent observations; historical capacity 16 for FIFO and Reservoir'],
    ['Inputs', 'Visual features, actual observation ages, recent or history role, and prediction horizon'],
    ['Independent training runs', 'Seeds 17, 29 and 43; 30 epochs and 9,690 optimizer steps per seed'],
    ['Balanced batch', '16 identical base queries repeated under all 4 conditions; 64 condition examples per step'],
    ['Regularization', 'Model dropout 0.10; retained-history dropout 0.15'],
    ['Optimizer and learning rate', 'AdamW; learning rate 0.0003; 5 percent warmup followed by cosine decay'],
    ['Single change in v6', 'Weight decay increased from 0.05 to 0.20; all other scientific settings retained'],
    ['Checkpoint selection', 'Earliest strict minimum of validation MSE averaged equally over all 4 conditions and 4 horizons'],
], [4.1, 13.1], 'model', 9.5)

page()
heading('3 Experimental design')
paragraph('For each seed, one shared reader is trained with all four memory conditions. The comparisons change the available context while keeping that reader and the future target fixed.')
table('Memory conditions', ['Condition', 'Information available'], [
    ['Recent', 'The 8 most recent observations, including the current observation.'],
    ['FIFO', 'The same recent observations plus up to 16 older observations. When historical memory fills, the oldest historical observation is deleted.'],
    ['Reservoir', 'The same recent observations plus up to 16 older observations retained by Algorithm R. Older eligible observations can be sampled and a new eligible item can be rejected.'],
    ['Full', 'The recent observations and all earlier available observations in the episode. No future observation is supplied.'],
], [2.5, 14.7], 'conditions', 10)
heading('3 1 Experiment A  Future visual prediction', 2)
paragraph('We evaluate all 2,356 validation queries under all 4 conditions and all 4 horizons for each of the 3 frozen readers. This gives 113,088 prediction scores. The primary comparison is Full versus Recent at 8 and 32 seconds; the 0.5 and 2 second results are also retained. Persistence predicts that the current feature remains unchanged. The fixed-basis oracle projects the known future residual into the fixed output basis and gives a diagnostic reconstruction floor; it has access to the future target.')
heading('3 2 Fixed comparison against the previous trial', 2)
paragraph('The v6 Full prediction is compared with the already executed v5 Recent and v5 Full predictions for exactly the same query keys. The old models were not run again. All four horizons are reported. The predeclared main comparisons are at 8 and 32 seconds for each seed and the pooled result.')
heading('3 3 Experiment B  Deletion choice and future consequences', 2)
paragraph('We select 4 valid full-memory states per recording under each of 2 source trajectories, FIFO and Reservoir. All 78 development recordings are included, giving 624 states. At a state, the existing 16 older records and the next eligible record create 17 legal deletion choices. Every choice is evaluated immediately and through the same 32 subsequent Reservoir updates, with shared random draws. All 3 readers are evaluated, producing 31,824 candidate records. The reported confidence intervals use the validation participants only, covering 160 of the 624 states.')
paragraph('One joint oracle deletion is selected for each state and reader. Its objective is the equally weighted mean of the immediate four-horizon MSE and the mean four-horizon MSE over the 32 continuation steps. The same selected deletion is then used for the joint, 8 second and 32 second comparisons against FIFO and Reservoir deletion baselines. Exact ties delete the oldest timestamp. The oracle sees future errors, so this experiment measures available deletion headroom rather than the performance of a deployable learned selector.')
heading('3 4 Experiment C  Persistent irreversible memory', 2)
paragraph('All 198 episodes are replayed with FIFO, Reservoir and Uniform policies at historical capacities 8, 16 and 32. Uniform deletes one of the 17 incoming or retained candidates uniformly when capacity is 16, and uses the analogous rule at other capacities. Unlike Reservoir, Uniform gives the incoming item the same deletion chance as each currently retained item. Deleted visual payloads cannot return from an external archive. The 1,782 retention trajectories are shared by all 3 readers, giving 812,592 prediction scores at the same 38 queries and 4 horizons. Both training and validation scores are reported separately.')

page()
heading('4 Evaluation scores and statistical interpretation')
heading('4 1 Normalized mean squared error', 2)
paragraph('Let z be the 5,120-dimensional encoded feature and let mu and sigma be its training-only mean and standard deviation. Each dimension is normalized as x = (z - mu) / max(sigma, 0.000001). The reader predicts the future normalized feature x_hat. The score is the mean of (x_hat[d] - x_future[d]) squared over d = 1, ..., 5120. Lower MSE means a more accurate feature prediction. This is a visual feature error, not action classification accuracy or a percentage of correct videos.')
paragraph('The training loss averages this MSE equally across the four horizons and four conditions. The reader predicts output-basis coefficients; the full feature reconstruction is current normalized feature + mean future residual + predicted coefficients multiplied by the transpose of the fixed output basis.')
heading('4 2 Absolute and relative improvement', 2)
paragraph('Absolute improvement Delta = baseline MSE - candidate MSE. Positive Delta means lower candidate error; negative Delta means higher candidate error. Relative improvement in percent = 100 times Delta divided by baseline MSE. The reported percentages are descriptive error reductions, not accuracy gains. For Experiment B, Delta is the reduction in its joint or horizon-specific deletion cost; no percentage is supplied because the compact summary reports differences rather than absolute baseline costs.')
heading('4 3 Joint deletion score', 2)
paragraph('For deletion d and horizon h, J_h(d) = 0.5 times immediate MSE_h(d) + 0.5 times the mean of the 32 continuation MSE_h(d) values. The joint score J(d) is the mean of J_h(d) over h = 1, 4, 16 and 64. The oracle minimizes J(d) once. In the results, Delta is J(baseline) - J(oracle), or J_h(baseline) - J_h(the same oracle) for the horizon rows. A positive joint improvement is partly expected from an oracle selected on this objective; it does not show that a learned online policy can identify the choice.')
heading('4 4 Aggregation and confidence intervals', 2)
paragraph('Evaluation scores first average queries within an episode, then episodes within a recording, then recordings within a participant. Participants receive equal weight. This prevents a long recording or a person with many queries from dominating the result. The validation comparisons use 10 participants. Pooled results average the three readers within each participant; they do not create 30 independent participants.')
paragraph('A paired participant bootstrap resamples these 10 participants 10,000 times with random seed 0. The 2.5 and 97.5 percentiles form the two-sided 95 percent confidence interval for Delta. The original intervals are reproduced directly from the executed summaries. They are not adjusted for multiple comparisons. Rows share videos and readers and must not be interpreted as independent success trials.')
table('Meaning of the evidence column', ['Code', 'Meaning'], [
    ['I', 'Improvement supported: the entire 95 percent interval for Delta is above zero.'],
    ['U', 'Uncertain direction: the interval includes zero, even if the point estimate is positive.'],
    ['W', 'Worse supported: the entire interval is below zero.'],
], [1.3, 15.9], 'legend', 10)
paragraph('The original adoption rules required positive intervals for all primary seed and pooled comparisons in A and B, with the additional fixed old-reference checks in v6. Those rules and their historical machine outputs are retained. The overall research judgment is deferred, and the final section separates supportive evidence from uncertainty and missing semantic evaluation.')
paragraph('Additional policy contrasts without intervals are descriptive. The original intervals are retained, with no new statistical tests or predictions. Values are displayed to 8 decimal places, or 10 for the small deletion effects; calculations use the full stored precision.')

page()
heading('5 Training results')
rows = []
for s in SEEDS:
    old_best = min(old_curves[s], key=lambda r: float(r['validation_full_mse']))
    new_best = min(new_curves[s], key=lambda r: float(r['validation_full_mse']))
    a, b = float(old_best['validation_full_mse']), float(new_best['validation_full_mse'])
    rows.append([s, new_best['epoch'], fmt(a), fmt(b), fmt(a-b, signed=True), pct(a-b, a),
                 fmt(float(new_curves[s][-1]['train_full_mse'])), fmt(float(new_curves[s][-1]['validation_full_mse']))])
table('Training and validation MSE for each independently initialized reader',
      ['Seed', 'Best epoch', 'v5 best val', 'v6 best val', 'Delta best val', 'Gain percent', 'v6 final train', 'v6 final val'],
      rows, [1.0, 1.4, 2.5, 2.5, 2.4, 1.7, 2.85, 2.85], 'training_best', 8.7)
paragraph('Each run completed 30 epochs and 9,690 optimizer steps. Best validation MSE decreased slightly in all three runs. Selection still occurred at epochs 5, 6 and 7. The final training errors are lower than the final validation errors, and validation errors rose after the selected epochs. This is consistent with a generalization gap; it does not identify its sole cause. The selected best checkpoints, rather than the final checkpoints, were used for the scientific results below.')
paragraph('Training validation MSE is query-weighted and averaged across conditions and horizons. The scientific tables use participant-balanced aggregation and show separate horizons. These two kinds of means should not be directly subtracted.')
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(str(figure), width=Cm(16.8))
doc.add_paragraph('Figure 1. All 30 training and validation epochs. Dashed lines show v5 and solid lines show v6. Open circles mark the v6 selected validation minima. Exact epoch values appear in Section 10.', 'Caption')

page()
heading('6 Experiment A  Prediction scores')
paragraph('All values in this section use the validation participants and the selected best checkpoints. The Full, FIFO and Reservoir predictions share the same trained reader for a given seed.')
absolute = []
for h in HORIZONS:
    for s in GROUPS:
        absolute.append([s, str(h*.5).rstrip('0').rstrip('.') if h*.5 != 0.5 else '0.5',
                         *[fmt(mean_mse(s,h,p)) for p in POLICIES],
                         fmt(measure(s,h,'persistence_mse')), fmt(measure(s,h,'basis_oracle_mse'))])
table('Absolute normalized MSE under every validation condition',
      ['Seed', 'Seconds', 'Recent', 'FIFO', 'Reservoir', 'Full', 'Persistence', 'Basis oracle'],
      absolute, [1.2, 1.2, 2.45, 2.45, 2.45, 2.45, 2.5, 2.5], 'A_absolute', 8.6)
paragraph('Persistence and the basis oracle are identical across the three seeds because they use the same features and fixed fit. The oracle uses the true future feature and is only a diagnostic floor.')
gain = []
for h in HORIZONS:
    for s in GROUPS:
        c=health['full_vs_recent_paired_checks'][str(h)][s]
        base,new=mean_mse(s,h,'recent'),mean_mse(s,h,'full')
        assert math.isclose(base-new,c['point'][0],abs_tol=1e-14)
        gain.append([s, h*.5, fmt(base), fmt(new), fmt(c['point'][0],signed=True),
                     pct(c['point'][0],base), interval(c), evidence(c)])
table('Full history improvement over Recent for every seed and horizon',
      ['Seed', 'Seconds', 'Recent MSE', 'Full MSE', 'Delta', 'Gain percent', '95 percent CI of Delta', 'Evidence'],
      gain, [1.1, 1.0, 2.0, 2.0, 2.2, 1.4, 6.0, 1.5], 'A_full_improvements', 8.6)
paragraph('I means a positive interval, U means the interval includes zero, and W means a negative interval. In particular, the positive point estimates for seed 29 at 8 and 32 seconds retain their uncertainty.')

page()
heading('6 1 Other memory conditions relative to Recent', 2)
paragraph('This table exposes the size and direction of the FIFO and Reservoir effects at every horizon. These are arithmetic differences of the participant-balanced means. The executed Experiment A summary did not provide confidence intervals for these additional policy contrasts, so they are descriptive rather than additional statistical adoption tests.')
policy_rows=[]
for h in HORIZONS:
    for s in GROUPS:
        base=mean_mse(s,h,'recent')
        for p in ('FIFO','Reservoir'):
            candidate=mean_mse(s,h,p);delta=base-candidate
            policy_rows.append([s,h*.5,p,fmt(base),fmt(candidate),fmt(delta,signed=True),pct(delta,base)])
table('Every FIFO and Reservoir point improvement over Recent',
      ['Seed','Seconds','Candidate','Recent MSE','Candidate MSE','Delta','Gain percent'],
      policy_rows,[1.1,1.1,2.4,3.1,3.1,3.1,3.3],'A_policy_point_improvements',9)
paragraph('The gain is an error reduction relative to Recent. A negative number means the candidate policy had higher MSE. A larger memory condition is not assumed to be better merely because it contains older observations.')

page()
heading('7 Comparison with the preceding real video trial')
paragraph('The baseline is the matching v5 reader and condition. The candidate is v6 Full. Both trials used the same fixed development data and target features. These comparisons describe a revision on an already used development set, not a new test set. The saved old scalar predictions were reused, with identical query keys and no old-model reinference.')
old_rows=[]
for h in HORIZONS:
    for old_policy in ('recent','full'):
        for s in GROUPS:
            c=reference['checks'][f'old_{old_policy}_vs_new_full_h{h}'][s]
            base,new=mean_mse(s,h,old_policy,True),mean_mse(s,h,'full')
            assert math.isclose(base-new,c['point'][0],abs_tol=1e-14)
            old_rows.append([s,h*.5,old_policy.title(),fmt(base),fmt(new),
                             fmt(c['point'][0],signed=True),pct(c['point'][0],base),interval(c),evidence(c)])
table('All matched v5 baseline and v6 Full comparisons',
      ['Seed','Seconds','v5 baseline','v5 MSE','v6 Full MSE','Delta','Gain percent','95 percent CI of Delta','Evidence'],
      old_rows,[1.0,1.0,1.7,2.0,2.0,2.1,1.3,4.8,1.3],'fixed_old_reference',8.1)
paragraph('The original main comparisons were at 8 and 32 seconds. Shorter-horizon rows are retained here so that the reader can see favorable and unfavorable changes alike. Interpretation codes use the same paired participant interval as Experiment A.')

page()
heading('8 Experiment B  Value of selecting a deletion')
paragraph('Source refers to the trajectory that produced the current historical memory. Baseline refers to the deletion made at that state. All alternatives then use the same 32 Reservoir continuation steps. The joint score averages all four horizons; the 8 and 32 second scores retain the deletion selected by that joint score. The tables show every reported validation comparison, including the small uncertain effect. Delta is baseline cost minus oracle cost.')
for gi,group in enumerate(('FIFO_FIFO','FIFO_Reservoir','Reservoir_FIFO','Reservoir_Reservoir')):
    if gi==2:
        page();heading('8 1 Reservoir source states',2)
    source,baseline=group.split('_')
    b_rows=[]
    for s in GROUPS:
        c=deletion['checks'][group][s]
        for i,label in enumerate(('Joint','8 seconds','32 seconds')):
            b_rows.append([s,label,fmt(c['point'][i],10,True),interval(c,i,10),evidence(c,i)])
    table(f'Current source {source} and deletion baseline {baseline}',
          ['Seed','Score','Delta cost','95 percent CI of Delta','Evidence'],
          b_rows,[1.7,2.4,3.9,7.7,1.5],f'B_{group}',9.2)
paragraph('The future-informed oracle measures headroom available in the evaluated states. Its joint advantage follows from choosing the best joint objective among all 17 candidates. Whether an online learned policy can obtain that advantage, using only available information, remains a separate research question.')

page()
heading('9 Experiment C  Continuous irreversible replay')
paragraph('Each table reports every seed, horizon and policy for a fixed split and capacity. FIFO MSE, Reservoir MSE and Uniform MSE are participant-balanced means. The two Delta columns subtract the alternative from FIFO; positive values favor that alternative. These are descriptive comparisons with no confidence interval supplied by the executed replay summary. They should not be described as statistically reliable improvements.')
paragraph('The main historical capacity is 16; capacities 8 and 32 are fixed secondary checks. The separate recent buffer contains 8 observations. Training scores use 29 participants and validation scores use 10. Pooled rows average the three readers within the same split. Training and validation scores are not directly paired with one another.')
replay_table_index=0
for split in ('validation','train'):
    for cap in (16,8,32):
        if replay_table_index in (2,4):
            page();heading('9 Experiment C  Continued replay scores')
        replay_table_index+=1
        c_rows=[]
        for h in HORIZONS:
            for s in GROUPS:
                f,r,u=(replay_mse(s,split,p,cap,h) for p in ('FIFO','Reservoir','Uniform'))
                c_rows.append([s,h*.5,fmt(f),fmt(r),fmt(u),fmt(f-r,signed=True),fmt(f-u,signed=True)])
        table(f'{split.title()} replay with historical capacity {cap}',
              ['Seed','Seconds','FIFO MSE','Reservoir MSE','Uniform MSE','FIFO minus Reservoir','FIFO minus Uniform'],
              c_rows,[1.2,1.2,2.8,2.8,2.8,3.2,3.2],f'C_{split}_{cap}',9)
paragraph('The execution audit checked all 456,192 trace rows: fixed random draws, capacity, chronological ages and no reentry of deleted payloads. This establishes faithful execution of the replay. It is distinct from an effect size or a test of action understanding.')

page()
heading('9 1 Evaluation that remains unavailable',2)
paragraph('The raw RGB videos are at 60 frames per second, while the fine action annotations refer to a 30 frame per second sequence. The original producer and its exact starting frame have not been established. A selected raw video contains 21,277 frames, whereas the associated capture timestamp metadata contain 21,276 entries and no embedded capture timecode resolves the annotation origin. A consumer that reads 30 fps labels is not proof of the raw frame used for annotation frame zero.')
table('Alignment sensitivity under two unverified sampling phases',
      ['Split','Different verb sets','Observations','Fraction percent'],
      [['Training','1,234','34,816',f'{1234/34816*100:.3f}'],
       ['Validation','493','15,872',f'{493/15872*100:.3f}'],
       ['All development','1,727','50,688',f'{1727/50688*100:.3f}']],
      [4.3,4.3,4.3,4.3],'alignment_sensitivity',10)
paragraph('This sensitivity calculation does not establish which phase is correct. Therefore no exact label-to-feature join was selected. Verb prediction and negative log likelihood, within-activity drift, response to an activity switch, different-activity donor corruption, and recovery at 0, 1, 2, 4, 8 and 16 steps remain unavailable. The visual MSE experiments above do not require action labels and their numerical results remain valid within that scope.')

heading('9 2 Execution and evidence records',2)
paragraph('The v6 training job completed successfully on one NVIDIA A100 with four scientific threads in 50 minutes 14 seconds. The fixed evaluation completed successfully in 28 minutes 10 seconds using four scientific CPU threads. Before inference, all six best and latest checkpoints were checked for completion, finite weights, source and cache binding, optimizer settings, schedule, random state and the earliest selected validation minimum. The final audit repeated the saved proof and verified all 78 feature files and the reported output hashes.')
paragraph('Independent numerical replay of the stored raw scores reproduced the existing summaries and intervals. This was a numerical audit of the same predictions, not another experiment or an independent replication. Raw model weights, features and large scalar or trace files remain on the course HPC. No additional science was run to produce this report.')
table('Traceable compact source records', ['Record','SHA256'],[
    ['Training report',SOURCES[str((NEW/'training_report.json').relative_to(PROJECT))]],
    ['Diagnostic protocol',SOURCES[str((NEW/'protocol.json').relative_to(PROJECT))]],
    ['Visual health summary',SOURCES[str((NEW/'visual_health.json').relative_to(PROJECT))]],
    ['Fixed previous trial comparison',SOURCES[str((NEW/'fixed_old_reference_comparison.json').relative_to(PROJECT))]],
    ['Deletion summary',SOURCES[str((NEW/'deletion_summary.json').relative_to(PROJECT))]],
    ['Replay participant summary',SOURCES[str((NEW/'replay_participant_summary.json').relative_to(PROJECT))]],
    ['Final source and numeric audit',SOURCES[str((NEW/'root_final_audit_7268.json').relative_to(PROJECT))]],
], [4.3,12.9],'sources',8.5)
paragraph('Source folder: runs/phase1_v6_revision/native7251_7268. Previous-trial means and training curves are in runs/phase1_v5_real_evaluation/native_completed7209_diagnostic7226. The epoch tables that follow retain every training and validation mean before the final interpretation.',size=9)

for s in SEEDS:
    page()
    heading(f'10 Exact training trajectory for seed {s}')
    paragraph('Each row is one completed epoch. Train and validation columns are the full normalized MSE averaged across the four conditions and horizons. Delta is v5 minus v6 at the same epoch; positive values favor v6. These point differences have no confidence interval. An asterisk marks the selected validation minimum within that trial.')
    old_best=min(old_curves[s],key=lambda r:float(r['validation_full_mse']))['epoch']
    new_best=min(new_curves[s],key=lambda r:float(r['validation_full_mse']))['epoch']
    e_rows=[]
    for a,b in zip(old_curves[s],new_curves[s]):
        at,bt,av,bv=(float(r[k]) for r,k in ((a,'train_full_mse'),(b,'train_full_mse'),
                                              (a,'validation_full_mse'),(b,'validation_full_mse')))
        e_rows.append([a['epoch'],fmt(at),fmt(bt),fmt(at-bt,signed=True),
                       fmt(av)+('*' if a['epoch']==old_best else ''),
                       fmt(bv)+('*' if b['epoch']==new_best else ''),fmt(av-bv,signed=True)])
    table(f'All 30 epochs for seed {s}',
          ['Epoch','v5 train','v6 train','Delta train','v5 validation','v6 validation','Delta validation'],
          e_rows,[1.2,2.5,2.5,2.7,2.8,2.8,2.7],f'epochs_{s}',9.2)
    paragraph('The selected checkpoint uses the original earliest validation minimum rule. The same-epoch differences in this table describe optimization and generalization behavior; they do not replace the paired participant evaluation of selected checkpoints.',size=10)

page()
heading('11 Summary and interpretation')
paragraph('The experiments provide measurable evidence that historical visual context can help longer-horizon feature prediction on the fixed validation recordings. For the pooled v6 readers, the 8 second MSE decreases from '+fmt(mean_mse('pooled',16,'recent'))+' to '+fmt(mean_mse('pooled',16,'full'))+', an absolute decrease of '+fmt(mean_mse('pooled',16,'recent')-mean_mse('pooled',16,'full'))+' and a relative decrease of '+pct(mean_mse('pooled',16,'recent')-mean_mse('pooled',16,'full'),mean_mse('pooled',16,'recent'))+' percent. Its 95 percent interval is '+interval(health['full_vs_recent_paired_checks']['16']['pooled'])+'. At 32 seconds the pooled error decreases from '+fmt(mean_mse('pooled',64,'recent'))+' to '+fmt(mean_mse('pooled',64,'full'))+', with relative decrease '+pct(mean_mse('pooled',64,'recent')-mean_mse('pooled',64,'full'),mean_mse('pooled',64,'recent'))+' percent and interval '+interval(health['full_vs_recent_paired_checks']['64']['pooled'])+'.')
paragraph('Training variation still matters. Seeds 17 and 43 have positive history-benefit intervals at both longer horizons; seed 29 has small positive estimates with intervals crossing zero. At 0.5 and 2 seconds, Full has no supported positive advantage over Recent. The pooled 0.5 second difference is negative, and the pooled 2 second interval crosses zero. Historical context therefore has a horizon-dependent effect in this experiment.')
paragraph('The deletion experiment shows that many evaluated memory states have room for a better deletion choice. For example, under a Reservoir source and Reservoir baseline, the pooled oracle decreases joint cost by '+fmt(deletion['checks']['Reservoir_Reservoir']['pooled']['point'][0],10)+', with interval '+interval(deletion['checks']['Reservoir_Reservoir']['pooled'],0,10)+'. The smallest uncertain comparison is Reservoir source, FIFO baseline, seed 43 at 32 seconds: Delta '+fmt(deletion['checks']['Reservoir_FIFO']['43']['point'][2],10)+', interval '+interval(deletion['checks']['Reservoir_FIFO']['43'],2,10)+'. These effects concern a future-informed oracle and do not yet establish that a learned online selector can obtain them.')
paragraph('Increasing weight decay from 0.05 to 0.20 slightly improves the best validation mean in every training run, but the matched previous-trial tables show that the change is not uniformly better across seeds and horizons. Replay is complete and respects irreversible deletion; its policy scores are descriptive and often close. A faithful replay audit should not be interpreted as a separate performance success.')
paragraph('As a compact count of the detailed evidence, 6 of the 8 original longer-horizon history comparisons have positive intervals; 9 of the 16 original longer-horizon previous-trial comparisons do; and 47 of the 48 deletion comparisons do. These rows share readers and participants. The counts are not an empirical probability of success and are not a substitute for the effect sizes, denominators and intervals reported above.')
paragraph('The present interpretation is partial support for longer-horizon visual memory and for oracle deletion headroom, with incomplete robustness across training runs and unresolved action-semantic evaluation. We defer an overall pass or fail judgment for the research. The original numerical outputs and adoption rules remain preserved. Completion of the visual calculations does not establish action understanding, a deployable learned deletion policy, or independent held-out confirmation.',bold=True)

doc.core_properties.title='Visual Memory Experiments'
doc.core_properties.subject='Experimental design evaluation scores and detailed results'
doc.core_properties.author=''
doc.core_properties.comments=''
output=ROOT/'CV_Visual_Memory_Experiment_Report_English.docx'
doc.save(output)

# Validate data coverage and retain unrounded records for reproducibility.
assert TABLE_COUNTS['A_absolute']==16
assert TABLE_COUNTS['A_full_improvements']==16
assert TABLE_COUNTS['A_policy_point_improvements']==32
assert TABLE_COUNTS['fixed_old_reference']==32
assert sum(TABLE_COUNTS[f'B_{g}'] for g in deletion['checks'])==48
assert sum(TABLE_COUNTS[f'C_{split}_{cap}'] for split in ('train','validation') for cap in (8,16,32))==96
assert sum(TABLE_COUNTS[f'epochs_{s}'] for s in SEEDS)==90
assert hashlib.sha256(original_report.read_bytes()).hexdigest()==original_report_hash
(OUT/'table_content.json').write_text(json.dumps(TABLE_DATA,ensure_ascii=False,indent=2)+'\n')
manifest={
    'created_utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),
    'output':str(output),'sha256':hashlib.sha256(output.read_bytes()).hexdigest(),
    'language':'English','style':'Monochrome academic Times New Roman no header fill or colored graphics',
    'summary_last':True,'overall_decision':'DEFERRED_AS_REQUESTED',
    'table_row_counts':TABLE_COUNTS,'source_sha256':SOURCES,
    'all_source_transfer_hashes_rechecked_locally':True,
    'existing_machine_outputs_and_protocols_modified':False,
    'new_model_inference_training_or_scientific_audit':False,
    'new_CIs_created':False,'relative_percentages_and_pooled_MSE_are_descriptive_arithmetic':True,
    'all_C_original_216_group_means_displayed_in_72_nonpooled_comparison_rows':True,
    'pooled_C_rows_added_by_arithmetic':24,
    'raw_large_payloads_downloaded':False,
    'original_report_sha256_unchanged':original_report_hash,
    'figure':str(figure),'render_QA_pending':True,
}
(OUT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'output':str(output),'sha256':manifest['sha256'],'tables':table_number,'table_rows':sum(TABLE_COUNTS.values())}))
