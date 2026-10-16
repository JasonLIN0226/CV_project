"""Build a figure-led academic report from already audited local evidence only.

No model, data generator, training code, remote connection or experiment is run.
"""
from pathlib import Path
import csv
import hashlib
import json
import statistics
import subprocess
import textwrap

from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.pdfbase.pdfmetrics import stringWidth
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

BASE = Path(__file__).resolve().parent
ROOT = BASE.parents[1]
OUT = BASE / 'report' / 'academic'
FIG = OUT / 'figures'
FIG.mkdir(parents=True, exist_ok=True)
NATIVE = BASE / 'native7590/output/job_7590'
PRED = ROOT / 'runs/phase1_predictor_repair_v2'
PN = PRED / 'native_final7579'
PDFTOPPM = '/Users/linxiangyu/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/pdftoppm'

def read(p):
    return json.loads(Path(p).read_text())

def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()

summary = read(NATIVE / 'summary.json')
audit = read(NATIVE / 'independent_audit.json')
outcome = read(PN / 'evaluation/job_7575/outcome.json')
nativeproof = read(PN / 'final_native_verification.json')
baselines = read(NATIVE / 'baselines.json')['results']
fixed = list(csv.DictReader((PRED / 'report/all_64_fixed_checks.csv').open()))
assert audit['status'] == 'INDEPENDENT_SCALAR_AUDIT_PASS'
assert audit['summary_sha256'] == sha(NATIVE / 'summary.json')
assert summary['raw_rows'] == 55296
assert outcome['status'] == 'DEVELOPMENT_REJECTED'
assert all(nativeproof['jobs'][str(j)]['state'] == 'COMPLETED' for j in (7574, 7575, 7576))
assert not outcome['confirmation_opened']

def txt(c, x, y, text, size=11, bold=False, align='left', gray=0):
    c.setFillGray(gray)
    c.setFont('Helvetica-Bold' if bold else 'Helvetica', size)
    {'left': c.drawString, 'center': c.drawCentredString, 'right': c.drawRightString}[align](x, y, text)

def wrapped(c, x, y, text, width, size=11, leading=None, bold=False):
    leading = leading or size * 1.25
    words = text.split()
    line = ''
    for word in words:
        proposed = (line + ' ' + word).strip()
        if stringWidth(proposed, 'Helvetica-Bold' if bold else 'Helvetica', size) > width and line:
            txt(c, x, y, line, size, bold)
            y -= leading
            line = word
        else:
            line = proposed
    if line:
        txt(c, x, y, line, size, bold)
    return y - leading

def line(c, x1, y1, x2, y2, gray=0, width=1, dash=None):
    c.setStrokeGray(gray)
    c.setLineWidth(width)
    c.setDash(dash or [])
    c.line(x1, y1, x2, y2)
    c.setDash()

def arrow(c, x1, y1, x2, y2, gray=0):
    line(c, x1, y1, x2, y2, gray, 1)
    import math
    a = math.atan2(y2-y1, x2-x1)
    for d in (-.45, .45):
        line(c, x2, y2, x2-7*math.cos(a+d), y2-7*math.sin(a+d), gray, 1)

def box(c, x, y, w, h, label, fill=1, size=11, bold=False):
    c.setFillGray(fill)
    c.setStrokeGray(.15)
    c.setLineWidth(.8)
    c.rect(x, y, w, h, fill=1, stroke=1)
    rows = label.split('\n')
    for i, row in enumerate(rows):
        txt(c, x+w/2, y+h/2+(len(rows)-1)*size*.62-i*size*1.25-3, row, size, bold, 'center')

def figure(name, width, height, draw):
    pdf = FIG / f'{name}.pdf'
    c = canvas.Canvas(str(pdf), pagesize=(width, height))
    draw(c)
    c.showPage()
    c.save()
    subprocess.run([PDFTOPPM, '-png', '-singlefile', '-r', '220', str(pdf), str(FIG/name)], check=True, stdout=subprocess.DEVNULL)
    return FIG / f'{name}.png'

def scene(c, x, y, kind, state=0):
    """Conceptual, drawn scene, explicitly not a test output."""
    c.setFillGray(1);c.setStrokeGray(.5);c.rect(x,y,115,45,fill=1)
    if kind == 'cue':
        line(c,x+55,y+3,x+55,y+42,.7)
        c.setFillGray(.2)
        c.rect(x+15 if state==0 else x+78,y+25,22,9,fill=1,stroke=0)
        c.circle(x+55,y+10,4,fill=1,stroke=0)
    elif kind == 'blank':
        line(c,x+55,y+3,x+55,y+42,.7)
        c.setFillGray(.2);c.circle(x+55,y+20,4,fill=1,stroke=0)
    elif kind == 'lane':
        line(c,x+55,y+3,x+55,y+42,.7)
        c.setFillGray(.2);c.circle(x+(25 if state==0 else 88),y+30,4,fill=1,stroke=0)
    elif kind == 'move':
        c.setFillGray(.85);c.rect(x+65,y+3,37,39,fill=1,stroke=0)
        c.setFillGray(.2);c.circle(x+17,y+(12 if state==0 else 33),4,fill=1,stroke=0)
        arrow(c,x+25,y+(12 if state==0 else 33),x+52,y+23)
    elif kind == 'hidden':
        c.setFillGray(.85);c.rect(x+38,y+3,46,39,fill=1,stroke=0)
        txt(c,x+61,y+18,'hidden',9,False,'center',.2)
    elif kind == 'return':
        c.setFillGray(.85);c.rect(x+15,y+3,42,39,fill=1,stroke=0)
        c.setFillGray(.2);c.circle(x+92,y+(12 if state==0 else 33),4,fill=1,stroke=0)

def task_diagram(c):
    for x, label in ((125,'Earlier observed evidence'),(315,'Same recent images'),(505,'Different future outcome')):
        txt(c,x,257,label,11,True,'center')
    txt(c,10,223,'A  Gate',12,True)
    for s,y in [(0,190),(1,137)]:
        scene(c,68,y,'cue',s);scene(c,258,y,'blank');scene(c,448,y,'lane',s)
        arrow(c,188,y+22,251,y+22);arrow(c,378,y+22,441,y+22)
        txt(c,48,y+17,'A' if s==0 else 'B',10,True)
    txt(c,10,110,'B  Motion',12,True)
    for s,y in [(0,60),(1,7)]:
        scene(c,68,y,'move',s);scene(c,258,y,'hidden');scene(c,448,y,'return',s)
        arrow(c,188,y+22,251,y+22);arrow(c,378,y+22,441,y+22)
        txt(c,48,y+17,'A' if s==0 else 'B',10,True)

taskfig = figure('01_tasks', 620, 278, task_diagram)

def pipeline(c):
    box(c,8,80,110,54,'Observed\nimages',1)
    box(c,151,80,119,54,'Frozen I-JEPA\nencoder',.94)
    box(c,303,80,125,54,'Cached visual\nfeatures',1)
    box(c,461,80,148,54,'Small predictor\nfuture features',.94)
    for a,b in [(118,151),(270,303),(428,461)]:arrow(c,a+3,107,b-4,107)
    box(c,309,8,112,43,'Recent buffer\n4 frames',1,10)
    box(c,439,8,170,43,'Historical evidence\nnone  all  or 2 entries',1,10)
    arrow(c,366,54,490,75);arrow(c,522,54,540,75)
    txt(c,10,35,'Known actions and timestamps',11,True)
    txt(c,10,16,'No future images or hidden states enter prediction',10)
    txt(c,315,151,'Predict at 3 and 6 future steps',11,True,'center')

pipefig = figure('02_pipeline', 620, 162, pipeline)

splits = [('test_id','Ordinary'),('test_long_gap','Long gap'),('test_distractor_shift','Distractor shift')]
gainrows = []
condition_mse = []
for task in ('gate','motion'):
    metric_rows = read(PN / f'evaluation/job_7575/{task}/evaluation/metrics.json')['per_seed']
    for split, label in splits:
        means = {}
        for cond in ('recent_only','full_history','privileged_small_budget'):
            vals = [r['metrics']['normalized_mse']['mean'] for r in metric_rows if r['split']==split and r['horizon']==6 and r['condition']==cond]
            assert len(vals)==3
            means[cond]=statistics.mean(vals)
        condition_mse.append(dict(task=task,split=split,**means))
        for h in (3,6):
            r = next(r for r in fixed if r['task']==task and r['kind']=='candidate_vs_reference_recent' and r['slice']==f'{task}/{split}/h{h}/seedall_seeds_mean')
            gainrows.append(dict(task=task,split=split,horizon=h,mean=float(r['mean'])*100,low=float(r['low'])*100,high=float(r['high'])*100))

def gains(c):
    for panel,task in enumerate(('gate','motion')):
        l=101+panel*306;r=l+190;b=43;t=185
        xp=lambda z:l+(z+5)/45*(r-l)
        txt(c,(l+r)/2,217,task.title(),13,True,'center')
        for tick in (0,10,20,30,40):
            line(c,xp(tick),b,xp(tick),t,.85 if tick!=10 else .4,.6,[3,3] if tick==10 else None)
            txt(c,xp(tick),b-16,str(tick),9,False,'center')
        for i,(split,label) in enumerate(splits):
            y=166-i*51
            txt(c,l-8,y-2,label,10,False,'right')
            for h,off in [(3,8),(6,-8)]:
                v=next(x for x in gainrows if x['task']==task and x['split']==split and x['horizon']==h)
                yy=y+off
                line(c,xp(v['low']),yy,xp(v['high']),yy,0,1)
                for z in (v['low'],v['high']):line(c,xp(z),yy-2,xp(z),yy+2)
                c.setFillGray(0 if h==3 else 1);c.setStrokeGray(0)
                if h==3:c.rect(xp(v['mean'])-2.7,yy-2.7,5.4,5.4,fill=1,stroke=1)
                else:c.circle(xp(v['mean']),yy,3,fill=1,stroke=1)
        txt(c,(l+r)/2,5,'Feature MSE reduction  percent',10,False,'center')
    c.setFillGray(0);c.rect(187,241,5,5,fill=1);txt(c,199,240,'3 steps',10)
    c.setFillGray(1);c.circle(286,243,3,fill=1,stroke=1);txt(c,297,240,'6 steps',10)
    txt(c,385,240,'Dashed line  10 percent floor',10)

gainfig = figure('03_prediction_gains',620,260,gains)

def deletion_method(c):
    txt(c,10,207,'Common parent',12,True)
    box(c,10,137,175,47,'Gate  cue  +  frame  +  frame',1,10)
    box(c,10,72,175,47,'Motion  start  +  end  +  shutter',1,10)
    txt(c,10,35,'3 real historical entries',11,True)
    for y,label in [(157,'Delete entry 1'),(92,'Delete entry 2'),(27,'Delete entry 3')]:
        box(c,270,y,126,38,label,.94,11)
        box(c,457,y,150,38,'2 entries retained',1,11)
        arrow(c,192,122,263,y+19)
        arrow(c,401,y+19,450,y+19)
    txt(c,445,207,'Same memory capacity',12,True)
    txt(c,315,5,'Same 4 recent frames  actions  predictor  and future',10,False,'center')

deletefig=figure('04_equal_capacity',620,225,deletion_method)

def reliability(c):
    labels=['Ordinary','Long gap','Distractor shift']
    for panel,task in enumerate(('gate','motion')):
        l=95+panel*310;r=l+185;b=39;t=162
        xp=lambda z:l+z/100*(r-l)
        txt(c,(l+r)/2,193,task.title(),13,True,'center')
        for k in (0,50,100):
            line(c,xp(k),b,xp(k),t,.85,.6)
            txt(c,xp(k),b-15,str(k),9,False,'center')
        for i,((split,_),lab) in enumerate(zip(splits,labels)):
            v=summary['tasks'][task]['scope'][split];y=141-i*43
            txt(c,l-7,y-1,lab,10,False,'right')
            c.setFillGray(.8);c.rect(l,y-8,185,18,fill=1,stroke=0)
            c.setFillGray(.12);c.rect(l,y-8,185*v['contrasts_passed']/v['contrasts_total'],18,fill=1,stroke=0)
            txt(c,r+8,y-2,f"{v['contrasts_passed']}/{v['contrasts_total']}",11,True)
        txt(c,(l+r)/2,7,'Critical deletion checks passed  percent',9,False,'center')

reliablefig=figure('05_reliability',650,210,reliability)

id6 = [r for r in summary['contrasts'] if r['split']=='test_id' and r['horizon']==6]
def effect_forest(c):
    for panel,task in enumerate(('gate','motion')):
        l=79+panel*310;r=l+214;b=47;t=286
        xp=lambda z:l+(z+22)/127*(r-l)
        txt(c,(l+r)/2,320,task.title(),13,True,'center')
        for tick in (-20,0,20,40,60,80,100):
            line(c,xp(tick),b,xp(tick),t,.85 if tick!=0 else .4,.6)
            txt(c,xp(tick),b-16,str(tick),9,False,'center')
        line(c,xp(10),b,xp(10),t,.3,.7,[3,3])
        selected=[x for x in id6 if x['task']==task]
        selected=sorted(selected,key=lambda x:(x['role']!='cue' and x['role']!='start', x['seed']))
        for i,row in enumerate(selected):
            y=(246-i*71) if task=='gate' else (267-i*40)
            lab=f"Seed {row['seed']}" if task=='gate' else f"{row['role'].title()} {row['seed']}"
            txt(c,l-7,y-3,lab,10,False,'right')
            for metric,off in [('feature_mse',4),('task_error',-4)]:
                v=row['metrics'][metric]['relative_reduction'];yy=y+off
                line(c,xp(v['low']*100),yy,xp(v['high']*100),yy,.1,1)
                for z in (v['low']*100,v['high']*100):line(c,xp(z),yy-2,xp(z),yy+2)
                c.setFillGray(0 if metric=='feature_mse' else 1);c.setStrokeGray(0)
                c.circle(xp(v['mean']*100),yy,2.9,fill=1,stroke=1)
        txt(c,(l+r)/2,7,'Retaining critical evidence reduces error  percent',9,False,'center')
    c.setFillGray(0);c.circle(128,347,3,fill=1,stroke=1);txt(c,138,344,'Feature error',10)
    c.setFillGray(1);c.circle(245,347,3,fill=1,stroke=1);txt(c,255,344,'Lane or position error',10)
    txt(c,410,344,'Dashed  10 percent effect floor',10)

forestfig=figure('06_deletion_effects',635,363,effect_forest)

bnames=[('FIFO_delete_oldest','FIFO'),('uniform_delete','Uniform'),('reject_newest','Reject newest'),('future_feature_oracle','Future oracle')]
bdata=[]
for key,label in bnames:
    vals=[r for r in baselines if r['task']=='motion' and r['split']=='test_id' and r['horizon']==6 and r['baseline']==key]
    assert len(vals)==3
    bdata.append(dict(key=key,label=label,feature=statistics.mean(r['metrics']['feature_mse']['mean'] for r in vals),task=statistics.mean(r['metrics']['task_error']['mean'] for r in vals),seeds=[r['metrics']['feature_mse']['mean'] for r in vals]))

def baseline_figure(c):
    l=68;r=589;b=54;t=246
    yp=lambda z:b+z/.6*(t-b)
    for tick in (0,.1,.2,.3,.4,.5,.6):
        line(c,l,yp(tick),r,yp(tick),.85,.6)
        txt(c,l-10,yp(tick)-3,f'{tick:.1f}',10,False,'right')
    for i,v in enumerate(bdata):
        x=125+i*130
        c.setFillGray(.82);c.setStrokeGray(.15);c.rect(x-30,b,60,yp(v['feature'])-b,fill=1,stroke=1)
        for j,z in enumerate(v['seeds']):
            c.setFillGray(0);c.circle(x+(j-1)*8,yp(z),2.4,fill=1,stroke=0)
        txt(c,x,yp(max(v['seeds']))+12,f"{v['feature']:.4f}",11,True,'center')
        txt(c,x,b-22,v['label'],11,False,'center')
    txt(c,10,263,'Feature MSE',11,True)
    txt(c,330,282,'Motion   Ordinary condition   6 future steps',12,True,'center')
    txt(c,329,7,'Lower is better   Bars are three seed means   Dots are individual seed means',10,False,'center')

baselinefig=figure('07_baseline_headroom',630,300,baseline_figure)

# A transparent machine-readable figure ledger is retained alongside the report.
figdata={'prediction_gains_percent':gainrows,'six_step_condition_mse':condition_mse,'deletion_contrasts':summary['contrasts'],'baseline_summary':bdata}
(OUT/'figure_data.json').write_text(json.dumps(figdata,indent=2)+'\n')

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.left_margin=sec.right_margin=Inches(.72)
sec.top_margin=Inches(.64);sec.bottom_margin=Inches(.62)
sec.header_distance=Inches(.2);sec.footer_distance=Inches(.26)
for b in list(doc.styles.element.iter(qn('w:pBdr'))): b.getparent().remove(b)
for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Caption'):
    st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.09
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(23)
doc.styles['Title'].paragraph_format.space_after=Pt(8)
doc.styles['Subtitle'].font.size=Pt(11)
doc.styles['Subtitle'].paragraph_format.space_after=Pt(10)
doc.styles['Heading 1'].font.size=Pt(15)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(6)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(8)
doc.styles['Heading 2'].font.size=Pt(11.5)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(7)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(5)
doc.styles['Caption'].font.size=Pt(9)
doc.styles['Caption'].font.italic=False
doc.styles['Caption'].font.bold=False
doc.styles['Caption'].paragraph_format.line_spacing=1.04
doc.styles['Caption'].paragraph_format.space_after=Pt(8)

footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Phase One Visual Memory Study  |  ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

def p(text='',boldlead=None):
    q=doc.add_paragraph()
    if boldlead and text.startswith(boldlead):
        q.add_run(boldlead).bold=True;q.add_run(text[len(boldlead):])
    else:q.add_run(text)
    return q

new_page_pending=False
def h(text):
    global new_page_pending
    q=doc.add_heading(text,level=1)
    if new_page_pending:q.paragraph_format.page_break_before=True
    new_page_pending=False
    return q
def h2(text):return doc.add_heading(text,level=2)
def page():
    global new_page_pending
    new_page_pending=True

def image(path,caption,width=6.95):
    q=doc.add_paragraph();q.paragraph_format.space_after=Pt(4)
    q.paragraph_format.keep_with_next=True
    r=q.add_run();r.add_picture(str(path),width=Inches(width))
    r._r.xpath('.//wp:docPr')[0].set('descr',caption)
    doc.add_paragraph(caption,style='Caption')

def table(headers,rows,widths,small=9.5):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        z=OxmlElement('w:'+side);z.set(qn('w:val'),'single');z.set(qn('w:sz'),'5');z.set(qn('w:color'),'D9D9D9');borders.append(z)
    t._tbl.tblPr.append(borders)
    for i,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
            for side,value in [('top',75),('bottom',75),('left',95),('right',95)]:
                z=OxmlElement('w:'+side);z.set(qn('w:w'),str(value));z.set(qn('w:type'),'dxa');marg.append(z)
            tcpr.append(marg)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'333333' if i==0 else ('F4F4F4' if i%2==0 else 'FFFFFF'));tcpr.append(sh)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.05
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or (i>0 and (widths[j]>=4 or len(c.text)>45)) else WD_ALIGN_PARAGRAPH.CENTER
                for rr in para.runs:
                    rr.font.size=Pt(small);rr.bold=i==0;rr.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def cap(text):doc.add_paragraph(text,style='Caption')

# PAGE 1
doc.add_paragraph('Phase One Evidence for Bounded Visual Memory',style='Title')
doc.add_paragraph('Controlled prediction experiments and equal capacity deletion tests\nAcademic progress report  |  16 October 2026',style='Subtitle')
h('1  Research objective')
p('Our proposal asks which historical observations a model should retain when its memory is limited. The intended learning signal is the prediction cost of deleting an observation: a memory should matter when forgetting it makes useful future prediction worse. The research goal is a better retention decision, rather than a predictor that solves every possible future.')
p('Phase One tests the prerequisites: whether stored visual features contain the missing evidence, and whether a predictor can use it. The later deletion assay adds a bridge toward policy learning, while keeping the predictor and memory capacity fixed. No retention policy has been trained.')
image(taskfig,'Figure 1. Why the past is needed. Paired episodes A and B share recent images and actions but have different earlier evidence and future outcomes. Gate requires a transient cue; Motion requires pre-occlusion positions. These drawings are conceptual illustrations, not sampled test images or model outputs.',6.95)
h2('Three questions guide the experiments')
table(['Question','Evidence needed'],[
    ['Representation','Can the relevant cue or motion evidence be recovered from frozen image features?'],
    ['Prediction','Does supplying historical evidence improve forecasts over recent observations alone?'],
    ['Deletion signal','At the same memory capacity, does losing critical evidence cause more harm than losing a nuisance frame?']
],[1.35,5.55],9.5)
p('Scope. The original proposal places controlled prediction before selector training. The deletion assay is an additional transition experiment; it does not establish an online retention algorithm or replace the earlier prediction acceptance tests.',boldlead='Scope.')

# PAGE 2
page();h('2  Experimental design')
image(pipefig,'Figure 2. Shared prediction pipeline. A frozen I-JEPA encoder [1] maps images into cached features. The predictor sees recent context, legal history, actions and timestamps. Simulator truth is used only to score outputs. All retained entries are read; there is no separate retrieval or learned retention selector.',6.95)
table(['Component','Specification'],[
    ['Controlled tasks','Gate cue and occluded Motion'],
    ['Data per task','512 training pairs; 96 validation pairs; 128 test pairs in each of 3 conditions'],
    ['Test conditions','Ordinary; longer gaps or occlusion; changed visual distractors'],
    ['Prediction context','4 recent frames; historical budget B = 2; horizons h = 3 and h = 6'],
    ['Complete candidate models','2 tasks x seeds 17, 29, 43 = 6 models; 120 epochs and 7680 steps each'],
    ['Visual representation','Frozen I-JEPA; train-only normalization and 128-dimensional basis; full 5120-dimensional scored output'],
    ['Model selection','Earliest minimum of validation MSE, equally weighting 4 conditions and both horizons'],
    ['Data separation','Pairs stay in one split; no simulator state, sibling identity or future image is a predictor input']
],[1.6,5.3],9.3)
h2('What was compared')
table(['Condition','Information provided'],[
    ['Recent only','Recent buffer and actions; no historical evidence'],
    ['Full history','All available legal history; a diagnostic upper information reference'],
    ['Privileged small budget','Two important historical frames, chosen using task construction'],
    ['Content intervention','Replace the relevant historical content to test whether prediction depends on it']
],[1.6,5.3],9.3)
p('The last complete candidate used centered visual differences, attention pooling over observed history and explicit elapsed-time transport. These components were motivated by Set Transformer [2] and NLinear [3]; the candidate is an adaptation, not a reproduction or a novelty claim. Its empty-history output equals the frozen recent-only reference.')
p('Statistics. Feature error is mean squared error (MSE); task error measures lane or position. Reported prediction intervals use paired bootstrap resampling over 128 base trajectory pairs. Seed-averaged intervals are conditional on the three tested training seeds. Forks, siblings and seeds are not treated as extra independent episodes.',boldlead='Statistics.')

# PAGE 3
page();h('3  History helps prediction but support is uneven')
image(gainfig,'Figure 3. Relative feature-MSE reduction from the history-conditioned candidate versus frozen recent only. Points and paired 95 percent intervals average seeds within each pair. Positive is better; the dashed line is the 10 percent material-effect floor. Data: D1.',6.95)
table(['Task and condition','Recent only\nMSE','Full history\nMSE','Budget 2\nMSE','Gain and 95 percent CI'],[
    [task.title()+'  '+lab,
     *[f"{next(v for v in condition_mse if v['task']==task and v['split']==sp)[co]:.4f}" for co in ('recent_only','full_history','privileged_small_budget')],
     (lambda v:f"{v['mean']:.1f}%  [{v['low']:.1f}, {v['high']:.1f}]")(next(v for v in gainrows if v['task']==task and v['split']==sp and v['horizon']==6))]
    for task in ('gate','motion') for sp,lab in splits
],[1.95,.85,.85,.95,2.3],8.9)
cap('Six-step normalized feature errors, averaged over three seeds. The last column compares privileged two-memory prediction with recent only. Full history uses more memory and is an information reference, not a capacity-matched competitor. Data: D1.')
table(['Acceptance check','Gate','Motion'],[
    ['Evidence readable from features','3 / 3','3 / 3'],
    ['Privileged history benefit under original task rules','23 / 24','8 / 24'],
    ['Response to historical content intervention','24 / 24','24 / 24'],
    ['Fixed recent-reference feature comparisons','24 / 24','8 / 24'],
    ['Critical error improvement over prior predictor at the same budget','8 / 8','0 / 8']
],[4.3,1.3,1.3],9.1)
cap('Table 1. Counts of prespecified checks, not independent sample counts or percentages of project completion. The original task checks total 51 per task; the additional fixed-reference checks total 32 per task. All required checks had to pass. Data: D1.')
p('Interpretation. Gate passed 50/51 original and 32/32 additional checks, but the ordinary h6 task-benefit interval for seed 43 was inconclusive. Motion passed 35/51 and 8/32: its ordinary gains did not extend reliably to long occlusion and shifted distractors. Both tasks therefore failed the complete development acceptance protocol; fresh confirmation was not opened.',boldlead='Interpretation.')

# PAGE 4
page();h('4  Does deletion cost identify important memory')
p('We reused the six complete models and their existing feature caches. From the same parent set of three real observations, we removed each entry in turn and compared the resulting forecasts. Every fork retained exactly two historical entries, so a gain cannot be explained by a larger memory budget.')
image(deletefig,'Figure 4. Same-capacity intervention. Each branch is reset to its common parent and differs only in the deleted memory. Gate pools contain one cue and two post-cue frames; Motion pools contain start and end observations plus a shutter frame. Two fixed nuisance-time variants were tested. Pools are privileged diagnostics, not online policy states.',6.95)
p('Deletion contrast = loss after deleting critical evidence minus loss after deleting a nuisance frame. Positive values favor keeping the critical observation. The prospective entry rule required a positive 95 percent paired interval and at least 10 percent mean relative benefit, for both feature error and task error in every seed, horizon and critical role.')
image(reliablefig,'Figure 5. Reliability of critical-versus-nuisance deletion contrasts. Gate has one critical role, giving 6 checks per condition; Motion has two roles, giving 12. Each check uses 128 base pairs. Different denominators reflect the design, not different dataset sizes. Data: D2.',6.95)
table(['Task','Ordinary','Long gap','Distractor shift','Cross horizon checks'],[
    ['Gate','4 / 6','0 / 6','2 / 6','18 / 18'],
    ['Motion','12 / 12','12 / 12','12 / 12','18 / 18']
],[.85,1.15,1.15,1.45,2.3],9.4)
cap('Table 2. Cross-horizon tests choose a deletion using feature loss at h3 and score it at h6, and vice versa, against uniform deletion. This future-informed diagnostic is not a deployable past-only policy. Ranking alone is insufficient: Gate fails the critical-deletion requirement. Data: D2.')

# PAGE 5
page();h('5  Effect sizes and uncertainty')
image(forestfig,'Figure 6. Ordinary-condition deletion effects at h6, shown separately for every training seed and critical memory role. Filled points are feature error; open points are lane or position error. Bars are paired 95 percent intervals. The dashed line is the 10 percent mean-effect floor, not a requirement that the entire interval exceed 10 percent. Data: D2.',6.95)
selected=[next(r for r in id6 if r['task']=='gate' and r['seed']==43),next(r for r in id6 if r['task']=='motion' and r['seed']==17 and r['role']=='start'),next(r for r in id6 if r['task']=='motion' and r['seed']==17 and r['role']=='end')]
rows=[]
for r in selected:
    v=r['metrics']['feature_mse'];ci=v['difference']
    rows.append([f"{r['task'].title()} {r['role']}  seed {r['seed']}",f"{v['critical_mean']:.4f}",f"{v['nuisance_mean']:.4f}",f"{ci['mean']:.4f}\n[{ci['low']:.4f}, {ci['high']:.4f}]"])
table(['Selected contrast','Delete critical\nfeature MSE','Delete nuisance\nfeature MSE','Difference and 95 percent CI'],rows,[2.1,1.35,1.35,2.1],9.1)
cap('Table 3. Illustrative numeric contrasts at h6. Within each row, only the deletion differs. Losses across Gate and Motion are not calibrated to the same task difficulty. All seeds and both Motion roles are shown in Figure 6.')
p('Motion. Keeping the start observation reduced feature error by 76.3 to 77.8 percent across seeds; keeping the end observation reduced it by 28.3 to 30.5 percent. Position-error contrasts agreed. Both observations therefore carry complementary useful evidence, and all tested Motion deletion contrasts passed.',boldlead='Motion.')
p('Gate. Seed 43 failed both ordinary deletion checks. At h3 the lane-error interval crossed zero; at h6 feature benefit was only 6.5 percent and lane-error benefit 6.3 percent, with both intervals crossing zero. A seed-average result would hide this instability.',boldlead='Gate.')
p('Why Motion can pass this test while failing prediction acceptance. Reliable deletion cost asks which forgetting choice is worse. Absolute prediction quality asks whether the forecast is sufficiently good. The first can be reliable while the second remains inadequate; neither result cancels the other.',boldlead='Why Motion can pass this test while failing prediction acceptance.')

# PAGE 6
page();h('6  A useful signal does not prove a better policy')
image(baselinefig,'Figure 7. Descriptive one-decision baseline costs in the ordinary Motion pools at h6. Bars average the same three seeds; dots show their separate means. The future-feature oracle chooses using unavailable future information and is only an optimistic reference. This is not a streaming-policy benchmark. Data: D2.',6.95)
table(['Deletion rule','Feature MSE','Position MSE'],[[r['label'],f"{r['feature']:.6f}",f"{r['task']:.6f}"] for r in bdata],[3.3,1.8,1.8],9.6)
p('The oracle feature MSE is only 0.65 percent lower than rejecting the newest frame: 0.115305 versus 0.116061. This small headroom matters. These constructed pools show which evidence is useful, but a simple rule already finds most of that value; they do not yet demonstrate a need for a learned selector.')
h2('What the experiments have achieved')
p('We have a working visual-feature prediction pipeline, paired causal input comparisons, auditable model checkpoints, and a fast same-capacity deletion evaluation. Motion supplies a reproducible signal about relative memory value. Gate supplies a concrete counterexample to assuming that feature-level utility is always stable across training seeds.')
h2('Limits that should shape the next experiment')
for lead,body in [
    ('Adaptive development data. ','The same development trajectories informed predictor repairs. No untouched-data confirmation has been run.'),
    ('Privileged pools. ','Candidate times were constructed using known task evidence. This does not show that a policy can discover them from past inputs alone.'),
    ('Single deletion. ','Every fork resets to its parent. Continuous irreversible forgetting, accumulated errors and long-term redundancy remain untested.'),
    ('Uncertain learning headroom. ','A learned policy must beat strong past-only rules in more varied memory states, not merely beat random deletion in an easy diagnostic pool.')
]: p(lead+body,boldlead=lead)

# PAGE 7
page();h('7  Evidence summary and research decision')
table(['Research claim','Gate','Motion'],[
    ['Relevant evidence exists in frozen features','Supported in tested splits','Supported in tested splits'],
    ['History improves prediction robustly under all acceptance checks','Not fully supported','Not fully supported'],
    ['Same-capacity deletion cost is reliable in ordinary conditions','Not supported for every seed','Supported in all tested seeds and roles'],
    ['A learned retention policy beats strong baselines','Not tested','Not tested'],
    ['Untouched-data independent confirmation','Not run','Not run']
],[3.15,1.85,1.9],9.4)
h2('Evidence and reproducibility')
p('D1 is the last complete six-model predictor experiment, with 720 epoch records and 69120 evaluation scalar records. A separate audit checked the raw axes, finite values and all 64 fixed-reference comparisons. D2 is the subsequent equal-capacity deletion assay: 55296 fork-horizon records and 180 independently recomputed scalar comparisons. The 128 trajectory pairs per task and condition, rather than these raw row totals, are the statistical units. Transferred compact evidence matched its native hashes.')
p('D2 performed no new fitting, feature encoding or GPU allocation; CPU inference took 163.7 seconds. Total settled GPU use across the proposal validation and repair attempts was 12723 seconds, or 3 h 32 min 03 s, within the approved 8 h cap. Incomplete later repair trials are excluded from the performance results above. No retention policy was trained.')
h2('Sources')
refs=[
    '[1] Assran et al. Self-Supervised Learning from Images with a Joint-Embedding Predictive Architecture. CVPR, 2023. https://arxiv.org/abs/2301.08243',
    '[2] Lee et al. Set Transformer A Framework for Attention-based Permutation-Invariant Neural Networks. ICML, 2019. https://proceedings.mlr.press/v97/lee19d.html',
    '[3] Zeng et al. Are Transformers Effective for Time Series Forecasting. AAAI, 2023. https://arxiv.org/abs/2205.13504',
    'D1  Frozen predictor experiment 7574, evaluation 7575 and independent audit 7576. Archived evaluation outcome, fixed-reference intervals and final native verification in phase1_predictor_repair_v2.',
    'D2  Equal-capacity deletion assay 7590. Archived summary, baseline estimates, independent audit and pair-level scalar contrasts in phase1_deletion_assay. The assay protocol was fixed before its outcomes; the model and data were already development-selected.'
]
for ref in refs:
    q=p(ref);q.paragraph_format.space_after=Pt(5)
    for rr in q.runs:rr.font.size=Pt(8.7)
h2('Conclusion')
p('Phase One has established useful evidence for the proposal, but it has not established a complete two-task pass. Both tasks retain relevant information in frozen visual features, and history improves prediction in important settings. Motion also provides reliable deletion-cost contrasts, supporting a narrowly scoped learnability pilot. Gate remains unstable, Motion prediction under long occlusion remains insufficient, and neither an online retention advantage nor fresh independent confirmation has been shown. The defensible achievement is a tested mechanism and a clear next experiment, not a completed memory-retention method.')

path=OUT/'Phase1_Academic_Report_with_Figures.docx'
doc.save(path)
sources=[NATIVE/'summary.json',NATIVE/'independent_audit.json',NATIVE/'baselines.json',PN/'evaluation/job_7575/outcome.json',PN/'final_native_verification.json',PRED/'report/all_64_fixed_checks.csv',PN/'evaluation/job_7575/gate/evaluation/metrics.json',PN/'evaluation/job_7575/motion/evaluation/metrics.json']
manifest={'report':str(path),'sha256':sha(path),'figures':7,'intended_pages':7,'all_pages_visually_verified':False,'sources':{str(x.relative_to(ROOT)):sha(x) for x in sources},'no_new_model_execution':True,'new_GPU_seconds':0}
(OUT/'report_manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
print(path)
