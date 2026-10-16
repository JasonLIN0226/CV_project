"""Build a figure-led academic report from already audited local evidence only.

No model, data generator, training code, remote connection or experiment is run.
"""
from pathlib import Path
import csv
import hashlib
import json
import statistics
import subprocess
import textwrap

from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.pdfbase.pdfmetrics import stringWidth
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

BASE = Path(__file__).resolve().parent
ROOT = BASE.parents[1]
OUT = BASE / 'report' / 'proposal_structured'
FIG = OUT / 'figures'
FIG.mkdir(parents=True, exist_ok=True)
NATIVE = BASE / 'native7590/output/job_7590'
PRED = ROOT / 'runs/phase1_predictor_repair_v2'
PN = PRED / 'native_final7579'
PDFTOPPM = '/Users/linxiangyu/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/pdftoppm'

def read(p):
    return json.loads(Path(p).read_text())

def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()

summary = read(NATIVE / 'summary.json')
audit = read(NATIVE / 'independent_audit.json')
outcome = read(PN / 'evaluation/job_7575/outcome.json')
nativeproof = read(PN / 'final_native_verification.json')
baselines = read(NATIVE / 'baselines.json')['results']
fixed = list(csv.DictReader((PRED / 'report/all_64_fixed_checks.csv').open()))
assert audit['status'] == 'INDEPENDENT_SCALAR_AUDIT_PASS'
assert audit['summary_sha256'] == sha(NATIVE / 'summary.json')
assert summary['raw_rows'] == 55296
assert outcome['status'] == 'DEVELOPMENT_REJECTED'
assert all(nativeproof['jobs'][str(j)]['state'] == 'COMPLETED' for j in (7574, 7575, 7576))
assert not outcome['confirmation_opened']

def txt(c, x, y, text, size=11, bold=False, align='left', gray=0):
    c.setFillGray(gray)
    c.setFont('Helvetica-Bold' if bold else 'Helvetica', size)
    {'left': c.drawString, 'center': c.drawCentredString, 'right': c.drawRightString}[align](x, y, text)

def wrapped(c, x, y, text, width, size=11, leading=None, bold=False):
    leading = leading or size * 1.25
    words = text.split()
    line = ''
    for word in words:
        proposed = (line + ' ' + word).strip()
        if stringWidth(proposed, 'Helvetica-Bold' if bold else 'Helvetica', size) > width and line:
            txt(c, x, y, line, size, bold)
            y -= leading
            line = word
        else:
            line = proposed
    if line:
        txt(c, x, y, line, size, bold)
    return y - leading

def line(c, x1, y1, x2, y2, gray=0, width=1, dash=None):
    c.setStrokeGray(gray)
    c.setLineWidth(width)
    c.setDash(dash or [])
    c.line(x1, y1, x2, y2)
    c.setDash()

def arrow(c, x1, y1, x2, y2, gray=0):
    line(c, x1, y1, x2, y2, gray, 1)
    import math
    a = math.atan2(y2-y1, x2-x1)
    for d in (-.45, .45):
        line(c, x2, y2, x2-7*math.cos(a+d), y2-7*math.sin(a+d), gray, 1)

def box(c, x, y, w, h, label, fill=1, size=11, bold=False):
    c.setFillGray(fill)
    c.setStrokeGray(.15)
    c.setLineWidth(.8)
    c.rect(x, y, w, h, fill=1, stroke=1)
    rows = label.split('\n')
    for i, row in enumerate(rows):
        txt(c, x+w/2, y+h/2+(len(rows)-1)*size*.62-i*size*1.25-3, row, size, bold, 'center')

def figure(name, width, height, draw):
    pdf = FIG / f'{name}.pdf'
    c = canvas.Canvas(str(pdf), pagesize=(width, height))
    draw(c)
    c.showPage()
    c.save()
    subprocess.run([PDFTOPPM, '-png', '-singlefile', '-r', '220', str(pdf), str(FIG/name)], check=True, stdout=subprocess.DEVNULL)
    return FIG / f'{name}.png'

def scene(c, x, y, kind, state=0):
    """Conceptual, drawn scene, explicitly not a test output."""
    c.setFillGray(1);c.setStrokeGray(.5);c.rect(x,y,115,45,fill=1)
    if kind == 'cue':
        line(c,x+55,y+3,x+55,y+42,.7)
        c.setFillGray(.2)
        c.rect(x+15 if state==0 else x+78,y+25,22,9,fill=1,stroke=0)
        c.circle(x+55,y+10,4,fill=1,stroke=0)
    elif kind == 'blank':
        line(c,x+55,y+3,x+55,y+42,.7)
        c.setFillGray(.2);c.circle(x+55,y+20,4,fill=1,stroke=0)
    elif kind == 'lane':
        line(c,x+55,y+3,x+55,y+42,.7)
        c.setFillGray(.2);c.circle(x+(25 if state==0 else 88),y+30,4,fill=1,stroke=0)
    elif kind == 'move':
        c.setFillGray(.85);c.rect(x+58,y+3,30,39,fill=1,stroke=0)
        c.setFillGray(.2);c.circle(x+17,y+(12 if state==0 else 33),4,fill=1,stroke=0)
        start_y = 12 if state == 0 else 33
        slope = (21 if state == 0 else -21) / 84
        arrow(c,x+25,y+start_y+8*slope,x+52,y+start_y+35*slope)
    elif kind == 'hidden':
        c.setFillGray(.85);c.rect(x+58,y+3,30,39,fill=1,stroke=0)
        txt(c,x+73,y+18,'hidden',8,False,'center',.2)
    elif kind == 'return':
        c.setFillGray(.85);c.rect(x+58,y+3,30,39,fill=1,stroke=0)
        c.setFillGray(.2);c.circle(x+101,y+(33 if state==0 else 12),4,fill=1,stroke=0)

def task_diagram(c):
    for x, label in ((125,'Earlier observed evidence'),(315,'Same recent images'),(505,'Different future outcome')):
        txt(c,x,257,label,11,True,'center')
    txt(c,10,223,'A  Gate',12,True)
    for s,y in [(0,190),(1,137)]:
        scene(c,68,y,'cue',s);scene(c,258,y,'blank');scene(c,448,y,'lane',s)
        arrow(c,188,y+22,251,y+22);arrow(c,378,y+22,441,y+22)
        txt(c,48,y+17,'A' if s==0 else 'B',10,True)
    txt(c,10,110,'B  Motion',12,True)
    for s,y in [(0,60),(1,7)]:
        scene(c,68,y,'move',s);scene(c,258,y,'hidden');scene(c,448,y,'return',s)
        arrow(c,188,y+22,251,y+22);arrow(c,378,y+22,441,y+22)
        txt(c,48,y+17,'A' if s==0 else 'B',10,True)

taskfig = figure('01_tasks', 620, 278, task_diagram)

def condition_timeline(c):
    txt(c,315,168,'What changes before the prediction is made',12,True,'center')
    for label,y,gap in [('Ordinary',111,'5 to 10 steps'),('Long gap',65,'20 to 32 steps'),('Distractor shift',19,'5 to 10 steps')]:
        txt(c,2,y+13,label,10,True)
        for x,w,lab in [(109,88,'Earlier clue'),(218,104,gap),(344,113,'4 recent frames'),(480,137,'Predict 3 or 6 ahead')]:
            box(c,x,y,w,32,lab,.94 if x==218 else 1,9)
        for x1,x2 in [(198,215),(323,341),(458,477)]:arrow(c,x1,y+16,x2,y+16)
    txt(c,316,3,'Gap is before the recent window   Prediction horizon is after the query',9,False,'center')

conditionfig = figure('11_condition_timeline',630,185,condition_timeline)

def pipeline(c):
    box(c,8,80,110,54,'Observed\nimages',1)
    box(c,151,80,119,54,'Frozen I-JEPA\nencoder',.94)
    box(c,303,80,125,54,'Cached visual\nfeatures',1)
    box(c,461,80,148,54,'Small predictor\nfuture features',.94)
    for a,b in [(118,151),(270,303),(428,461)]:arrow(c,a+3,107,b-4,107)
    box(c,309,8,112,43,'Recent buffer\n4 frames',1,10)
    box(c,439,8,170,43,'Historical evidence\nnone  all  or 2 entries',1,10)
    arrow(c,366,54,490,75);arrow(c,522,54,540,75)
    txt(c,10,35,'Known actions and timestamps',11,True)
    txt(c,10,16,'No future images or hidden states enter prediction',10)
    txt(c,315,151,'Predict at 3 and 6 future steps',11,True,'center')

pipefig = figure('02_pipeline', 620, 162, pipeline)

splits = [('test_id','Ordinary'),('test_long_gap','Long gap'),('test_distractor_shift','Distractor shift')]
gainrows = []
condition_mse = []
for task in ('gate','motion'):
    metric_rows = read(PN / f'evaluation/job_7575/{task}/evaluation/metrics.json')['per_seed']
    for split, label in splits:
        means = {}
        for cond in ('recent_only','full_history','privileged_small_budget'):
            vals = [r['metrics']['normalized_mse']['mean'] for r in metric_rows if r['split']==split and r['horizon']==6 and r['condition']==cond]
            assert len(vals)==3
            means[cond]=statistics.mean(vals)
        condition_mse.append(dict(task=task,split=split,**means))
        for h in (3,6):
            r = next(r for r in fixed if r['task']==task and r['kind']=='candidate_vs_reference_recent' and r['slice']==f'{task}/{split}/h{h}/seedall_seeds_mean')
            gainrows.append(dict(task=task,split=split,horizon=h,mean=float(r['mean'])*100,low=float(r['low'])*100,high=float(r['high'])*100))

def gains(c):
    for panel,task in enumerate(('gate','motion')):
        l=101+panel*306;r=l+190;b=43;t=185
        xp=lambda z:l+(z+5)/45*(r-l)
        txt(c,(l+r)/2,217,task.title(),13,True,'center')
        for tick in (0,10,20,30,40):
            line(c,xp(tick),b,xp(tick),t,.85 if tick!=10 else .4,.6,[3,3] if tick==10 else None)
            txt(c,xp(tick),b-16,str(tick),9,False,'center')
        for i,(split,label) in enumerate(splits):
            y=166-i*51
            txt(c,l-8,y-2,label,10,False,'right')
            for h,off in [(3,8),(6,-8)]:
                v=next(x for x in gainrows if x['task']==task and x['split']==split and x['horizon']==h)
                yy=y+off
                line(c,xp(v['low']),yy,xp(v['high']),yy,0,1)
                for z in (v['low'],v['high']):line(c,xp(z),yy-2,xp(z),yy+2)
                c.setFillGray(0 if h==3 else 1);c.setStrokeGray(0)
                if h==3:c.rect(xp(v['mean'])-2.7,yy-2.7,5.4,5.4,fill=1,stroke=1)
                else:c.circle(xp(v['mean']),yy,3,fill=1,stroke=1)
        txt(c,(l+r)/2,5,'Feature MSE reduction  percent',10,False,'center')
    c.setFillGray(0);c.rect(187,241,5,5,fill=1);txt(c,199,240,'3 steps',10)
    c.setFillGray(1);c.circle(286,243,3,fill=1,stroke=1);txt(c,297,240,'6 steps',10)
    txt(c,385,240,'Dashed line  10 percent benchmark',10)

gainfig = figure('03_prediction_gains',620,260,gains)

def deletion_method(c):
    txt(c,10,207,'Common parent',12,True)
    box(c,10,137,175,47,'Gate  cue  +  frame  +  frame',1,10)
    box(c,10,72,175,47,'Motion  start  +  end  +  shutter',1,10)
    txt(c,10,35,'3 real historical entries',11,True)
    for y,label in [(157,'Delete entry 1'),(92,'Delete entry 2'),(27,'Delete entry 3')]:
        box(c,270,y,126,38,label,.94,11)
        box(c,457,y,150,38,'2 entries retained',1,11)
        arrow(c,192,122,263,y+19)
        arrow(c,401,y+19,450,y+19)
    txt(c,445,207,'Same memory capacity',12,True)
    txt(c,315,5,'Same 4 recent frames  actions  predictor  and future',10,False,'center')

deletefig=figure('04_equal_capacity',620,225,deletion_method)

def reliability(c):
    labels=['Ordinary','Long gap','Distractors']
    for panel,task in enumerate(('gate','motion')):
        l=95+panel*310;r=l+185;b=39;t=162
        txt(c,(l+r)/2,193,task.title(),13,True,'center')
        for i,((split,_),lab) in enumerate(zip(splits,labels)):
            v=summary['tasks'][task]['scope'][split];y=141-i*43
            txt(c,l-7,y-1,lab,10,False,'right')
            for j in range(v['contrasts_total']):
                x=l+8+j*169/(v['contrasts_total']-1)
                c.setFillGray(.1 if j<v['contrasts_passed'] else 1)
                c.setStrokeGray(.25);c.setLineWidth(.8)
                c.circle(x,y+2,4.4,fill=1,stroke=1)
            txt(c,r+8,y-2,f"{v['contrasts_passed']}/{v['contrasts_total']} pass",10,True)
    c.setFillGray(.1);c.circle(156,32,4.4,fill=1,stroke=1);txt(c,167,28,'Criterion met',10)
    c.setFillGray(1);c.circle(293,32,4.4,fill=1,stroke=1);txt(c,305,28,'Criterion not met',10)
    txt(c,326,5,'Each dot is one seed  horizon  and evidence role comparison',10,False,'center')

reliablefig=figure('05_reliability',650,210,reliability)

id6 = [r for r in summary['contrasts'] if r['split']=='test_id' and r['horizon']==6]
def effect_forest(c):
    for panel,task in enumerate(('gate','motion')):
        l=79+panel*310;r=l+214;b=47;t=286
        xp=lambda z:l+(z+22)/127*(r-l)
        txt(c,(l+r)/2,320,task.title(),13,True,'center')
        for tick in (-20,0,20,40,60,80,100):
            line(c,xp(tick),b,xp(tick),t,.85 if tick!=0 else .4,.6)
            txt(c,xp(tick),b-16,str(tick),9,False,'center')
        line(c,xp(10),b,xp(10),t,.3,.7,[3,3])
        selected=[x for x in id6 if x['task']==task]
        selected=sorted(selected,key=lambda x:(x['role']!='cue' and x['role']!='start', x['seed']))
        for i,row in enumerate(selected):
            y=(246-i*71) if task=='gate' else (267-i*40)
            lab=f"Seed {row['seed']}" if task=='gate' else f"{row['role'].title()} {row['seed']}"
            txt(c,l-7,y-3,lab,10,False,'right')
            for metric,off in [('feature_mse',4),('task_error',-4)]:
                v=row['metrics'][metric]['relative_reduction'];yy=y+off
                line(c,xp(v['low']*100),yy,xp(v['high']*100),yy,.1,1)
                for z in (v['low']*100,v['high']*100):line(c,xp(z),yy-2,xp(z),yy+2)
                c.setFillGray(0 if metric=='feature_mse' else 1);c.setStrokeGray(0)
                c.circle(xp(v['mean']*100),yy,2.9,fill=1,stroke=1)
        txt(c,(l+r)/2,7,'Retaining critical evidence reduces error  percent',9,False,'center')
    c.setFillGray(0);c.circle(128,347,3,fill=1,stroke=1);txt(c,138,344,'Feature error',10)
    c.setFillGray(1);c.circle(245,347,3,fill=1,stroke=1);txt(c,255,344,'Lane or position error',10)
    txt(c,410,344,'Dashed  10 percent benchmark',10)

forestfig=figure('06_deletion_effects',635,363,effect_forest)

bnames=[('FIFO_delete_oldest','FIFO'),('uniform_delete','Uniform'),('reject_newest','Reject newest'),('future_feature_oracle','Future oracle')]
bdata=[]
for key,label in bnames:
    vals=[r for r in baselines if r['task']=='motion' and r['split']=='test_id' and r['horizon']==6 and r['baseline']==key]
    assert len(vals)==3
    bdata.append(dict(key=key,label=label,feature=statistics.mean(r['metrics']['feature_mse']['mean'] for r in vals),task=statistics.mean(r['metrics']['task_error']['mean'] for r in vals),seeds=[r['metrics']['feature_mse']['mean'] for r in vals]))

def baseline_figure(c):
    l=68;r=589;b=54;t=246
    yp=lambda z:b+z/.6*(t-b)
    for tick in (0,.1,.2,.3,.4,.5,.6):
        line(c,l,yp(tick),r,yp(tick),.85,.6)
        txt(c,l-10,yp(tick)-3,f'{tick:.1f}',10,False,'right')
    for i,v in enumerate(bdata):
        x=125+i*130
        c.setFillGray(.82);c.setStrokeGray(.15);c.rect(x-30,b,60,yp(v['feature'])-b,fill=1,stroke=1)
        for j,z in enumerate(v['seeds']):
            c.setFillGray(0);c.circle(x+(j-1)*8,yp(z),2.4,fill=1,stroke=0)
        txt(c,x,yp(max(v['seeds']))+12,f"{v['feature']:.4f}",11,True,'center')
        txt(c,x,b-22,v['label'],11,False,'center')
    txt(c,10,263,'Feature MSE',11,True)
    txt(c,330,282,'Motion   Ordinary condition   6 future steps',12,True,'center')
    txt(c,329,7,'Shorter bars mean less prediction error   Dots show the three training runs',10,False,'center')

baselinefig=figure('07_baseline_headroom',630,300,baseline_figure)

# A transparent machine-readable figure ledger is retained alongside the report.
figdata={'prediction_gains_percent':gainrows,'six_step_condition_mse':condition_mse,'deletion_contrasts':summary['contrasts'],'baseline_summary':bdata}
(OUT/'figure_data.json').write_text(json.dumps(figdata,indent=2)+'\n')

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.left_margin=sec.right_margin=Inches(.72)
sec.top_margin=Inches(.64);sec.bottom_margin=Inches(.62)
sec.header_distance=Inches(.2);sec.footer_distance=Inches(.26)
for b in list(doc.styles.element.iter(qn('w:pBdr'))): b.getparent().remove(b)
for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Caption'):
    st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.09
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(23)
doc.styles['Title'].paragraph_format.space_after=Pt(8)
doc.styles['Subtitle'].font.size=Pt(11)
doc.styles['Subtitle'].paragraph_format.space_after=Pt(10)
doc.styles['Heading 1'].font.size=Pt(15)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(6)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(8)
doc.styles['Heading 2'].font.size=Pt(11.5)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(7)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(5)
doc.styles['Caption'].font.size=Pt(9)
doc.styles['Caption'].font.italic=False
doc.styles['Caption'].font.bold=False
doc.styles['Caption'].paragraph_format.line_spacing=1.04
doc.styles['Caption'].paragraph_format.space_after=Pt(8)

footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Phase One Visual Memory Study  |  ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

def p(text='',boldlead=None):
    q=doc.add_paragraph()
    if boldlead and text.startswith(boldlead):
        q.add_run(boldlead).bold=True;q.add_run(text[len(boldlead):])
    else:q.add_run(text)
    return q

new_page_pending=False
def h(text):
    global new_page_pending
    q=doc.add_heading(text,level=1)
    if new_page_pending:q.paragraph_format.page_break_before=True
    new_page_pending=False
    return q
def h2(text):return doc.add_heading(text,level=2)
def page():
    global new_page_pending
    new_page_pending=True

def image(path,caption,width=6.95):
    q=doc.add_paragraph();q.paragraph_format.space_after=Pt(4)
    q.paragraph_format.keep_with_next=True
    r=q.add_run();r.add_picture(str(path),width=Inches(width))
    r._r.xpath('.//wp:docPr')[0].set('descr',caption)
    doc.add_paragraph(caption,style='Caption')

def table(headers,rows,widths,small=9.5):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        z=OxmlElement('w:'+side);z.set(qn('w:val'),'single');z.set(qn('w:sz'),'5');z.set(qn('w:color'),'D9D9D9');borders.append(z)
    t._tbl.tblPr.append(borders)
    for i,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
            for side,value in [('top',75),('bottom',75),('left',95),('right',95)]:
                z=OxmlElement('w:'+side);z.set(qn('w:w'),str(value));z.set(qn('w:type'),'dxa');marg.append(z)
            tcpr.append(marg)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'333333' if i==0 else ('F4F4F4' if i%2==0 else 'FFFFFF'));tcpr.append(sh)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.05
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or (i>0 and (widths[j]>=4 or len(c.text)>45)) else WD_ALIGN_PARAGRAPH.CENTER
                for rr in para.runs:
                    rr.font.size=Pt(small);rr.bold=i==0;rr.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def cap(text):doc.add_paragraph(text,style='Caption')

# Additional figures for the user's test-by-test report structure.
probe_verification = read(ROOT/'runs/phase1_v3_dev/verification.json')
probedata = []
metrics_by_task = {}
for task in ('gate','motion'):
    source = ROOT/f'runs/phase1_v3_dev/{task}/reference/probes/metrics.json'
    assert sha(source)==probe_verification['actual_hpc_compact_sha256'][f'{task}/reference/probes/metrics.json']
    pr = read(source)
    metrics_by_task[task]=read(PN/f'evaluation/job_7575/{task}/evaluation/metrics.json')
    assert pr['cache_signature']==metrics_by_task[task]['cache_signature']
    variant='gate_cue' if task=='gate' else 'motion_two_frame'
    for split,label in splits:
        evidence=next(r for r in pr['results'] if r['task']==task and r['variant']==variant and r['split']==split)
        recent=next(r for r in pr['results'] if r['task']==task and r['variant']=='recent_only' and r['split']==split)
        probedata.append(dict(task=task,split=split,label=label,evidence=evidence['accuracy']['mean'],recent=recent['accuracy']['mean'],wilson=evidence['wilson_effective_pairs']))

def probe_plot(c):
    for panel,task in enumerate(('gate','motion')):
        l=55+panel*310;r=l+245;b=52;t=212
        yp=lambda z:b+z/100*(t-b)
        txt(c,(l+r)/2,240,task.title(),13,True,'center')
        for tick in (0,25,50,75,100):
            line(c,l,yp(tick),r,yp(tick),.85,.6,[3,3] if tick==75 else None)
            txt(c,l-8,yp(tick)-3,str(tick),9,False,'right')
        for i,(split,label) in enumerate(splits):
            row=next(r for r in probedata if r['task']==task and r['split']==split)
            x=l+40+i*83
            for off,key,color in [(-12,'recent',.8),(12,'evidence',.12)]:
                c.setFillGray(color);c.setStrokeGray(.2);c.rect(x+off-10,b,20,yp(row[key]*100)-b,fill=1,stroke=1)
            txt(c,x,b-17,['Ordinary','Long gap','Distractors'][i],9,False,'center')
        txt(c,(l+r)/2,13,'Classification accuracy  percent',10,False,'center')
    c.setFillGray(.8);c.rect(180,264,12,7,fill=1,stroke=1);txt(c,200,263,'Recent only',10)
    c.setFillGray(.12);c.rect(315,264,12,7,fill=1,stroke=1);txt(c,335,263,'Historical evidence',10)

probefig=figure('08_feature_readout',635,285,probe_plot)

swapdata=[]
for task in ('gate','motion'):
    for split,label in splits:
        for horizon in (3,6):
            r=next(r for r in metrics_by_task[task]['comparisons'] if r['split']==split and r['horizon']==horizon and r['seed']=='all_seeds_mean' and r['alternative']=='wrong_evidence')
            swapdata.append(dict(task=task,split=split,horizon=horizon,**r['toward_wrong_future']))

def swap_plot(c):
    maximum=max(r['high'] for r in swapdata if r['horizon']==6)*1.10
    for panel,task in enumerate(('gate','motion')):
        l=104+panel*310;r=l+180;b=40;t=207
        xp=lambda z:l+z/maximum*(r-l)
        txt(c,(l+r)/2,236,task.title(),13,True,'center')
        for tick in (0,.25,.5,.75,1):
            if tick>maximum:continue
            line(c,xp(tick),b,xp(tick),t,.85,.6);txt(c,xp(tick),b-15,f'{tick:.2f}',9,False,'center')
        for i,(split,label) in enumerate(splits):
            v=next(r for r in swapdata if r['task']==task and r['split']==split and r['horizon']==6);y=185-i*59
            txt(c,l-8,y-3,label,10,False,'right')
            line(c,xp(v['low']),y,xp(v['high']),y,0,1)
            for z in (v['low'],v['high']):line(c,xp(z),y-3,xp(z),y+3)
            c.setFillGray(0);c.circle(xp(v['mean']),y,3,fill=1,stroke=1)
        txt(c,(l+r)/2,6,'Shift toward the paired wrong future',9.5,False,'center')
    txt(c,320,266,'6 steps ahead  Three training runs averaged within each pair',11,True,'center')

swapfig=figure('09_history_swap',640,282,swap_plot)

harddata=[]
for task in ('gate','motion'):
    for horizon in (3,6):
        x=next(x for x in outcome['tasks'][task]['fixed_reference_checks'] if x['kind']=='critical_absolute_mse_reduction' and x['slice'].endswith(f'/h{horizon}/seedall_seeds_mean'))
        harddata.append(dict(task=task,horizon=horizon,reference=x['reference_privileged_mse'],candidate=x['candidate_privileged_mse'],**x['absolute_mse_reduction']))

def fixed_plot(c):
    ymax=.22
    for panel,task in enumerate(('gate','motion')):
        l=56+panel*310;r=l+240;b=48;t=226
        yp=lambda z:b+z/ymax*(t-b)
        txt(c,(l+r)/2,256,task.title()+'  '+('Distractors' if task=='gate' else 'Long occlusion'),12,True,'center')
        for tick in (0,.05,.1,.15,.2):
            line(c,l,yp(tick),r,yp(tick),.85,.6);txt(c,l-8,yp(tick)-3,f'{tick:.2f}',9,False,'right')
        for i,horizon in enumerate((3,6)):
            v=next(r for r in harddata if r['task']==task and r['horizon']==horizon);x=l+60+i*120
            for off,key,color in [(-18,'reference',.8),(18,'candidate',.1)]:
                c.setFillGray(color);c.setStrokeGray(.2);c.rect(x+off-14,b,28,yp(v[key])-b,fill=1,stroke=1)
                txt(c,x+off,yp(v[key])+8,f'{v[key]:.3f}',9,False,'center')
            txt(c,x,b-18,f'{horizon} steps ahead',10,False,'center')
        txt(c,(l+r)/2,8,'Feature MSE   Lower is better',10,False,'center')
    c.setFillGray(.8);c.rect(170,282,12,7,fill=1,stroke=1);txt(c,190,281,'Earlier predictor',10)
    c.setFillGray(.1);c.rect(335,282,12,7,fill=1,stroke=1);txt(c,355,281,'Current candidate',10)

fixedfig=figure('10_fixed_predictor',635,302,fixed_plot)

rankdata=summary['rankings']
def ranking_plot(c):
    for panel,task in enumerate(('gate','motion')):
        rows=[r for r in rankdata if r['task']==task and r['split']=='test_id']
        v=[r['gains']['feature_mse'] for r in rows]
        maximum=max(z['high'] for z in v)*1.15
        l=99+panel*310;r=l+185;b=47;t=277
        xp=lambda z:l+z/maximum*(r-l)
        txt(c,(l+r)/2,307,task.title(),13,True,'center')
        for k in range(5):
            z=maximum*k/4;line(c,xp(z),b,xp(z),t,.85,.6);txt(c,xp(z),b-15,f'{z:.3f}',9,False,'center')
        for i,rr in enumerate(sorted(rows,key=lambda a:(a['seed'],a['select_h']))):
            v=rr['gains']['feature_mse'];y=257-i*37
            txt(c,l-7,y-3,f"{rr['seed']}  {rr['select_h']} to {rr['test_h']}",10,False,'right')
            line(c,xp(v['low']),y,xp(v['high']),y,0,1)
            for z in (v['low'],v['high']):line(c,xp(z),y-3,xp(z),y+3)
            c.setFillGray(0);c.circle(xp(v['mean']),y,3,fill=1,stroke=1)
        txt(c,(l+r)/2,7,'Feature error saved versus random deletion',9,False,'center')
    txt(c,320,337,'Ordinary condition   Seed and selection time to scoring time',11,True,'center')

rankfig=figure('11_cross_horizon',640,350,ranking_plot)

def future_work(c):
    txt(c,7,202,'Retention policy study   NOT RUN',12,True)
    steps=[('Offline deletion\ncost targets',8),('Past-only\nscorer',214),('Streaming comparison\nagainst baselines',420)]
    for label,x in steps:box(c,x,130,185,45,label,1,11)
    arrow(c,197,152,208,152);arrow(c,403,152,414,152)
    txt(c,7,100,'Independent confirmation   NOT RUN',12,True)
    steps=[('Freeze protocol\nand resource plan',8),('New train and validation\nFit and lock models',214),('Untouched test data\nNo retuning',420)]
    for label,x in steps:box(c,x,27,185,45,label,1,11)
    arrow(c,197,49,208,49);arrow(c,403,49,414,49)

futurefig=figure('12_unrun_studies',620,225,future_work)

# The four requested main sections. Each completed test receives its own page.
def title_page(text,level=1):
    global new_page_pending
    q=doc.add_heading(text,level=level)
    q.paragraph_format.page_break_before=True
    new_page_pending=False
    return q

def sub(text):return doc.add_heading(text,level=3)
doc.styles['Heading 2'].font.size=Pt(13)
doc.styles['Heading 3'].font.name='Arial';doc.styles['Heading 3'].font.size=Pt(11)
doc.styles['Heading 3'].font.color.rgb=RGBColor(0,0,0)

doc.add_paragraph('Phase One Study of Bounded Visual Memory',style='Title')
doc.add_paragraph('Research goals  Experimental design  Individual tests  Findings and next tasks\nI-JEPA prototype  |  16 October 2026',style='Subtitle')
h('1  Research objective')
p('Overall objective. When a model cannot remember everything it has seen, which observations should it keep? Our proposal tests whether the extra prediction error caused by deleting a memory can teach a small retention policy to make better choices.',boldlead='Overall objective.')
p('Objective of this stage. Establish the prerequisites before learning a policy: useful evidence must survive visual encoding, the predictor must respond to that evidence, and equal-sized memory alternatives must produce informative differences in prediction cost. The original Phase One covers the first two questions; the deletion assay is a transition toward the later retention stage.',boldlead='Objective of this stage.')
image(taskfig,'Figure 1. Read left to right: earlier evidence, recent view, future outcome. Paired scenes A and B have the same recent view but different futures. The earlier Gate cue or Motion positions distinguish them. These are explanatory drawings, not sampled test frames or model predictions.',6.95)
table(['Question','Why it matters'],[
    ['Can the features preserve the clue?','A policy cannot keep useful information if encoding has already removed it.'],
    ['Can prediction use that clue?','Deletion cost is meaningful only when the predictor responds to what remains.'],
    ['Does the cost distinguish forgetting choices?','This is the proposed supervision signal for a future retention scorer.']
],[2.2,4.7],9.8)
p('The intended contribution is a better bounded-memory decision. A perfect general dynamics model is not the research objective. Current results concern I-JEPA and these controlled tasks; no learned retention-policy advantage, V-JEPA transfer or fresh confirmation has been demonstrated.')

title_page('2  Experimental design')
h2('2.1  Data and preprocessing')
p('We use generated, controlled visual trajectories rather than a public video benchmark. Each matched pair contains two episodes with identical recent images and action sequences, but different earlier evidence and later outcomes. Gate tests a transient cue; Motion tests pre-occlusion movement. Historical split names contain the word test, but these data are now development evidence because predictor repairs were informed by them.')
table(['Split per task','Matched pairs','Episodes per task','Purpose'],[
    ['Training','512','1024','Fit predictors and train-only preprocessing'],
    ['Validation','96','192','Select readout settings and predictor checkpoint'],
    ['Ordinary test condition','128','256','Conditions similar to training'],
    ['Long gap or occlusion','128','256','Stress longer missing-information intervals'],
    ['Distractor shift','128','256','Stress changed visual distractions'],
    ['Total per task','992','1984','Two tasks give 3968 episodes in total']
],[1.85,1,1.3,2.75],9.3)
p('Processing. Rendered 96 x 96 RGB frames pass through the pinned official I-JEPA image processor without a center crop. The frozen encoder is spatially pooled into four cells, producing a 5120-dimensional vector per frame. Mean and scale are fitted using training frames only; a train-only 128-dimensional basis reduces predictor input size. Prediction error is still scored in the complete normalized 5120-dimensional output.',boldlead='Processing.')
title_page('2.2  Understanding the three test conditions',2)
p('These are three versions of each task, not three different tasks. Both Gate and Motion are evaluated in all three. Each condition has its own 128 matched pairs. Within a pair, recent images and actions match, while earlier evidence distinguishes the two possible futures.')
sub('Ordinary means a familiar delay and familiar scenery')
p('Motion example: the model sees two observations of the moving object, then a shutter hides it. It must use that earlier motion and known actions to predict a later position. Ordinary already includes occlusion; it does not mean that the object stays visible. The last evidence frame is followed by 5 to 10 intervening steps before the 4-frame recent window begins. These settings match the training distribution.')
p('Gate example: a brief geometric cue appears and disappears. Later, the model must combine the old cue with the action to predict which lane the object takes. The question is whether history helps under familiar conditions.')
sub('Long gap means the same evidence must help after a longer delay')
p('The intervening gap increases to 20 to 32 steps. Gate must use a cue that disappeared longer ago; Motion must predict after a longer period of hidden movement. The task rules, recent-window length and evaluated future horizons remain the same. This tests delayed use of evidence, rather than a different prediction target.')
sub('Distractor shift means changed irrelevant scenery')
p('The ordinary gap and task rules are retained. Only the irrelevant decorative objects outside the arena change: 3 circles become 7 rectangles and the color palette changes. This does not change the movement rule or replace the correct cue. It tests whether unfamiliar visual distractions interfere with using the relevant evidence.')
image(conditionfig,'Timing guide. Read each row from left to right. The gap measures how long ago the clue occurred before the recent window. The prediction horizon measures how far ahead to forecast after that window. Distractor shift shares ordinary timing; its difference is the irrelevant scenery described above.',6.9)
sub('Why the same condition can be easy for one test and hard for another')
p('T1 directly reads the stored features of the earlier evidence frames. Waiting longer afterward does not erase those cached features, so long gap need not reduce T1 accuracy. T2 and T4 instead require a future forecast using the evidence, actions and elapsed time. They can fail even when T1 is perfect. Condition names describe the generated episodes; they do not guarantee that every test becomes harder by the same amount.')

title_page('2.3  Tests and the questions they answer',2)
table(['Test','Actual comparison','What it can establish'],[
    ['T1  Evidence readout','Historical feature versus recent-only feature readout','Did the frozen features retain the cue or movement direction?'],
    ['T2  History benefit','Two key memories versus recent observations alone','Does supplying historical evidence improve future prediction?'],
    ['T3  History replacement','Correct history versus the paired episode history','Does changing historical content change the predicted future?'],
    ['T4  Fixed reference','Current candidate versus an earlier frozen predictor','Did repair improve prediction against a stable reference?'],
    ['T5  Critical deletion','Delete a key frame versus delete nuisance content; both retain 2','Does deletion cost recognize useful evidence at equal capacity?'],
    ['T6  Cross horizon','Choose a deletion at 3 steps and score at 6, then reverse','Does the usefulness of the deletion choice transfer across time?'],
    ['T7  Policy advantage','A learned past-only scorer versus strong streaming baselines','Can a deployable policy exploit the signal?  Not run.'],
    ['T8  Fresh confirmation','Locked procedure evaluated with untouched data','Does the result repeat beyond inspected development data?  Not run.']
],[1.35,2.8,2.75],9.2)
cap('Legacy labels: Original prediction checks combine T1, T2 and T3. Additional fixed-reference checks are T4. The previously combined last row contains two different future studies, T7 and T8. They have no completed result.')

title_page('2.4  Structure and recorded data',2)
image(pipefig,'Figure 2. Images become stored features. A small predictor reads recent context, legal historical entries, timestamps and action continuations, then predicts future features. Simulator state is used by diagnostics and scoring, not supplied as predictor input. I-JEPA is frozen; no retention selector is present.',6.95)
table(['Block','What it does','What is recorded'],[
    ['Controlled data','Create paired Gate and Motion episodes','RGB, actions, timestamps, pair and split IDs, evidence times; state and labels reserved for diagnostics'],
    ['Frozen encoder and cache','Encode each observed frame once','Feature arrays, shape, precision, checkpoint and processor identity, per-file hashes'],
    ['Train-only preprocessing','Normalize features and fit the input basis','Training episode membership, mean, scale, basis and their hashes'],
    ['Recent reference and candidate','Predict features with different history conditions','Epoch and step, training loss, validation error, best checkpoint, optimizer and RNG evidence'],
    ['Prediction evaluation','Compare forecasts with targets and paired alternatives','Feature MSE, task error, pair margin, condition, seed, horizon, memory count and bytes'],
    ['Deletion evaluation','Fork the same parent into alternative removals','Removed and retained entry IDs, equal-capacity checks, feature and task losses per fork'],
    ['Statistical audit','Check pairing and recompute contrasts','128-pair aggregates, uncertainty intervals, x/y criterion counts and provenance ledger']
],[1.6,2.05,3.25],9.3)
sub('Shared settings')
p('The candidate starts with a frozen recent-only forecast and adds a learned correction from centered historical features. It reads every retained entry, weights their content, and uses elapsed time with observed feature differences to transport evidence forward. This is a prediction mechanism, not a trained memory-retention policy; no retrieval or selector chooses the stored entries.')
p('Recent context is 4 frames; the bounded history contains 2 real entries; future horizons are 3 and 6 steps. Full history is an additional information reference, not an equal-capacity competitor. The complete candidate experiment trained seeds 17, 29 and 43 for each task: six models, 120 epochs and 7680 steps each. Checkpoints were selected by the earliest minimum of validation MSE across four balanced conditions and both horizons.')
sub('Reading the numbers')
p('MSE means average squared error: lower is better. A positive error reduction means history helped. For prediction tests, 95 percent confidence intervals use 2000 paired bootstrap replicates of 128 base pairs; the deletion assay uses 10000 replicates. Siblings, deletion forks and seed means are not additional independent examples. An uncertainty interval crossing zero leaves the direction unresolved.')

title_page('3  Results for each test')
h2('3.1  T1  Can the features preserve the evidence')
p('Question. Before asking a predictor to use memory, did the visual representation preserve what the task needs?',boldlead='Question.')
p('Comparison. A simple linear readout was fitted on training features, selected on validation, and scored on each test condition. It reads the Gate cue or the two-frame Motion evidence. A recent-only readout is the control. These previously fitted probes are reused; no new fitting was performed for this report.',boldlead='Comparison.')
image(probefig,'Figure 3. Taller bars mean the clue is easier to decode. Black bars use historical evidence and gray bars use recent frames only. History reaches 100 percent in all six task-condition combinations; recent-only stays at 50 percent, as expected for the matched pairs. These are classification results, not future-prediction accuracies.',6.95)
table(['Condition','Gate history / recent','Motion history / recent'],[
    [label, f"{next(r for r in probedata if r['task']=='gate' and r['split']==sp)['evidence']:.0%} / 50%",f"{next(r for r in probedata if r['task']=='motion' and r['split']==sp)['evidence']:.0%} / 50%"] for sp,label in splits
],[2.3,2.3,2.3],10)
p('Coverage. Gate 3/3 pass; Motion 3/3 pass. Each uses 128 base pairs per condition. The frozen criterion required at least 75 percent accuracy and a Wilson lower bound above 50 percent; the observed historical readout lower bound is 97.1 percent.',boldlead='Coverage.')
p('Why exactly 50 and 100 percent. The original values are 0.5 and 1.0, not rounded approximations. Paired recent inputs are identical but their labels are opposite: the recent-only classifier gives the same answer to both, getting exactly one right. Historical inputs are explicitly chosen evidence frames, and a supervised linear readout has learned to decode them. It correctly classified all evaluated examples in these controlled conditions. This is an observed ceiling on a deliberately simple diagnostic, not guaranteed perfection on new data.',boldlead='Why exactly 50 and 100 percent.')
p('Meaning. The relevant discrete evidence survives this encoding and pooling pipeline, so there is something useful for memory to preserve. The Motion readout distinguishes movement direction; it does not establish exact speed recovery, future position accuracy or success after long occlusion. The splits were held out from probe fitting but have since been inspected during development.',boldlead='Meaning.')

title_page('3.2  T2  Does historical evidence improve prediction',2)
p('Question. Does the same predictor forecast better when given important past observations in addition to the identical recent context?',boldlead='Question.')
p('Comparison. Recent-only forecasts are compared with privileged two-memory forecasts. Full history provides an additional reference. A counted history-benefit check requires sufficient improvement and positive uncertainty bounds in both feature error and the task metric: Gate lane error or Motion position MSE. The original rules retain their absolute-gain fallback when the baseline is near zero.',boldlead='Comparison.')
image(gainfig,'Figure 4. Further right means a bigger reduction in prediction error from two key memories. Squares are 3-step forecasts; circles are 6-step forecasts. Horizontal lines show 95 percent uncertainty. A line crossing zero leaves the direction unresolved. The dashed 10 percent line is the material-effect benchmark. Means average the three seeds within each pair.',6.95)
table(['Task and condition','Recent only MSE','Full history MSE','Two-memory MSE','Two-memory gain at 6 steps'],[
    [task.title()+'  '+lab,*[f"{next(v for v in condition_mse if v['task']==task and v['split']==sp)[co]:.4f}" for co in ('recent_only','full_history','privileged_small_budget')],(lambda v:f"{v['mean']:.1f}% [{v['low']:.1f}, {v['high']:.1f}]")(next(v for v in gainrows if v['task']==task and v['split']==sp and v['horizon']==6))]
    for task in ('gate','motion') for sp,lab in splits
],[1.95,.95,.95,.95,2.1],8.9)
p('Coverage. Gate 23/24 pass; Motion 8/24 pass. The 24 comparisons are 3 conditions x 2 horizons x 4 summaries: three seeds and their mean. These are not 24 independent datasets.',boldlead='Coverage.')
p('Meaning. Gate shows broad history benefit, with an inconclusive ordinary 6-step task benefit for seed 43. Motion improves substantially in ordinary conditions, but the longer-gap 6-step feature gain is only 1.5 percent with an interval from -1.5 to 4.6 percent. The harder conditions limit the generalization claim.',boldlead='Meaning.')

title_page('3.3  T3  Does changing history change the forecast',2)
p('Question. Is the model using the content of the old observation, or merely responding to the presence of extra memory entries?',boldlead='Question.')
p('Comparison. Keep recent images and actions fixed, but replace relevant memory features with those from the paired episode. For Gate, swap the cue associated with the other outcome; for Motion, swap the evidence for the other movement. Measure how much the predicted feature moves toward the paired alternative future.',boldlead='Comparison.')
image(swapfig,'Figure 5. Points further right show a larger shift toward the other episode future after swapping history. Lines show paired 95 percent intervals. Positive shifts in both tasks show that historical content affects prediction. The score is a normalized feature-distance margin change, not a percentage accuracy. The plot shows 6-step, seed-averaged results.',6.95)
table(['Task and condition','Shift at 3 steps','Shift at 6 steps and 95 percent CI'],[
    [task.title()+'  '+label,
     f"{next(r for r in swapdata if r['task']==task and r['split']==sp and r['horizon']==3)['mean']:.3f}",
     (lambda r:f"{r['mean']:.3f} [{r['low']:.3f}, {r['high']:.3f}]")(next(r for r in swapdata if r['task']==task and r['split']==sp and r['horizon']==6))]
    for task in ('gate','motion') for sp,label in splits
],[2.5,1.25,3.15],9.4)
p('Coverage. Gate 24/24 pass; Motion 24/24 pass. The original criterion requires a mean margin shift of at least 0.05 and a positive lower confidence bound. The 24 summaries have the same condition, horizon and seed structure as T2.',boldlead='Coverage.')
p('Meaning. The predictor responds to which history it receives. This supports historical-content use, even where overall future prediction is weak. It does not establish an intervention on the physical environment, a reliable deletion ranking or a learned retention policy. Those are different questions.',boldlead='Meaning.')

title_page('3.4  T4  Does repair improve a fixed predictor reference',2)
p('Question. Did the current candidate improve against an unchanged earlier reference, rather than simply change its own comparison baseline?',boldlead='Question.')
p('Comparison. There are 24 candidate two-memory versus frozen recent-only comparisons, plus 8 comparisons where both earlier and current predictors receive two memories. The latter target Gate distractor shift and Motion long occlusion. They require a positive paired error reduction; the recent-only comparisons retain the original material-benefit rule.',boldlead='Comparison.')
image(fixedfig,'Figure 6. Shorter bars mean lower prediction error. Gray is the earlier predictor; black is the current candidate, both with the same two key memories. Gate improves in the designated hard condition. Motion becomes worse than this reference under long occlusion. Bars average three seeds; the table gives paired intervals for the difference.',6.95)
table(['Task and horizon','Reference MSE','Candidate MSE','Reference minus candidate and CI'],[
    [r['task'].title()+f"  {r['horizon']} steps",f"{r['reference']:.4f}",f"{r['candidate']:.4f}",f"{r['mean']:.4f} [{r['low']:.4f}, {r['high']:.4f}]"] for r in harddata
],[1.55,1.25,1.25,2.85],9.3)
table(['Fixed-reference comparison','Gate','Motion'],[
    ['Candidate with 2 memories versus frozen recent only','24/24 pass','8/24 pass'],
    ['Both predictors with 2 memories on the hard condition','8/8 pass','0/8 pass'],
    ['Combined descriptive coverage','32/32 pass','8/32 pass']
],[4.1,1.4,1.4],9.3)
p('Meaning. The repair has broad fixed-reference support for Gate. For Motion, ordinary gains coexist with a clear long-occlusion regression against the earlier two-memory predictor. This is a specific limitation of the candidate; it does not imply that Motion evidence is absent from the encoder or that every memory mechanism is ineffective.',boldlead='Meaning.')

title_page('3.5  T5  Does deletion cost identify useful memories',2)
p('Question. At the same storage capacity, does losing critical evidence harm prediction more than losing a nuisance observation?',boldlead='Question.')
p('Comparison. Each parent has three real historical frames. Gate contains the cue and two post-cue nuisance frames; Motion contains start, end and a shutter frame. Remove each candidate in turn, always leaving two entries and the same recent context. Two fixed parent-time variants are averaged within each base pair. No predictor or policy is retrained.',boldlead='Comparison.')
image(forestfig,'Figure 7. Further right means keeping that memory is more useful at 6 steps in ordinary conditions. Black points measure feature error; white points measure lane or position error. Lines show 95 percent uncertainty. Motion start and end memories help in every seed. Gate seed 43 is near zero and uncertain. The dashed line marks the 10 percent mean-benefit requirement.',6.95)
table(['Task','Ordinary','Long gap','Distractor shift','Total'],[
    ['Gate','4/6 pass','0/6 pass','2/6 pass','6/18 pass'],
    ['Motion','12/12 pass','12/12 pass','12/12 pass','36/36 pass']
],[1,1.4,1.4,1.7,1.4],9.8)
p('A counted contrast requires at least 10 percent relative benefit from retaining the critical observation, plus a positive paired interval, for both feature and task error. Gate has one critical role; Motion has two, explaining the different denominators. Each contrast uses 128 base pairs, averaging siblings and parent variants first.')
p('Meaning. Motion supplies consistent relative deletion value in the examined conditions. At ordinary h6, preserving start reduces feature error by 76.3 to 77.8 percent and preserving end by 28.3 to 30.5 percent across seeds. Gate provides partial support, with both ordinary exceptions coming from seed 43. Useful deletion value and accurate absolute forecasting are distinct capabilities.',boldlead='Meaning.')

title_page('3.6  T6  Does a deletion choice transfer across future times',2)
p('Question. Is a deletion choice useful beyond the single future time used to choose it?',boldlead='Question.')
p('Comparison. Select the removal with lowest feature error at 3 steps, then score it at 6 steps against uniform-random deletion. Reverse the times and repeat. Each counted comparison requires positive paired uncertainty bounds for both feature and task advantage. This uses future outcomes to select the removal; it is an oracle diagnostic, not a past-only policy.',boldlead='Comparison.')
image(rankfig,'Figure 8. All points and their uncertainty lines are right of zero in these ordinary-condition examples: the cross-time choice saves feature error compared with random deletion. A label such as 17 3 to 6 means seed 17, choose at step 3, score at step 6. The two tasks use different horizontal scales. Full condition coverage is below.',6.95)
table(['Condition','Gate coverage','Motion coverage','Mean choice agreement Gate / Motion'],[
    [label,'6/6 pass','6/6 pass',f"{statistics.mean(r['rank_agreement'] for r in rankdata if r['task']=='gate' and r['split']==sp):.1%} / {statistics.mean(r['rank_agreement'] for r in rankdata if r['task']=='motion' and r['split']==sp):.1%}"] for sp,label in splits
]+[['Total','18/18 pass','18/18 pass','Agreement is descriptive, not a readiness criterion']],[1.55,1.15,1.15,3.05],9.3)
p('The 18 comparisons per task are 3 conditions x 3 seeds x 2 selection directions. Choice agreement describes how often the best removal agrees across horizons, averaged over the tested seeds and directions; near ties are averaged, not resolved using evidence labels.')
p('Meaning. The deletion signal is not confined to one scoring time. Gate has strong cross-horizon coverage despite weaker critical-versus-nuisance contrasts in T5: the two tests ask different questions. This transfer does not prove usefulness beyond the tested 3- and 6-step futures, or that a scorer can choose well without seeing the future.',boldlead='Meaning.')

title_page('3.7  T7  Can a learned policy beat strong baselines',2)
p('Status. Not run: no retention scorer has been trained and no learned irreversible streaming policy has been evaluated.',boldlead='Status.')
p('The intended study would train a scorer on offline deletion costs, restrict its inputs to past and present information, and compare it with FIFO, reservoir sampling and other strong baselines under the same predictor and memory budget. Its result would establish deployable policy value rather than privileged evidence value.')
table(['Existing descriptive Motion comparison at ordinary h6','Feature MSE','Position MSE'],[[r['label'],f"{r['feature']:.6f}",f"{r['task']:.6f}"] for r in bdata],[4,1.45,1.45],9.2)
p('The available one-deletion baseline data are preparation for T7, not a T7 result. Reject-newest reaches feature MSE 0.116061, close to the future-feature oracle at 0.115305, only 0.65 percent lower. A learned advantage is therefore not established; more varied memory states are needed to test whether learning offers useful headroom.')
h2('3.8  T8  Does the finding repeat on untouched data')
p('Status. Not run: fresh independent confirmation data have not been opened or generated.',boldlead='Status.')
p('The registered confirmation plan requires fresh data, new train-only preprocessing, reference and candidate models for both tasks and all three seeds, and locked model selection before test outcomes are opened. Existing development weights evaluated on another test set would not satisfy that plan. The complete 12-model confirmation and its compute guard must be frozen before execution.')
image(futurefig,'Figure 9. Planned studies, not completed experiments. The top row asks whether a past-only learned scorer beats streaming baselines. The bottom row asks whether locked findings repeat with fresh data and newly fitted models. These answer two distinct questions; neither is included in the reported pass counts.',6.95)

title_page('4  Overall results analysis and next tasks')
table(['Evidence category','Gate','Motion'],[
    ['T1  Historical evidence readable','3/3 pass','3/3 pass'],
    ['T2  History improves prediction','23/24 pass','8/24 pass'],
    ['T3  Forecast responds to historical content','24/24 pass','24/24 pass'],
    ['T4  Fixed-reference coverage','32/32 pass','8/32 pass'],
    ['T5  Critical deletion value','6/18 pass; ordinary 4/6','36/36 pass; ordinary 12/12'],
    ['T6  Cross-horizon deletion value','18/18 pass','18/18 pass'],
    ['T7  Learned-policy advantage','Not run','Not run'],
    ['T8  Fresh confirmation','Not run','Not run']
],[3.5,1.7,1.7],9.3)
h2('4.1  What has been achieved')
p('The pipeline preserves the relevant discrete evidence, and both predictors respond to its content. Gate has broad prediction coverage. Motion provides consistent deletion value and cross-horizon transfer, even though its long-occlusion forecast quality remains limited. This is meaningful mechanism evidence for a bounded-memory project, with task- and condition-specific limits.')
p('Counts are not votes of equal weight or probabilities of success. The legacy original total is T1 + T2 + T3: Gate 50/51 and Motion 35/51. It excludes other guards. In particular, Gate full-history predictions at long-gap h6 triggered mean-level degradation warnings for seeds 17 and 29, with intervals spanning no change. The earlier strict all-checks acceptance record remains unmet; this descriptive summary does not change it.')
image(baselinefig,'Figure 10. A useful deletion signal does not automatically leave room for a learned policy. In these constructed Motion pools, rejecting the newest entry is already close to the future-informed oracle. Bars show mean feature error; dots show the three training seeds. This is a one-deletion diagnostic, not a streaming-policy evaluation.',6.6)
title_page('4.2  Next tasks and unresolved questions',2)
table(['Priority','Bounded next task','What would be learned'],[
    ['1','Motion pilot with varied memory states and strong simple rules','Can a past-only scorer exploit deletion costs beyond an easy reject-newest heuristic?'],
    ['2','Targeted Gate seed and full-history diagnostics','Which conditions make its useful-history signal unstable?'],
    ['3','Freeze a separate confirmation and resource plan before fresh data','Do the selected claims repeat without further adaptation?'],
    ['Scope','Treat any V-JEPA study as a separate backbone validation','Does the mechanism transfer beyond this I-JEPA implementation?']
],[.75,3.2,2.95],9.2)
p('No further experiment is initiated by this report. Existing predictors and caches were reused for the deletion assay, adding zero GPU time. Settled proposal-validation and repair GPU use is 12723 seconds, or 3 h 32 min 03 s, within the approved 8 h cap. Learned retention, streaming policy success and new-data confirmation remain future evidence.')
h2('4.3  Sources and verification')
refs=[
    '[1] Assran et al. Self-Supervised Learning from Images with a Joint-Embedding Predictive Architecture. CVPR 2023. https://arxiv.org/abs/2301.08243',
    '[2] Lee et al. Set Transformer. ICML 2019. https://proceedings.mlr.press/v97/lee19d.html',
    '[3] Zeng et al. Are Transformers Effective for Time Series Forecasting. AAAI 2023. https://arxiv.org/abs/2205.13504',
    'Evidence: complete predictor run 7574, evaluation 7575 and independent audit 7576; verified inherited feature probes; deletion assay 7590 and its independent 180-comparison audit. The predictor evaluation contains 69120 scalar rows and deletion 55296 rows, not that many independent episodes. Compact evidence hashes were verified.'
]
for ref in refs:
    q=p(ref);q.paragraph_format.space_after=Pt(4)
    for rr in q.runs:rr.font.size=Pt(8.5)
h2('Conclusion')
p('Phase One has produced substantial but uneven evidence. Historical information survives encoding, the predictors use it, and Motion deletion costs consistently distinguish useful evidence in the tested pools. Gate deletion stability and Motion hard-condition prediction remain limitations. These findings support a scoped next-stage feasibility study, while leaving learned-policy advantage, V-JEPA transfer and independent confirmation unproven.')

path=OUT/'Phase1_Structured_Academic_Report.docx'
doc.save(path)
sources=[NATIVE/'summary.json',NATIVE/'independent_audit.json',NATIVE/'baselines.json',PN/'evaluation/job_7575/outcome.json',PN/'final_native_verification.json',PRED/'report/all_64_fixed_checks.csv',PN/'evaluation/job_7575/gate/evaluation/metrics.json',PN/'evaluation/job_7575/motion/evaluation/metrics.json',ROOT/'runs/phase1_v3_dev/verification.json',ROOT/'runs/phase1_v3_dev/gate/reference/probes/metrics.json',ROOT/'runs/phase1_v3_dev/motion/reference/probes/metrics.json']
data={'probes':probedata,'prediction_gains':gainrows,'condition_MSE':condition_mse,'history_swap':swapdata,'fixed_reference_hard_conditions':harddata,'deletion_contrasts':summary['contrasts'],'cross_horizon':rankdata,'descriptive_baselines':bdata}
(OUT/'figure_data.json').write_text(json.dumps(data,indent=2)+'\n')
manifest={'report':str(path),'sha256':sha(path),'expected_sections':['Research objective','Experimental design','Results for each test','Overall results analysis and next tasks'],'expected_pages':14,'all_pages_visually_verified':False,'sources':{str(x.relative_to(ROOT)):sha(x) for x in sources},'original_criteria_and_outcomes_unchanged':True,'new_model_execution':False,'new_GPU_seconds':0,'T7_and_T8_not_run':True}
(OUT/'report_manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
print(path)

